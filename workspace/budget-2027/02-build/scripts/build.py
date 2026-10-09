"""Build the collated 2027 COO budget workbook.

Usage:
  python3 -I build.py COO_BOOK SD_BOOK OUT_XLSX EXTREFS_JSON [NOTES_JSON] [--mark-frozen-zeros]

- Base = COO book (loaded with formulas).
- COO FO / FE are replaced by SD FO / FE (cell-by-cell copy: value/formula + style).
- SD sheets that FO/FE depend on (directly or indirectly) are copied the same way.
- Formulas that reference an external workbook ("[n]Sheet") are replaced by the SD
  book's cached value; every such cell is logged to EXTREFS_JSON.
- If NOTES_JSON is given, a "Collation Notes" sheet is inserted at the front.
- --mark-frozen-zeros (v2): every frozen external-link cell whose stored value is not an error
  (e.g. cached zeros) gets a cell comment FROZEN_NOTE, so it is not read as data.
  Without the flag the output is the v1 build.
Inputs are only read, never written.
"""
import sys, re, json, copy
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

FLAGS = {a for a in sys.argv[1:] if a.startswith("--")}
ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
COO_PATH, SD_PATH, OUT_PATH, EXT_JSON = ARGS[0:4]
NOTES_JSON = ARGS[4] if len(ARGS) > 4 else None
MARK_FROZEN = "--mark-frozen-zeros" in FLAGS
FROZEN_NOTE = "stale external-link value \u2013 not real data"

REPLACE = ["FO", "FE"]
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_\.\[\]]+))!")
EXT_REF = re.compile(r"\[\d+\]")
# Google Sheets' "#ERROR!" (formula parse error) has no Excel equivalent.
# An unresolvable external link shows as #REF! in Excel, so it is stored as #REF!.
ERROR_MAP = {"#ERROR!": "#REF!"}
EXCEL_ERRORS = {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}


def formula_text(v):
    if isinstance(v, ArrayFormula):
        return v.text
    if isinstance(v, str) and v.startswith("="):
        return v
    return None


def sheet_deps(ws):
    deps = set()
    for row in ws.iter_rows():
        for c in row:
            f = formula_text(c.value)
            if f:
                for m in SHEET_REF.finditer(f):
                    s = (m.group(1) or m.group(2)).replace("''", "'")
                    if not EXT_REF.search(s):
                        deps.add(s)
    return deps


def closure(wb, roots):
    seen, todo = [], list(roots)
    while todo:
        s = todo.pop(0)
        if s in seen or s not in wb.sheetnames:
            continue
        seen.append(s)
        todo.extend(sorted(sheet_deps(wb[s]) - set(seen)))
    return seen


def copy_sheet(src, dst, src_cached, ext_log):
    # cells
    for row in src.iter_rows():
        for c in row:
            if isinstance(c, MergedCell):
                continue
            v = c.value
            f = formula_text(v)
            d = dst.cell(row=c.row, column=c.column)
            if f and EXT_REF.search(f):
                cached = src_cached[c.coordinate].value
                stored = ERROR_MAP.get(cached, cached) if isinstance(cached, str) else cached
                ext_log.append({"sheet": src.title, "cell": c.coordinate, "formula": f,
                                "sd_cached": cached if not hasattr(cached, "isoformat") else cached.isoformat(),
                                "stored": stored if not hasattr(stored, "isoformat") else stored.isoformat()})
                d.value = stored
                if isinstance(stored, str) and stored in EXCEL_ERRORS:
                    d.data_type = "e"
            elif isinstance(v, ArrayFormula):
                d.value = ArrayFormula(v.ref, v.text)
            else:
                d.value = v
            if c.has_style:
                d.font = copy.copy(c.font)
                d.fill = copy.copy(c.fill)
                d.border = copy.copy(c.border)
                d.alignment = copy.copy(c.alignment)
                d.number_format = c.number_format
                d.protection = copy.copy(c.protection)
            if c.hyperlink:
                d.hyperlink = copy.copy(c.hyperlink)
            if c.comment:
                d.comment = copy.copy(c.comment)
    # merged ranges
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    # column widths / hidden
    for key, dim in src.column_dimensions.items():
        nd = dst.column_dimensions[key]
        nd.width = dim.width
        nd.hidden = dim.hidden
        nd.outlineLevel = dim.outlineLevel
        nd.min, nd.max = dim.min, dim.max
        if dim.has_style:
            nd.font = copy.copy(dim.font); nd.fill = copy.copy(dim.fill)
            nd.number_format = dim.number_format; nd.alignment = copy.copy(dim.alignment)
            nd.border = copy.copy(dim.border)
    for key, dim in src.row_dimensions.items():
        nd = dst.row_dimensions[key]
        nd.height = dim.height
        nd.hidden = dim.hidden
        nd.outlineLevel = dim.outlineLevel
    # sheet-level settings
    dst.sheet_format = copy.copy(src.sheet_format)
    dst.sheet_properties.tabColor = copy.copy(src.sheet_properties.tabColor)
    dst.sheet_properties.outlinePr = copy.copy(src.sheet_properties.outlinePr)
    dst.freeze_panes = src.freeze_panes
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    dst.sheet_view.zoomScale = src.sheet_view.zoomScale
    dst.sheet_state = src.sheet_state
    # data validations
    for dv in src.data_validations.dataValidation:
        dst.add_data_validation(copy.deepcopy(dv))
    # conditional formatting (rule.dxf is bound on load; writer re-registers it)
    for cf in src.conditional_formatting:
        for rule in cf.rules:
            dst.conditional_formatting.add(str(cf.sqref), copy.deepcopy(rule))


def add_notes_sheet(wb, notes):
    ws = wb.create_sheet("Collation Notes", 0)
    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="DDEBF7")
    r = 1
    ws.cell(r, 1, notes["title"]).font = Font(bold=True, size=14); r += 2
    for sec in notes["sections"]:
        ws.cell(r, 1, sec["heading"]).font = Font(bold=True, size=12); r += 1
        for line in sec.get("lines", []):
            c = ws.cell(r, 1, line)
            if isinstance(line, str) and line.startswith("="):
                c.data_type = "s"
            r += 1
        if sec.get("table"):
            t = sec["table"]
            for j, h in enumerate(t["header"], 1):
                c = ws.cell(r, j, h); c.font = bold; c.fill = hdr_fill
            r += 1
            fcols = set(t.get("formula_cols", []))   # 1-based columns whose "=..." text is a live formula
            for row in t["rows"]:
                for j, v in enumerate(row, 1):
                    c = ws.cell(r, j, v)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(v, str) and v.startswith("="):
                        if j in fcols:
                            c.number_format = "#,##0;(#,##0)"
                        else:
                            # LibreOffice turns inline text that starts with "=" into a formula on import,
                            # so quoted formula text is stored without its leading "=".
                            c.value = v[1:]
                            c.data_type = "s"
                    elif isinstance(v, str) and v in EXCEL_ERRORS:
                        c.data_type = "s"              # error code shown as text, not as an error value
                    elif isinstance(v, (int, float)):
                        c.number_format = "#,##0;(#,##0)"
                r += 1
        r += 1
    widths = notes.get("col_widths", [14, 10, 26, 60, 60, 50, 16, 16])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return ws


def main():
    coo = openpyxl.load_workbook(COO_PATH)
    sd = openpyxl.load_workbook(SD_PATH)
    sd_v = openpyxl.load_workbook(SD_PATH, data_only=True)

    for s in REPLACE:
        assert s in coo.sheetnames and s in sd.sheetnames, s
    carry = closure(sd, REPLACE)
    support = [s for s in sd.sheetnames if s in carry and s not in REPLACE]  # keep SD order
    clash = [s for s in support if s in coo.sheetnames]
    assert not clash, f"sheet name clash: {clash}"

    ext_log = []
    # replace FO / FE in place (same tab position)
    for s in REPLACE:
        idx = coo.sheetnames.index(s)
        coo.remove(coo[s])
        dst = coo.create_sheet(s, idx)
        copy_sheet(sd[s], dst, sd_v[s], ext_log)
    # supporting sheets appended after PE, in SD order
    for s in support:
        dst = coo.create_sheet(s)
        copy_sheet(sd[s], dst, sd_v[s], ext_log)

    marked = []
    if MARK_FROZEN:
        for e in ext_log:
            st = e["stored"]
            if isinstance(st, str) and st in EXCEL_ERRORS:
                continue
            c = coo[e["sheet"]][e["cell"]]
            assert c.comment is None, f"{e['sheet']}!{e['cell']} already has a comment"
            c.comment = Comment(FROZEN_NOTE, "collation v2")
            marked.append(f"{e['sheet']}!{e['cell']}")

    if NOTES_JSON:
        with open(NOTES_JSON) as fh:
            add_notes_sheet(coo, json.load(fh))

    # exactly one selected tab, active = first sheet
    for ws in coo.worksheets:
        ws.sheet_view.tabSelected = False
    coo.active = 0
    coo.worksheets[0].sheet_view.tabSelected = True

    coo.save(OUT_PATH)
    with open(EXT_JSON, "w") as fh:
        json.dump({"carried": carry, "support": support, "external_cells": ext_log,
                   "frozen_marked": marked, "frozen_note": FROZEN_NOTE if MARK_FROZEN else None}, fh, indent=1, default=str)
    print("carried:", carry)
    print("external-ref cells replaced by cached values:", len(ext_log), "| frozen non-error cells commented:", len(marked))
    print("sheets:", coo.sheetnames)


if __name__ == "__main__":
    main()
