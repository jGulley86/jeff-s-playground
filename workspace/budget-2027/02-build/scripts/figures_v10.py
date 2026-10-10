"""v10 figures: every number the v10 notes and Collation Notes quote, computed from the workbooks (no typed figure).

Usage:
  python3 -I figures_v10.py V9_XLSX V10_XLSX SD_BOOK CN9_JSON SC_DIR OUT_JSON

  V9_XLSX   v9 workbook (verified; cached values)
  V10_XLSX  v10 workbook (stage 2 or final; cached values = LibreOffice recalculation, patched)
  SD_BOOK   Service_Delivery_PL_worksheets_v2.xlsx (Logistics&WH Payroll cached values, for the copy check)
  CN9_JSON  v9 Collation Notes extracted by cnotes_v6.py (v9 caveat effects that v10 does not touch are carried from it;
            each carried item is listed with the reason its inputs are unchanged)
  SC_DIR    LibreOffice recalculations of scratch copies (scen_v8.py): base (v10 as is), cur (B131 = 0), idle (B135 = 0),
            idlecur (B131 = 0, B135 = 0), q4 (B132 = 1), q4cur (B132 = 1, B131 = 0), b13 (B13 = 1), v9cur (v9 with B131 = 0)

No name is read or written: people are '<tab> r<row> <title>' and seats are seat codes.
"""
import sys, json, re, itertools
import openpyxl

V9, V10, SD, CN9, SC, OUT = sys.argv[1:7]
LW = "Logistics&WH Payroll"
EPS = 1e-6
w9 = openpyxl.load_workbook(V9, data_only=True)
w10 = openpyxl.load_workbook(V10, data_only=True)
f10 = openpyxl.load_workbook(V10)
sd = openpyxl.load_workbook(SD, data_only=True)
cn9 = json.load(open(CN9))
SCN = {n: openpyxl.load_workbook(f"{SC}/{n}.xlsx", data_only=True) for n in ("base", "cur", "idle", "idlecur", "q4", "q4cur", "b13", "v9cur")}
M27 = range(13, 25); P27 = range(21, 33); S27 = range(6, 18)
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
F = {}


def n(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def row27(ws, r, cols=P27):
    return [n(ws.cell(r, c).value) for c in cols]


def contrib(w):
    return n(w["COO P&L"]["AG132"].value)


# ------------------------------------------------------------------ headline
HL = [("rev", 16), ("cogs", 63), ("gp", 64), ("opex", 116), ("noi", 117), ("other", 131), ("ni", 132)]
F["headline"] = {k: {"n": n(w10["COO P&L"].cell(r, 14).value), "ag": n(w10["COO P&L"].cell(r, 33).value),
                     "ag9": n(w9["COO P&L"].cell(r, 33).value), "n9": n(w9["COO P&L"].cell(r, 14).value), "row": r} for k, r in HL}
F["headline"]["dep"] = {"n": sum(n(w10["COO P&L"].cell(r, 14).value) for r in range(19, 23)),
                        "ag": sum(n(w10["COO P&L"].cell(r, 33).value) for r in range(19, 23)),
                        "ag9": sum(n(w9["COO P&L"].cell(r, 33).value) for r in range(19, 23)),
                        "n9": sum(n(w9["COO P&L"].cell(r, 14).value) for r in range(19, 23)), "row": "19-22"}
C10 = F["headline"]["ni"]["ag"]; C9 = F["headline"]["ni"]["ag9"]
F["c10"], F["c9"] = C10, C9
assert abs(contrib(SCN["base"]) - C10) < EPS, "base recalculation differs from the workbook"

# ------------------------------------------------------------------ Logistics&WH tab in the workbook
lw = w10[LW]; lwf = f10[LW]; lsd = sd[LW]
mx = 0.0; nonnum = 0; ncell = 0
for row in lsd.iter_rows():
    for c in row:
        a, b = lw[c.coordinate].value, c.value
        if b is None and a is None:
            continue
        ncell += 1
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            mx = max(mx, abs(a - b))
        elif a != b:
            nonnum += 1
errs = sum(1 for row in lw.iter_rows() for c in row if isinstance(c.value, str) and re.fullmatch(r"#(REF!|NAME\?|VALUE!|DIV/0!|N/A|NUM!|NULL!)", c.value))
nform = sum(1 for row in lwf.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
S = lambda r: sum(n(lw.cell(r, c).value) for c in S27)
tax, ben = n(lw["B5"].value), n(lw["B6"].value)
F["lw"] = {"cells_compared": ncell, "max_diff_vs_sd_cache": mx, "nonnum_diffs": nonnum, "errors": errs, "formulas": nform,
           "total_2027": n(lw["S47"].value), "sal": S(43), "tax_amt": S(44), "ben_amt": S(45), "nh": S(46),
           "tax": tax, "ben": ben, "check_B60": n(lw["B60"].value), "B59": n(lw["B59"].value),
           "log01_sal": S(34), "log01_cost": S(34) * (1 + tax + ben), "mh01_sal": S(38), "mh01_cost": S(38) * (1 + tax + ben),
           "nh_per_hire": n(lw["B10"].value), "hires_2027": S(24)}
L = F["lw"]
L["ex_sal"] = L["sal"] - L["log01_sal"]; L["ex_tax"] = L["tax_amt"] - L["log01_sal"] * tax; L["ex_ben"] = L["ben_amt"] - L["log01_sal"] * ben
L["ex_total"] = L["ex_sal"] + L["ex_tax"] + L["ex_ben"] + L["nh"]
assert abs(L["total_2027"] - L["log01_cost"] - L["ex_total"]) < 1e-6
# seats (labels without the names in brackets)
seats = []
for r in range(14, 23):
    lab = re.sub(r"\s*\(.*?\)\s*$", "", str(lw.cell(r, 1).value)).strip()
    code = lab.split()[0]
    sal27 = sum(n(lw.cell(r + 20, c).value) for c in S27)
    start = lw.cell(r, 21).value
    nh27 = L["nh_per_hire"] if (start is not None and start.year == 2027) else 0.0
    seats.append({"row": r, "seat": lab, "code": code, "team": lw.cell(r, 20).value, "start": start.strftime("%b-%y") if start else "on payroll",
                  "annual": n(lw.cell(r, 24).value), "sal_2027": sal27, "cost_2027": sal27 * (1 + tax + ben) + nh27, "nh_2027": nh27,
                  "hc_dec27": n(lw.cell(r, 17).value), "in_co": code != "LOG-01"})
F["lw"]["seats"] = seats
assert abs(sum(s["cost_2027"] for s in seats if s["in_co"]) - L["ex_total"]) < 1e-6
F["lw"]["hc_ex_by_month"] = [n(lw.cell(23, c).value) - n(lw.cell(14, c).value) for c in S27]
F["lw"]["cost_ex_by_month"] = [n(lw.cell(47, c).value) - n(lw.cell(34, c).value) * (1 + tax + ben) for c in S27]

# ------------------------------------------------------------------ HC-CO: new row 18 and roster totals
h10, h9 = w10["HC-CO"], w9["HC-CO"]
d6, d7 = n(h10["D6"].value), n(h10["D7"].value)
new_sal = [n(h10.cell(18, c).value) for c in M27]
F["newco"] = {"row": 18, "annual": n(h10["D18"].value), "start": h10["E18"].value.strftime("%b-%y"), "end": h10["F18"].value.strftime("%b-%y"),
              "raise": n(h10["G18"].value), "raise_month": h10["D8"].value.strftime("%b-%y"),
              "sal_2027": sum(new_sal), "tax_2027": sum(new_sal) * d6, "ben_2027": sum(new_sal) * d7,
              "cost_2027": sum(new_sal) * (1 + d6 + d7), "per_month_jan": new_sal[0] * (1 + d6 + d7),
              "sep_dec_26": [h10.cell(18, c).value for c in range(9, 13)], "co_d6": d6, "co_d7": d7,
              "Z18": n(h10["Z18"].value)}
assert abs(F["newco"]["Z18"] - F["newco"]["sal_2027"]) < EPS
people10 = sum(1 for r in range(14, 58) if n(h10.cell(r, 26).value) > 0)
people9 = sum(1 for r in range(14, 58) if n(h9.cell(r, 26).value) > 0)
F["hcco"] = {"people9": people9, "people10": people10, "Z70_9": n(h9["Z70"].value), "Z70_10": n(h10["Z70"].value),
             "Z67_10": n(h10["Z67"].value), "Z67_9": n(h9["Z67"].value)}

# ------------------------------------------------------------------ CO 6610 / 6630 / 6640 before -> after
co9, co10 = w9["CO"], w10["CO"]
F["co"] = {}
for r, gl in ((67, "6610"), (68, "6630"), (69, "6640")):
    hc_r = 67 + (r - 67)
    F["co"][gl] = {"v9": n(co9.cell(r, 33).value), "v10": n(co10.cell(r, 33).value),
                   "hc_v10": n(h10.cell(hc_r, 26).value), "hc_v9": n(h9.cell(hc_r, 26).value),
                   "months_v10": row27(co10, r), "months_v9": row27(co9, r)}
F["co"]["6610"]["lw"] = L["ex_sal"] + L["nh"]; F["co"]["6630"]["lw"] = L["ex_tax"]; F["co"]["6640"]["lw"] = L["ex_ben"]
F["co"]["6610"]["new"] = F["newco"]["sal_2027"]; F["co"]["6630"]["new"] = F["newco"]["tax_2027"]; F["co"]["6640"]["new"] = F["newco"]["ben_2027"]
for gl in ("6610", "6630", "6640"):
    e = F["co"][gl]
    e["check"] = e["v10"] - e["hc_v10"] - e["lw"]                      # LOG-01 not double counted: must be 0
    e["check_v9"] = e["hc_v10"] - e["new"] - e["v9"]                   # HC-CO existing 5 = v9 typed values: must be 0
    assert abs(e["check"]) < 1e-6 and abs(e["check_v9"]) < 1e-6, (gl, e["check"], e["check_v9"])
F["co"]["total"] = {k: sum(F["co"][g][k] for g in ("6610", "6630", "6640")) for k in ("v9", "v10", "lw", "new", "hc_v10")}
F["co"]["subtotals"] = {lab: {"v9": n(co9.cell(r, 33).value), "v10": n(co10.cell(r, 33).value)} for lab, r in
                        (("Total - 6000 - Labor Expense - SG&A (row 73)", 73), ("Total - Expense (row 116)", 116), ("Net Income (row 132)", 132))}

# ------------------------------------------------------------------ PO: the B131 base without MH-01 (PO HC r20)
hp = w10["HC-PO"]; pd6, pd7 = n(hp["D6"].value), n(hp["D7"].value)
r20 = [n(hp.cell(20, c).value) for c in M27]; r16 = [n(hp.cell(16, c).value) for c in M27]
po_cur10 = SCN["cur"]["PO"]; po_cur9 = SCN["v9cur"]["PO"]
F["po"] = {"r20_title": hp["C20"].value, "r16_title": hp["C16"].value, "pd6": pd6, "pd7": pd7,
           "r20_sal": sum(r20), "r20_cost": sum(r20) * (1 + pd6 + pd7), "r16_cost": sum(r16) * (1 + pd6 + pd7),
           "base_v9": sum(n(po_cur9.cell(r, 33).value) for r in (33, 34, 35)),
           "base_v10": sum(n(po_cur10.cell(r, 33).value) for r in (33, 34, 35)),
           "by_gl_v9": {gl: n(po_cur9.cell(r, 33).value) for gl, r in (("5110", 33), ("5130", 34), ("5140", 35))},
           "by_gl_v10": {gl: n(po_cur10.cell(r, 33).value) for gl, r in (("5110", 33), ("5130", 34), ("5140", 35))},
           "booked_v10": sum(n(w10["PO"].cell(r, 33).value) for r in (33, 34, 35)),
           "hcpo_Z70": n(hp["Z70"].value), "hcpo_people": sum(1 for r in range(14, 58) if n(hp.cell(r, 26).value) > 0)}
assert abs(F["po"]["base_v9"] - F["po"]["base_v10"] - F["po"]["r20_cost"]) < 1e-6, "PO base must drop exactly r20 at PO HC rates"
assert abs(F["po"]["booked_v10"]) < 1e-9
# the MH-01 seat in the model vs PO HC r20 (same person; different rates)
F["po"]["mh01_model_vs_hc_sal_diff"] = L["mh01_sal"] - F["po"]["r20_sal"]

# ------------------------------------------------------------------ scenarios
ds = lambda w, c: n(w["Depreciation Schedule"][c].value)
sc = {}
for k in ("base", "cur", "idle", "idlecur", "q4", "q4cur", "b13", "v9cur"):
    w = SCN[k]
    sc[k] = {"c": contrib(w), "capex": ds(w, "B236"), "dep": ds(w, "R195"), "po5100": sum(n(w["PO"].cell(r, 33).value) for r in (33, 34, 35)),
             "ppe": ds(w, "B238"), "cip": ds(w, "B239")}
F["sc"] = sc
D = {k: sc[k]["c"] - C10 for k in sc}
F["d"] = D
F["d_idle_under_cur"] = sc["idlecur"]["c"] - sc["cur"]["c"]
F["ramp"] = sc["cur"]["capex"]
F["policy_effect"] = sc["cur"]["po5100"] + sc["cur"]["capex"] - sc["base"]["dep"]
F["policy_effect_v9_check"] = (sc["v9cur"]["po5100"] + sc["v9cur"]["capex"] - ds(w9, "R195"))

# ------------------------------------------------------------------ new caveat quantities
F["burden"] = {"ex_sal": L["ex_sal"], "model": L["ex_sal"] * (tax + ben), "co_rates": L["ex_sal"] * (d6 + d7)}
F["burden"]["diff"] = F["burden"]["model"] - F["burden"]["co_rates"]
pe = w10["HC-PE"]
co_hc_sal = n(h10["Z67"].value)
F["i37"] = {"co_hc": co_hc_sal * (0.0765 - d6), "co_lw": L["ex_sal"] * (0.0765 - tax), "pe": n(pe["Z67"].value) * (0.0765 - n(pe["D6"].value))}
F["i37"]["total"] = F["i37"]["co_hc"] + F["i37"]["co_lw"] + F["i37"]["pe"]
v9_i37 = n(w9["HC-CO"]["Z67"].value) * (0.0765 - n(w9["HC-CO"]["D6"].value)) + n(w9["HC-PE"]["Z67"].value) * (0.0765 - n(w9["HC-PE"]["D6"].value))
F["i37"]["v9_reproduced"] = v9_i37
cur9 = n(w10["Depreciation Schedule"]["B214"].value)
F["overlap9"] = {"current_model": cur9, "per_head": cur9 / 9, "dep_rate": sc["base"]["dep"] / sc["base"]["capex"]}
F["overlap9"]["dep_effect"] = F["overlap9"]["per_head"] * F["overlap9"]["dep_rate"]
F["i41"] = {"r16_cost": F["po"]["r16_cost"], "v9_both": F["po"]["r16_cost"] + F["po"]["r20_cost"]}

# ------------------------------------------------------------------ caveat table (v10)
cav9 = {i: r for i, r in enumerate(cn9["sections"][3]["table"]["rows"])}
assert cav9[0][0].startswith("Decision #10 scope: current 9") and cav9[10][0] == "I-10" and cav9[22][0] == "I-37"
assert abs(cav9[4][3] + F["i41"]["v9_both"]) < 0.01, "v9 I-41 not reproduced"
assert abs(cav9[22][3] + v9_i37) < 0.01, "v9 I-37 CO/PE not reproduced"
assert abs(cav9[0][3] - (sc["v9cur"]["c"] - C9)) < 0.01, "v9 sensitivity not reproduced"
assert abs(cav9[28][3] - D["b13"]) < 0.01, "A-14 (B13 = 1) changed"
CARRY = {3: "labour block unchanged (diff: 0 cells on Depreciation Schedule / Prod Payroll)",
         5: "Prod Payroll and the labour block unchanged", 6: "Prod Payroll and PO HC rates unchanged",
         10: "FO / FE unchanged", 12: "Fld Ops Payroll / Fld Maint Budget unchanged", 13: "Fld Eng Budget unchanged",
         14: "PO depreciation rows unchanged", 15: "PO depreciation rows unchanged", 16: "Depreciation Schedule unchanged",
         17: "PO row 19 unchanged", 18: "PO row 19 unchanged", 19: "Depreciation Schedule unchanged", 20: "FE unchanged",
         21: "FE unchanged", 23: "FO / FE unchanged", 24: "PE unchanged", 25: "Fld Eng Budget unchanged",
         26: "Fld Eng Budget unchanged", 27: "Depreciation Schedule unchanged", 29: "revenue unchanged"}
F["carried"] = {str(i): {"item": cav9[i][0], "direction": cav9[i][1], "effect": cav9[i][3], "what": cav9[i][2], "src": cav9[i][5],
                         "why": CARRY[i]} for i in CARRY}
F["i26_text"] = cav9[11][2]
F["rev_breakeven_pct"] = C10 / F["headline"]["rev"]["ag"]

# ------------------------------------------------------------------ range and sign flip
DOWN = {"sens": ("the sensitivity (current 9 expensed)", D["cur"]), "idle": ("idle-month labour expensed (I-40)", D["idle"]),
        "i41": ("I-41 (7-overlap): Material Handling Team Lead stays expensed", -F["i41"]["r16_cost"]),
        "i10": ("I-10", cav9[10][3]), "i31": ("I-31 shortfall", cav9[14][3]), "i05": ("I-05 burden capitalised", cav9[5][3]),
        "o7": ("rate base of the current 9 (O7)", cav9[6][3])}
# I-05 burden at the B131 = 0 depreciation rate: v9's lower-bound-2 component (computed by figures_v9 from the labour block,
# which v10 does not touch), read from the v9 Collation Notes
m = re.search(r"I-05 burden capitalised, at the B131 = 0 rate \(([\d,]+)\)", " ".join(cn9["sections"][3]["table"].get("after", [])))
assert m, "v9 lower bound 2 I-05 component not found"
i05_cur = -float(m.group(1).replace(",", ""))
F["i05_cur"] = i05_cur
lb1_items = ["i41", "i10", "i31", "i05", "o7", "idle"]
F["lb1"] = {"items": [[DOWN[k][0], DOWN[k][1]] for k in lb1_items]}
F["lb1"]["value"] = C10 + sum(DOWN[k][1] for k in lb1_items)
lb2 = [["sensitivity: current 9 expensed", D["cur"]], ["I-10", cav9[10][3]], ["I-31 shortfall", cav9[14][3]],
       ["I-05 burden capitalised, at the B131 = 0 rate", i05_cur], ["idle-month labour expensed, at B131 = 0", F["d_idle_under_cur"]]]
F["lb2"] = {"items": lb2, "value": C10 + sum(v for _, v in lb2)}
F["upper"] = C10 + cav9[16][3]
F["v9_range"] = {"lb1": None, "lb2": None, "upper": None}
mm = re.search(r"Range \(v9\).*?\): \(([\d,]+)\)\. Lower bound 2.*?\): \(([\d,]+)\)\. Upper end, I-32 in full only: ([\d,]+)",
               " ".join(cn9["sections"][3]["table"]["after"]))
F["v9_range"] = {"lb1": -float(mm.group(1).replace(",", "")), "lb2": -float(mm.group(2).replace(",", "")), "upper": float(mm.group(3).replace(",", ""))}
singles = sorted(([k, DOWN[k][0], DOWN[k][1], C10 + DOWN[k][1]] for k in DOWN), key=lambda x: x[3])
F["singles"] = singles
F["singles_flip"] = [s for s in singles if s[3] < 0]
pairs = []
for a, b in itertools.combinations([k for k in DOWN if C10 + DOWN[k][1] >= 0], 2):
    v = C10 + DOWN[a][1] + DOWN[b][1]
    if v < 0:
        pairs.append([DOWN[a][0], DOWN[b][0], v])
F["pairs_flip"] = sorted(pairs, key=lambda x: x[2])
F["pair_sens_idle"] = sc["idlecur"]["c"]
json.dump(F, open(OUT, "w"), indent=1, default=str)
# recon-lite for paylist_v7.amounts(): the two Logistics&WH seats held by people on an HC roster (LOG-01, MH-01)
json.dump({"lw": {"seats": [{"annual": s["annual"], "cost_2027": s["cost_2027"], "overlap": s["code"] in ("LOG-01", "MH-01")} for s in seats]}},
          open(OUT[:-5] + "_recon.json", "w"), indent=1)
print(f"figures_v10: contribution {C10:,.2f} (v9 {C9:,.2f}); CO 6610-6640 {F['co']['total']['v9']:,.2f} -> {F['co']['total']['v10']:,.2f}; "
      f"logistics in CO {L['ex_total']:,.2f}; PO B131 base {F['po']['base_v9']:,.2f} -> {F['po']['base_v10']:,.2f}; "
      f"range {F['lb1']['value']:,.0f} / {F['lb2']['value']:,.0f} to {F['upper']:,.0f}; single flips {len(F['singles_flip'])}, pairs {len(F['pairs_flip'])}")
