"""v14 standard suite (self-contained; expected values are the verifier-confirmed v13 figures - v14 changes none).
No clock, no file properties, no dates (F1 rule).

  1. Headline and department figures (constants): COO P&L AG132/N132/AG16, CO 6610-6640 + 7700/8600, PE 6610-6640,
     PO 5110-5140 = 0 and 6610-6640, Depreciation Schedule R195, HC Z totals.
  2. COO P&L = sum of departments on every SUM(FO,PO,CO,FE,PE) cell (<= 1e-6).
  3. Zero error caches (deliverable + its recalculation); zero '#REF!' in formula texts AND in defined names;
     exactly 6 historical '#REF!' text mentions on Collation Notes (cells E343, D364, D365, E365).
  4. Recalculation self-consistency: a LibreOffice full recalc of v14 reproduces its cache (numeric <= 1e-6).
  5. Deliverable = v13 values outside the Collation Notes sheet (<= 1e-6), so every v13 figure/scenario holds.
  6. Collation Notes: C21:F28 live formulas identical to v13's (32); H78:H81 formulas identical to v13's and cache
     AND recalc to 0 (|x| <= 0.01); CONFIDENTIAL banner (A3); the v14 corrections landed (spot checks).
  7. Package: no shared formulas.
  8. Structure spot checks carried from v13 (HC-PO r20 blanked; PO!U33 live basis; PE/PO!U67 live links; CO!U67
     subtracts rows 34/36/38/39).

Usage:
  python3 -I suite_v14.py V14_XLSX V14_RECALC_XLSX V13_XLSX OUT_JSON
Exit 1 on any failure.
"""
import sys, json, re, zipfile
import openpyxl

V14, V14RC, V13, OUT = sys.argv[1:5]
TOL = 1e-6
ERR = {"#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!"}
isnum = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
fails, rep = [], {}


def check(ok, msg):
    print(("PASS " if ok else "FAIL ") + msg)
    if not ok:
        fails.append(msg)


v = openpyxl.load_workbook(V14, data_only=True)
f = openpyxl.load_workbook(V14)
vr = openpyxl.load_workbook(V14RC, data_only=True)
o13 = openpyxl.load_workbook(V13)
v13 = openpyxl.load_workbook(V13, data_only=True)
n = lambda x: float(x) if isnum(x) else 0.0

# 1. figures (verifier-confirmed v13 constants)
EXP = [("COO P&L", "AG132", -646753.85), ("COO P&L", "N132", -899222.05), ("COO P&L", "AG16", 15563645.54),
       ("CO", "AG67", 1032260.69), ("CO", "AG68", 56264.89), ("CO", "AG69", 129782.64),
       ("CO", "AG98", 24000), ("CO", "AG106", 78000),
       ("PE", "AG67", 1571694.96), ("PE", "AG68", 56132.55), ("PE", "AG69", 138674.30),
       ("PO", "AG67", 125937.50), ("PO", "AG68", 4580.21), ("PO", "AG69", 25988.80),
       ("PO", "AG33", 0), ("PO", "AG34", 0), ("PO", "AG35", 0),
       ("Depreciation Schedule", "R195", 104416.68),
       ("HC-CO", "Z67", 860846.29), ("HC-PO", "Z67", 600953.60), ("HC-FE", "Z67", 892313.53),
       ("HC-PE", "Z67", 1571694.96)]
for sh, ad, exp in EXP:
    got = n(v[sh][ad].value)
    check(abs(got - exp) <= 0.01, f"{sh}!{ad} = {got:,.2f} (expected {exp:,.2f})")
rep["AG132"] = n(v["COO P&L"]["AG132"].value)

# 2. roll-up
pat = re.compile(r"^=SUM\(FO!([A-Z]+\d+),PO!\1,CO!\1,FE!\1,PE!\1\)$")
cnt, mx = 0, 0.0
for row in f["COO P&L"].iter_rows():
    for c in row:
        if isinstance(c.value, str):
            m = pat.match(c.value.replace(" ", ""))
            if m:
                ad = m.group(1)
                s = sum(n(v[d][ad].value) for d in ("FO", "PO", "CO", "FE", "PE"))
                mx = max(mx, abs(n(v["COO P&L"][c.coordinate].value) - s)); cnt += 1
check(cnt > 2000 and mx <= TOL, f"COO P&L = FO+PO+CO+FE+PE on {cnt} cells (max dev {mx:.2e})")
rep["rollup"] = {"cells": cnt, "max_dev": mx}

# 3. error caches + #REF! in formulas and names; the 6 historical text mentions
err = sum(1 for ws in v.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value in ERR)
err_rc = sum(1 for ws in vr.worksheets for row in vr[ws.title].iter_rows() for c in row if isinstance(c.value, str) and c.value in ERR)
fref = sum(1 for ws in f.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith("=") and "#REF!" in c.value)
nref = sum(1 for k, d in f.defined_names.items() if "#REF!" in str(d.value))
check(err == 0 and err_rc == 0 and fref == 0 and nref == 0,
      f"0 error caches (v14 {err}, recalc {err_rc}); 0 #REF! in formulas ({fref}) and names ({nref}, of {len(f.defined_names)})")
mention_cells = {}
for row in f["Collation Notes"].iter_rows():
    for c in row:
        if isinstance(c.value, str) and not c.value.startswith("=") and "#REF!" in c.value:
            mention_cells[c.coordinate] = c.value.count("#REF!")
check(sum(mention_cells.values()) == 6 and set(mention_cells) == {"E343", "D364", "D365", "E365"},
      f"6 historical '#REF!' text mentions on Collation Notes ({mention_cells})")

# 4. recalc self-consistency + 5. deliverable = v13 outside Collation Notes
def cmp_books(a, b, skip_cn):
    mxd, cells, bad = 0.0, 0, 0
    for ws in a.worksheets:
        if skip_cn and ws.title == "Collation Notes":
            continue
        wb2 = b[ws.title]
        for row in ws.iter_rows():
            for c in row:
                x, y = c.value, wb2.cell(c.row, c.column).value
                if isnum(x) and isnum(y):
                    mxd = max(mxd, abs(x - y)); cells += 1
                elif (x if x != "" else None) != (y if y != "" else None):
                    bad += 1
    return mxd, cells, bad

mxd, cells, bad = cmp_books(v, vr, False)
check(mxd <= TOL and bad == 0, f"recalc reproduces the cache: {cells} numeric cells, max {mxd:.2e}, non-numeric diffs {bad}")
mxd2, cells2, bad2 = cmp_books(v13, v, True)
check(mxd2 <= TOL and bad2 == 0, f"deliverable = v13 values outside Collation Notes ({cells2} cells, max {mxd2:.2e}, other {bad2})")

# 6. Collation Notes
old_live = {c.coordinate: c.value for row in o13["Collation Notes"].iter_rows(min_row=21, max_row=28)
            for c in row if isinstance(c.value, str) and c.value.startswith("=")}
new_live = {c.coordinate: c.value for row in f["Collation Notes"].iter_rows(min_row=21, max_row=28)
            for c in row if isinstance(c.value, str) and c.value.startswith("=")}
check(old_live == new_live and len(new_live) == 32, "Collation Notes C21:F28 live formulas identical to v13 (32)")
h_old = [o13["Collation Notes"][f"H{r}"].value for r in (78, 79, 80, 81)]
h_new = [f["Collation Notes"][f"H{r}"].value for r in (78, 79, 80, 81)]
hv = [abs(n(v["Collation Notes"][f"H{r}"].value)) for r in (78, 79, 80, 81)]
hvr = [abs(n(vr["Collation Notes"][f"H{r}"].value)) for r in (78, 79, 80, 81)]
check(h_old == h_new and max(hv) <= 0.01 and max(hvr) <= 0.01,
      f"H78:H81 de-dup checks identical to v13 and 0 (cache max {max(hv):.2e}, recalc max {max(hvr):.2e})")
cn = f["Collation Notes"]
check(isinstance(cn["A3"].value, str) and "CONFIDENTIAL" in cn["A3"].value, "CONFIDENTIAL banner present")
SPOT = [("A33", "11 hires/moves (4 CO, 5 PE, 1 PO, 1 FE"), ("E17", "(1,205,750.10)"), ("A32", "(3,095,044)"),
        ("C43", "166,414"), ("C60", "CO HC 20,906 + logistics 1,415 + PE 64,102 + PM 5,054 = 91,477"),
        ("C61", "3,036,558"), ("F357", "(210,023) booked"), ("F377", "(91,477) CO/PE/PM"),
        ("D380", "sensitivity (558,996) -> (1,205,750)"), ("D381", "Three branches"),
        ("H381", "three-way"), ("A255", "twice in Sep-Nov-26"), ("G90", "net +38,342.16"),
        ("B857", "co-edit the H78:H80 check formulas"), ("A873", "11 hires/moves"),
        ("A876", "Changes from v13 (v14)")]
for ad, sub in SPOT:
    val = cn[ad].value
    check(isinstance(val, str) and sub in val, f"v14 text landed: {ad} contains {sub!r}")

# 7. package
z = zipfile.ZipFile(V14)
shared = sum(z.read(nm).decode("utf-8", "ignore").count('t="shared"') for nm in z.namelist() if nm.startswith("xl/worksheets/"))
check(shared == 0, "0 shared formulas in the package")

# 8. structure (carried from v13)
hcpo = f["HC-PO"]
check(all(hcpo[f"{c}20"].value in (None, "") for c in "BCD"), "HC-PO r20 blanked (B/C/D empty)")
check("SUM('HC-PO'!M14:M33)" in str(f["PO"]["U33"].value), "PO!U33 live 8-staff basis")
check(str(f["PE"]["U67"].value) == "='HC-PE'!M67", "PE!U67 live roster link")
check(str(f["PO"]["U67"].value) == "='HC-PO'!M41", "PO!U67 live PM link")
co_u67 = str(f["CO"]["U67"].value)
check(all(t in co_u67 for t in ("F34", "F36", "F38", "F39", "F46")), "CO!U67 subtracts rows 34/36/38/39 + hire cost")

rep["fails"] = fails
json.dump(rep, open(OUT, "w"), indent=1)
print(f"suite_v14: {'ALL GREEN' if not fails else str(len(fails)) + ' FAILURES'}")
sys.exit(1 if fails else 0)
