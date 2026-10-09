"""v5: decide which cached values change, from a full LibreOffice recalculation, limited to the static downstream set.

Usage:
  python3 -I patches_v5.py V4_XLSX STAGE1_LO_XLSX DEPS_JSON SHEET CELL VALUE OUT_PATCH_JSON OUT_NOISE_JSON

- STAGE1_LO_XLSX: LibreOffice full recalculation of stage 1 (= v4 with only SHEET!CELL = VALUE).
- DEPS_JSON: deps_v5.py on v4 from SHEET!CELL (every formula cell that can depend on the input).
- Patch list = downstream cells (excluding Collation Notes, which is regenerated) whose recalculated value differs
  from v4. The value written is LibreOffice's exact <v> text.
- Every other cell (not downstream): recalculated value vs v4 is recorded as floating-point noise. Any non-numeric
  difference, or a numeric difference > 1e-6, stops the script: it would mean the static trace missed a dependency.
"""
import sys, json, re, zipfile, collections
from xml.sax.saxutils import escape
import openpyxl

V4, LO1, DEPS, SH, CELL, VAL, OUT, NOISE = sys.argv[1:9]
VAL = float(VAL); VAL = int(VAL) if VAL.is_integer() else VAL
TOL = 1e-6
deps = json.load(open(DEPS))
assert [list(x) for x in deps["seeds"]] == [[SH, CELL]], deps["seeds"]
down = {tuple(x) for x in deps["cells"]}

a = openpyxl.load_workbook(V4, data_only=True)
b = openpyxl.load_workbook(LO1, data_only=True)
assert a.sheetnames == b.sheetnames
assert b[SH][CELL].value == VAL and a[SH][CELL].value is None, (a[SH][CELL].value, b[SH][CELL].value)

# raw <v> text per cell from the LibreOffice package
z = zipfile.ZipFile(LO1)
wbx = z.read("xl/workbook.xml").decode(); rel = z.read("xl/_rels/workbook.xml.rels").decode()


def raw_values(sheet):
    nm = escape(sheet, {'"': "&quot;"})
    rid = re.search(rf'<sheet name="{re.escape(nm)}"[^>]*r:id="(rId\d+)"', wbx).group(1)
    part = "xl/" + re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rel).group(1)
    x = z.read(part).decode()
    return dict(re.findall(r'<c r="([A-Z]+\d+)"[^>]*?(?<!/)>(?:<f[^>]*>[^<]*</f>|<f[^>]*/>)?<v>([^<]*)</v>', x))


isnum = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
patches = []; noise = collections.Counter(); noise_max = 0.0; noise_ex = []; bad = []
down_same = 0; down_noise_only = 0
raw_cache = {}
for s in a.sheetnames:
    if s == "Collation Notes":
        continue
    wa, wb_ = a[s], b[s]
    mr, mc = max(wa.max_row, wb_.max_row), max(wa.max_column, wb_.max_column)
    for row in wb_.iter_rows(min_row=1, max_row=mr, max_col=mc):
        for c in row:
            va, vb = wa[c.coordinate].value, c.value
            if (s, c.coordinate) == (SH, CELL):
                continue
            if va == vb:
                if (s, c.coordinate) in down: down_same += 1
                continue
            if (s, c.coordinate) in down:
                if not (isnum(va) and isnum(vb)):
                    bad.append([s, c.coordinate, str(va), str(vb), "downstream non-numeric change"]); continue
                if s not in raw_cache: raw_cache[s] = raw_values(s)
                raw = raw_cache[s][c.coordinate]
                assert float(raw) == float(vb), (s, c.coordinate, raw, vb)
                patches.append([s, c.coordinate, raw])
                if abs(va - vb) <= TOL: down_noise_only += 1
            else:
                if isnum(va) and isnum(vb) and abs(va - vb) <= TOL:
                    noise[s] += 1; noise_max = max(noise_max, abs(va - vb))
                    if len(noise_ex) < 5: noise_ex.append([s, c.coordinate, va, vb])
                else:
                    bad.append([s, c.coordinate, str(va), str(vb), "change outside the downstream set"])
res = {"input": [SH, CELL, VAL], "patched": len(patches),
       "patched_by_sheet": dict(collections.Counter(p[0] for p in patches)),
       "downstream_total": len(down), "downstream_unchanged": down_same, "downstream_noise_only_patched": down_noise_only,
       "downstream_collation_notes": sum(1 for s, _ in down if s == "Collation Notes"),
       "noise_outside": {"cells": sum(noise.values()), "max_abs": noise_max, "by_sheet": dict(noise), "examples": noise_ex},
       "bad": bad}
json.dump({"input": [SH, CELL, VAL], "values": patches}, open(OUT, "w"), indent=0)
json.dump(res, open(NOISE, "w"), indent=1, default=str)
print(f"patches_v5: {len(patches)} cached values to patch {res['patched_by_sheet']}; downstream {len(down)} "
      f"(unchanged {down_same}); noise outside downstream: {sum(noise.values())} cells, max {noise_max:.1e}; bad {len(bad)}")
if bad:
    print(bad[:10]); sys.exit(1)
