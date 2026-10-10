"""v9 build: one rule for Sep-Dec 26 labour (B132 = 0) and an idle-month labour switch (critic v8 B1/B2). Package patch of v8.

Usage:
  python3 -I build_v9.py stage1 V8_XLSX OUT_STAGE1_XLSX OUT_SPEC_JSON
      v8 package with the v9 edits only (cached values of edited formula cells are placeholders):
        Depreciation Schedule (labour block, rows 128+ only):
          1. B132 1 -> 0 (Sep-Dec 26 model labour is a 2026 period cost; the 2026 books expense 5100); label A132.
          2. New switch row 135: B135 'idle-month labour: 1 = roll forward (capitalise), 0 = expense', default 1. (The
             handoff suggested B133; B133 already holds the build-to-go-live lag, which 25 formulas read, so the switch
             goes in the first free row of the input block.)
          3. Rows 153 / 155 (labour waiting for the next build) x B135; new section 11 (rows 218+): idle labour expensed
             by month and its split to 5110 / 5130 / 5140 (Prod Payroll's own tax and benefit shares of the month); the
             Sep-26 builds (2026 go-lives) depreciated in 2027 when B132 = 1, so the Sep-26 labour follows the same switch
             and rule as Oct-Dec 26; CapEx and Dec-27 balance-sheet memo (in service, net; not yet in service = CIP).
          4. Rows 193 / 194 (+ Sep-26 builds' labour depreciation, 0 when B132 = 0); checks R196 and B206 extended;
             B213 restricted to the 2027 vintages.
          5. Labels: A129, A132, A153, A155, A160, A161, A196, A203, A204, A206, A208, A209, A213 ('inside the PO run
             rate' wording removed; 'CIP / inventory' -> 'PP&E not yet placed in service (CIP)').
        PO: U33:AF35 (5110-5140) = (1 - B131) x the v7 amount + the idle labour expensed to that GL (0 while B135 = 1);
            AI33:AI35 notes.
      OUT_SPEC_JSON lists every edited / new cell with the old and new content, and the row map.
  python3 -I build_v9.py final STAGE1_XLSX PATCH_JSON OUT_XLSX [NOTES_JSON]
      = build_v8.final (cached values from patches_v9.py; optional Collation Notes from NOTES_JSON).
Every other zip part is copied byte for byte. sharedStrings.xml and styles.xml are untouched (inline strings; existing styles).
"""
import sys, re, json, zipfile, os
from xml.sax.saxutils import escape, unescape
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_v8
from build_v8 import part_of, cell_xml, colkey, S_TXT, S_SEC, S_HDR, S_IN, S_DATE, S_UNIT, S_BOLD, S_MON, MCOL, Y26, Y27, PP

DS = "Depreciation Schedule"
SW_IDLE = 135
RUN_RATE_WORDS = ("run rate", "run-rate")

# ------------------------------------------------------------------ edits of existing Depreciation Schedule cells
LABELS = {
    "A129": ("Prod Payroll monthly cost (rows 41, 62, 79; model unchanged) is capitalised into the units built that month: split lift / tray "
             "(supervisors pro rata to the lift : tray labour), pooled by build month, divided by the units built, added to the go-live vintage "
             "(B133 months after the build) and depreciated with the section 1 lives, salvage and start rule. A month with no units built rolls "
             "into the next build month when B135 = 1, or is expensed (PO 5110-5140, section 11) when B135 = 0. Sep-Dec 26 labour: B132 (v9 "
             "booked 0: 2026 labour is a 2026 period cost; the 2026 books expense 5100). Section 9 feeds PO rows 20/21 (U:AF). PO 5110-5140 "
             "2027 = (1 - B131) x the v7 amount (+ idle labour expensed) = 0 (capitalised)."),
    "A132": ("Sep-Dec 26 model labour (Sep-26 to Dec-26 builds) capitalised: 0 = no (booked v9: 2026 labour is a 2026 period cost; the 2026 books "
             "expense 5100); 1 = capitalise all of it (sensitivity only: the 2026 PO forecast already expenses Oct-Dec 26 and the Q4-26 model "
             "hires are on no roster)"),
    "A153": "Lift labour waiting for the next lift build (B135 = 1 rolls a month with no lifts built forward; B135 = 0 expenses it, row 221)",
    "A155": "Tray labour waiting for the next tray build (B135 = 1 rolls a month with no trays built forward; B135 = 0 expenses it, row 222)",
    "A160": "Labour in builds that go live in 2028: PP&E not yet placed in service (CIP) at Dec-27, no 2027 depreciation",
    "A161": ("Labour in builds that go live in 2026 (Sep-26 builds, Nov-26 go-lives): 0 when B132 = 0 (2026 labour is a 2026 period cost; the "
             "2026 books expense 5100); with B132 = 1 depreciated in 2027 (rows 231-232)"),
    "A196": "Check: rows 193:194 less all labour vintage rows and rows 231:232 (must be 0; a non-0 means a GL is not 5011/5012)",
    "A203": "  in builds for 2027 go-lives (Feb-Oct 27 builds; Jan-27 rolls into Feb-27 when B135 = 1)",
    "A204": "  in builds for 2028 go-lives (Nov-Dec 27 builds): PP&E not yet placed in service (CIP) at Dec-27",
    "A206": "Check: row 202 less rows 203:205 less the idle labour expensed in 2027 (R223) (must be 0)",
    "A208": "  in Jan-Feb 27 go-lives (Nov-Dec 26 builds; 0 when B132 = 0)",
    "A209": ("  in Nov-26 go-lives (Sep-26 builds): 0 when B132 = 0 (2026 labour is a 2026 period cost; the 2026 books expense 5100); with "
             "B132 = 1 depreciated in 2027 (R233)"),
    "A213": "PP&E in service at Dec-27: capitalised labour in the 2027 go-live vintages, net of their 2027 depreciation (row 210 - (row 212 - R233))",
}
FORMULAS = {"R196": "R195-SUM(R165:R189)-R231-R232", "B206": "B202-SUM(B203:B205)-R223", "B213": "B210-(B212-R233)"}


def ds_wait(c, i, cap, units, wait):
    prev = "" if i == 0 else f"+{MCOL[i - 1]}{wait}"
    return f"IF({c}{units}>0,0,$B${SW_IDLE}*({c}{cap}{prev}))"


def new_rows():
    rows = {}; R = {}

    def put(r, col, kind, content, style):
        rows.setdefault(r, []).append((col, kind, content, style))

    def months(r, fn, style=S_MON, cols=MCOL):
        for c in cols:
            put(r, c, "f", fn(c), style)

    put(SW_IDLE, "A", "text", "Idle-month labour (months with no units built): 1 = roll forward into the next build month (capitalise; booked); "
        "0 = expense in the month (PO 5110-5140, section 11)", S_TXT)
    put(SW_IDLE, "B", "num", 1, S_IN); put(SW_IDLE, "C", "text", "switch", S_TXT)
    R["sw_idle"] = SW_IDLE

    put(218, "A", "text", "11. SWITCH LOGIC (added in v9): idle-month labour (B135), Sep-26 builds under B132, CapEx and Dec-27 balance sheet", S_SEC)
    put(219, "A", "text", "Rows 221-228: labour of a month with no units built, expensed when B135 = 0 (0 while B135 = 1), split to 5110 / 5130 / "
        "5140 with Prod Payroll's own payroll-tax and benefit shares of that month; rows 226-228 (Jan-Dec 27) feed PO rows 33-35. Rows 231-233: "
        "the Sep-26 builds go live Nov-26 (2026 vintage); their labour (0 when B132 = 0) is depreciated in 2027 like the 2027 vintages "
        "(full month from go-live + B13, 84 / 120 months, salvage B12) and feeds rows 193-194. Rows 236-241: memo.", S_TXT)
    put(220, "A", "text", "Month", S_HDR)
    months(220, lambda c: f"DeploySum!{c}4", S_DATE)
    put(220, "R", "text", "FY2027", S_HDR); put(220, "S", "text", "Source / formula", S_HDR)
    for rr, what, cap, units, wait in ((221, "Lift", 147, 150, 153), (222, "Tray", 148, 151, 155)):
        put(rr, "A", "text", f"{what} labour of a month with no {what.lower()}s built, expensed (B135 = 0)", S_TXT)
        for i, c in enumerate(MCOL):
            prev = "" if i == 0 else f"+{MCOL[i - 1]}{wait}"
            put(rr, c, "f", f"IF({c}{units}>0,0,(1-$B${SW_IDLE})*({c}{cap}{prev}))", S_MON)
        put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON)
    put(221, "S", "text", "Same pool as rows 152-155; with B135 = 0 nothing waits, so a month with no builds is expensed in full", S_TXT)
    put(223, "A", "text", "Idle-month labour expensed - total", S_BOLD)
    months(223, lambda c: f"{c}221+{c}222"); put(223, "R", "f", "SUM(F223:Q223)", S_MON)
    put(224, "A", "text", "Prod Payroll payroll-tax share of the month (rows 37 + 58 + 75) / row 81", S_TXT)
    months(224, lambda c: f"IF({PP}!{c}81=0,0,({PP}!{c}37+{PP}!{c}58+{PP}!{c}75)/{PP}!{c}81)", S_MON)
    put(225, "A", "text", "Prod Payroll benefits share of the month (rows 38 + 59 + 76) / row 81", S_TXT)
    months(225, lambda c: f"IF({PP}!{c}81=0,0,({PP}!{c}38+{PP}!{c}59+{PP}!{c}76)/{PP}!{c}81)", S_MON)
    for rr, gl, fn, porow in ((226, "5110 - Salaries and Wages - Production", lambda c: f"{c}223*(1-{c}224-{c}225)", 33),
                              (227, "5130 - Payroll Taxes - Production", lambda c: f"{c}223*{c}224", 34),
                              (228, "5140 - Benefits - Production", lambda c: f"{c}223*{c}225", 35)):
        put(rr, "A", "text", f"Idle labour expensed to {gl} (-> PO row {porow}, U:AF)", S_TXT)
        months(rr, fn, S_MON, Y27); put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON)
    put(226, "S", "text", "Jan-Dec 27 only (a 2026 month is outside the 2027 P&L); rows 226:228 sum to row 223 (F:Q)", S_TXT)
    put(229, "A", "text", "Check: rows 226:228 less row 223, Jan-Dec 27 (must be 0)", S_TXT)
    put(229, "R", "f", "SUM(R226:R228)-R223", S_MON)

    put(230, "A", "text", "Sep-26 builds (Nov-26 go-lives): 2027 depreciation of their labour (0 when B132 = 0; same rule as the 2027 vintages)", S_BOLD)
    for rr, pool, life in ((231, 152, "$B$9"), (232, 154, "$B$10")):
        nm = "SlipLift" if rr == 231 else "SlipTray"
        put(rr, "A", "text", f"{nm}s built Sep-Dec 26 that go live in 2026: labour depreciation", S_TXT)
        terms = "+".join(f"${c}${pool}*(YEAR(${c}$158)<2027)*(EOMONTH(${c}$158,$B$13)<={{c}}$27)" for c in Y26)
        months(rr, lambda c, terms=terms, life=life: "(" + terms.replace("{c}", c) + f")*(1-$B$12)/{life}", S_MON, Y27)
        put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON)
    put(231, "S", "text", "Pooled labour (rows 152 / 154, B:E) of builds whose go-live (row 158) is in 2026, from go-live + B13; lives B9 / B10; "
        "-> rows 193 / 194", S_TXT)
    put(233, "A", "text", "Total (rows 231:232)", S_BOLD)
    months(233, lambda c: f"{c}231+{c}232", S_MON, Y27); put(233, "R", "f", "SUM(F233:Q233)", S_MON)

    put(235, "A", "text", "Memo (v9): CapEx and the Dec-27 balance sheet of the capitalised labour", S_HDR)
    put(235, "B", "text", "Value", S_HDR)
    memo = [
        (236, "capex_2027", "CapEx increase in 2027 for production labour: 2027 payroll capitalised (R149 less the idle labour expensed, R223)", "R149-R223"),
        (237, "capex_2026", "Sep-Dec 26 labour capitalised (2026 CapEx; 0 when B132 = 0)", "SUM(B149:E149)-SUM(B221:E222)"),
        (238, "bs_in_service", "PP&E in service at Dec-27 (labour, net of 2027 depreciation) = B213", "B213"),
        (239, "bs_cip", "PP&E not yet placed in service (CIP) at Dec-27 (labour in builds for 2028 go-lives + labour waiting) = B204 + B205", "B204+B205"),
        (240, "idle_2027", "Idle-month labour expensed in 2027 (B135 = 0) = R223", "R223"),
        (241, "dep26_2027", "2027 depreciation of the Sep-26 builds' labour (B132 = 1) = R233", "R233"),
    ]
    for rr, key, lab, f in memo:
        put(rr, "A", "text", lab, S_TXT); put(rr, "B", "f", f, S_MON); R["m_" + key] = rr
    R.update(idle_lift=221, idle_tray=222, idle_total=223, tax_share=224, ben_share=225, idle_5110=226, idle_5130=227, idle_5140=228,
             idle_check=229, dep26_lift=231, dep26_tray=232, dep26_total=233)
    return rows, R


CELL = lambda coord: re.compile(rf'<c r="{coord}"(?: [^>]*?)?(?:/>|>.*?</c>)', re.S)


def replace_cell(x, coord, new_xml):
    m = list(CELL(coord).finditer(x))
    assert len(m) == 1, (coord, len(m))
    return x[:m[0].start()] + new_xml + x[m[0].end():], m[0].group(0)


UNESC = {"&apos;": "'", "&quot;": '"'}


def old_content(cxml):
    f = re.search(r"<f[^>]*>([^<]*)</f>", cxml)
    if f: return "=" + unescape(f.group(1), UNESC)
    t = re.search(r"<t[^>]*>([^<]*)</t>", cxml)
    if t: return unescape(t.group(1), UNESC)
    v = re.search(r"<v>([^<]*)</v>", cxml)
    return v.group(1) if v else ""


def style_of(cxml):
    return int(re.search(r' s="(\d+)"', cxml).group(1))


def stage1(v8, out, spec_out):
    z = zipfile.ZipFile(v8)
    p_ds, p_po = part_of(z, DS), part_of(z, "PO")
    spec = {"ds_edit": [], "ds_new": [], "po": [], "rowmap": None}
    x = z.read(p_ds).decode("utf-8")
    assert '<dimension ref="A1:S216"/>' in x and re.findall(r'<row r="(\d+)"', x)[-1] == "216"
    assert '<row r="135"' not in x
    # 1. B132 1 -> 0
    x, old = replace_cell(x, "B132", f'<c r="B132" s="{S_IN}" t="n"><v>0</v></c>')
    assert old_content(old) == "1"
    spec["ds_edit"].append({"cell": "B132", "kind": "num", "old": "1", "new": "0"})
    # labels
    for coord, txt in LABELS.items():
        m = CELL(coord).search(x); assert m, coord
        x, old = replace_cell(x, coord, cell_xml(coord, "text", txt, style_of(m.group(0))))
        spec["ds_edit"].append({"cell": coord, "kind": "text", "old": old_content(old), "new": txt})
    # wait rows 153 / 155
    for wait, cap, units in ((153, 147, 150), (155, 148, 151)):
        for i, c in enumerate(MCOL):
            coord = f"{c}{wait}"
            m = CELL(coord).search(x)
            nf = ds_wait(c, i, cap, units, wait)
            x, old = replace_cell(x, coord, cell_xml(coord, "f", nf, style_of(m.group(0))))
            spec["ds_edit"].append({"cell": coord, "kind": "f", "old": old_content(old), "new": "=" + nf})
    # output rows 193 / 194 + Sep-26 builds
    for rr, add in ((193, 231), (194, 232)):
        for c in Y27:
            coord = f"{c}{rr}"
            m = CELL(coord).search(x); of = old_content(m.group(0))
            assert of == f"=SUMIF($B$165:$B$189,$B{rr},{c}$165:{c}$189)", of
            nf = of[1:] + f"+{c}{add}"
            x, old = replace_cell(x, coord, cell_xml(coord, "f", nf, style_of(m.group(0))))
            spec["ds_edit"].append({"cell": coord, "kind": "f", "old": of, "new": "=" + nf})
    for coord, nf in FORMULAS.items():
        m = CELL(coord).search(x)
        x, old = replace_cell(x, coord, cell_xml(coord, "f", nf, style_of(m.group(0))))
        spec["ds_edit"].append({"cell": coord, "kind": "f", "old": old_content(old), "new": "=" + nf})
    # new rows: 135 and 218+
    rows, R = new_rows()

    def rowxml(r):
        cells = sorted(rows[r], key=lambda t: colkey(t[0]))
        assert len({c[0] for c in cells}) == len(cells)
        for c, k, v, s in cells:
            spec["ds_new"].append({"cell": f"{c}{r}", "kind": k, "content": ("=" + v) if k == "f" else v})
        return f'<row r="{r}">' + "".join(cell_xml(f"{c}{r}", k, v, s) for c, k, v, s in cells) + "</row>"
    i = x.index('<row r="136"')
    x = x[:i] + rowxml(SW_IDLE) + x[i:]
    last = max(rows)
    tail = "".join(rowxml(r) for r in sorted(rows) if r > 216)
    i = x.index("</sheetData>")
    x = x[:i] + tail + x[i:]
    x = x.replace('<dimension ref="A1:S216"/>', f'<dimension ref="A1:S{last}"/>', 1)
    for t in ("inside the PO run rate", "in the PO run rate", "CIP / inventory"):
        assert t not in x, t
    new_ds = x
    rm8 = build_v8.block()[1]
    rm8.update(R)
    spec["rowmap"] = rm8

    # PO rows 33-35: + idle labour expensed to that GL
    x = z.read(p_po).decode("utf-8")
    po_cols = [get_column_letter(c) for c in range(21, 33)]
    for row, src in ((33, 226), (34, 227), (35, 228)):
        for k, col in enumerate(po_cols):
            coord = f"{col}{row}"
            m = CELL(coord).search(x); of = old_content(m.group(0))
            assert re.fullmatch(r"=\(1-'Depreciation Schedule'!\$B\$131\)\*[-0-9.E]+", of), of
            nf = of[1:] + f"+'Depreciation Schedule'!{Y27[k]}{src}"
            vv = re.search(r"<v>([^<]*)</v>", m.group(0)).group(1)
            new = f'<c r="{coord}" s="{style_of(m.group(0))}" t="n"><f aca="false">{escape(nf, {chr(39): "&apos;"})}</f><v>{vv}</v></c>'
            x, _ = replace_cell(x, coord, new)
            spec["po"].append({"cell": coord, "old": of, "new": "=" + nf})
    for row, src in ((33, 226), (34, 227), (35, 228)):
        coord = f"AI{row}"
        m = CELL(coord).search(x); of = old_content(m.group(0))
        assert of.startswith("v8: 51"), of
        nt = of + (f" v9: + 'Depreciation Schedule' row {src} (idle-month labour expensed to this GL when B{SW_IDLE} = 0; 0 while B{SW_IDLE} = 1, "
                   f"booked).")
        x, _ = replace_cell(x, coord, cell_xml(coord, "text", nt, style_of(m.group(0))))
        spec["po"].append({"cell": coord, "old": of, "new": nt})
    new_po = x

    repl = {p_ds: new_ds, p_po: new_po}
    zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in z.infolist():
        data = z.read(info.filename)
        if info.filename in repl:
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = info.external_attr
            zo.writestr(zi, repl[info.filename].encode("utf-8"))
        else:
            zo.writestr(info, data)
    zo.close()
    spec["parts"] = {DS: p_ds, "PO": p_po}
    json.dump(spec, open(spec_out, "w"), indent=1)
    print(f"stage1 {out}: Depreciation Schedule cells edited {len(spec['ds_edit'])}, new {len(spec['ds_new'])} (row {SW_IDLE}, rows 218-{last}); "
          f"PO cells edited {len(spec['po'])}; parts changed {sorted(repl)}")


if __name__ == "__main__":
    if sys.argv[1] == "stage1":
        stage1(*sys.argv[2:5])
    elif sys.argv[1] == "final":
        build_v8.final(*sys.argv[2:6])
    else:
        raise SystemExit(__doc__)
