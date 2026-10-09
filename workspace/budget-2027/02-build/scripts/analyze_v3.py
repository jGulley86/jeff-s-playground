"""v3 checks. Reads only.

Usage:
  python3 -I analyze_v3.py V2_XLSX V3_XLSX NEW_DEPLOYSUM_XLSX EXT_JSON OUT_JSON

Checks: v2 -> v3 cell diff by sheet (formulas and cached values) with every changed PO / COO P&L cell classified;
COO P&L roll-up; revenue tie-out; independent recompute of the depreciation schedule (Python, from the raw
DeploySum values and PO 2026 run rates); error scan; package scan (externalLink parts, defined names); DeploySum
label map, tie-out rows 16-21 vs 54-64; unit-cost reconciliation; headline numbers.
Exit code 1 if a completion criterion fails.
"""
import sys, json, re, zipfile, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

V2, V3, NEW, EXT, OUT = sys.argv[1:6]
ERR = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#ERROR!"}
CORE = ["COO P&L", "FO", "PO", "CO", "FE", "PE"]
DEPTS = ["FO", "PO", "CO", "FE", "PE"]
ext = json.load(open(EXT))
lay = ext["layout"]


def norm(v):
    if isinstance(v, ArrayFormula):
        return ("array", v.ref, v.text)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


def iserr(v):
    return isinstance(v, str) and v in ERR


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


of = openpyxl.load_workbook(V2); ov = openpyxl.load_workbook(V2, data_only=True)
nf = openpyxl.load_workbook(V3); nv = openpyxl.load_workbook(V3, data_only=True)
sf = openpyxl.load_workbook(NEW); sv = openpyxl.load_workbook(NEW, data_only=True)
res = {"fail": []}

# ------------------------------------------------------------------ 1. diff v2 -> v3
sheets = [s for s in of.sheetnames if s in nf.sheetnames and s != "Collation Notes"]
diff = {}
changed = {}          # sheet -> list of (coord, kind)
comments = {"added": [], "removed": []}
for s in sheets:
    a_f, b_f, a_v, b_v = of[s], nf[s], ov[s], nv[s]
    mr = max(a_f.max_row, b_f.max_row); mc = max(a_f.max_column, b_f.max_column)
    n = df = dv = noise = 0; maxnoise = 0.0; lst = []
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            fa, fb = norm(a_f.cell(r, c).value), norm(b_f.cell(r, c).value)
            va, vb = norm(a_v.cell(r, c).value), norm(b_v.cell(r, c).value)
            ca, cb = a_f.cell(r, c).comment, b_f.cell(r, c).comment
            ta, tb = (ca.text if ca else None), (cb.text if cb else None)
            if ta != tb:
                k = f"{s}!{get_column_letter(c)}{r}"
                if ta is None: comments["added"].append(k)
                elif tb is None: comments["removed"].append(k)
                else: comments.setdefault("changed", []).append(k)
            if fa is None and fb is None and va is None and vb is None:
                continue
            n += 1
            kind = []
            if fa != fb: df += 1; kind.append("formula")
            if va != vb:
                if (isinstance(va, (int, float)) and isinstance(vb, (int, float)) and abs(va - vb) <= 1e-6):
                    noise += 1; maxnoise = max(maxnoise, abs(va - vb))   # recalculation-order float noise
                else:
                    dv += 1; kind.append("value")
            if kind:
                lst.append((r, c, "+".join(kind)))
    diff[s] = {"cells": n, "formula_diffs": df, "value_diffs": dv, "float_noise_cells": noise, "max_noise": maxnoise}
    changed[s] = lst
res["diff"] = diff
res["comments"] = comments
res["sheets_only_in_v3"] = [s for s in nf.sheetnames if s not in of.sheetnames]
res["sheets_only_in_v2"] = [s for s in of.sheetnames if s not in nf.sheetnames]

for s in ["FO", "CO", "FE", "PE", "Fld Ops Payroll", "Fld Maint Budget", "Fld Eng Budget"]:
    if diff[s]["formula_diffs"] or diff[s]["value_diffs"]:
        res["fail"].append(f"{s} changed vs v2")

# PO classification
PO_LINES = {12, 13, 15, 19, 20, 21, 22}
po_rows = {}
for r, c, k in changed["PO"]:
    po_rows.setdefault(r, set()).add(get_column_letter(c))
po_class = []
for r in sorted(po_rows):
    cols = sorted(po_rows[r], key=lambda x: (len(x), x))
    lab = nf["PO"].cell(r, 1).value
    if r in PO_LINES:
        cls = "edited line"
    elif set(cols) <= {"AH"}:
        cls = "ratio column only (AH % of total revenue = AG/AG16; revenue no longer 0)"
    else:
        cls = "subtotal / derived"
    po_class.append({"row": r, "label": lab, "cols": cols, "class": cls})
res["po_changed_rows"] = po_class
po_sub = [d for d in po_class if d["class"] == "subtotal / derived"]
# every subtotal/derived row must be a formula row whose U:AG formulas did not change
for d in po_sub:
    r = d["row"]
    for col in d["cols"]:
        a, b = of["PO"][f"{col}{r}"].value, nf["PO"][f"{col}{r}"].value
        if a != b:
            res["fail"].append(f"PO!{col}{r} formula changed on a non-edited row")
res["po_formula_changes_outside_lines"] = [f"PO!{get_column_letter(c)}{r}" for r, c, k in changed["PO"]
                                           if "formula" in k and r not in PO_LINES]
if res["po_formula_changes_outside_lines"]:
    res["fail"].append("PO formula changed outside edited lines")
coo_formula_changes = [f"{get_column_letter(c)}{r}" for r, c, k in changed["COO P&L"] if "formula" in k]
res["coo_formula_changes"] = coo_formula_changes
coo_rows = {}
for r, c, k in changed["COO P&L"]:
    coo_rows.setdefault(r, set()).add(get_column_letter(c))
res["coo_changed_rows"] = [{"row": r, "label": nf["COO P&L"].cell(r, 1).value,
                            "cols": sorted(v, key=lambda x: (len(x), x))} for r, v in sorted(coo_rows.items())]
if coo_formula_changes:
    res["fail"].append("COO P&L formulas changed")
CHAIN = {12, 13, 14, 15, 16, 19, 20, 21, 22, 37, 63, 64, 65, 117, 132}
res["coo_changes_outside_chain"] = [f"{x['row']}:{','.join(x['cols'])}" for x in res["coo_changed_rows"]
                                    if x["row"] not in CHAIN and set(x["cols"]) - {"AH"}]
if res["coo_changes_outside_chain"]:
    res["fail"].append("COO P&L values changed outside the revenue/depreciation chain")
pv_ = nv["Prod Payroll"]; pf_ = nf["Prod Payroll"]
nm_ = [c.coordinate for row in pv_.iter_rows() for c in row if c.value == "#NAME?"]
res["pp_name_xlookup"] = len([k for k in nm_ if "xlookup" in str(pf_[k].value).lower()])
res["pp_name_rows"] = sorted({pv_[k].row for k in nm_})

# ------------------------------------------------------------------ 2. roll-up
cp_f, cp_v = nf["COO P&L"], nv["COO P&L"]
cells = fails = 0; maxd = 0.0
for r in range(9, 142):
    for c in list(range(2, 15)) + list(range(21, 34)):
        f = cp_f.cell(r, c).value
        v = cp_v.cell(r, c).value
        if f is None and v is None:
            continue
        if isinstance(f, str) and "/" in f:
            continue
        if not isinstance(v, (int, float)) and v not in (None, ""):
            continue
        s = sum(num(nv[d].cell(r, c).value) for d in DEPTS)
        d = abs(num(v) - s)
        cells += 1; maxd = max(maxd, d)
        if d > 1: fails += 1
res["rollup"] = {"cells": cells, "fails": fails, "max_abs_diff": maxd}
if fails: res["fail"].append("COO P&L roll-up")
res["check_row_133_max"] = max(abs(num(cp_v.cell(133, c).value)) for c in list(range(2, 15)) + list(range(21, 34)))

# ------------------------------------------------------------------ 3. revenue
ds_new = sv["DeploySum"]
rev_src = [ds_new.cell(39, c).value for c in range(6, 18)]
rev_po = [nv["PO"].cell(12, c).value for c in range(21, 33)]
res["revenue"] = {"deploysum_F39_Q39_newfile": rev_src, "po_U12_AF12": rev_po,
                  "sum_src": sum(rev_src), "sum_po": sum(rev_po), "R39_newfile": ds_new["R39"].value,
                  "diff": sum(rev_po) - sum(rev_src),
                  "po_formulas": [nf["PO"].cell(12, c).value for c in range(21, 33)],
                  "sepdec26_deploysum": [ds_new.cell(39, c).value for c in range(2, 6)],
                  "sepdec26_po4100": [ov["PO"].cell(12, c).value for c in range(10, 14)],
                  "sepdec26_po_total_income": [ov["PO"].cell(16, c).value for c in range(10, 14)]}
if abs(res["revenue"]["diff"]) > 1: res["fail"].append("revenue tie")

# ------------------------------------------------------------------ 4. independent depreciation recompute
lifts_go = [ds_new.cell(19, c).value for c in range(6, 18)]
trays_go = [ds_new.cell(20, c).value for c in range(6, 18)]
R24, R25, R29 = ds_new["R24"].value, ds_new["R25"].value, ds_new["R29"].value
LIFT, TRAY = 65000.0, 8000.0                 # DeploySum note 3 (new file A46)
note3 = ds_new["A46"].value
assert "SlipLift $65,000" in note3 and "SlipTray $8,000" in note3, note3
PERIPH = (R29 - R25 * TRAY) / R24 - LIFT
LL, TL, PL = 84, 120, 84
po2026 = {r: [ov["PO"].cell(r, c).value for c in range(2, 14)] for r in (19, 20, 21, 22)}
runrate = {r: po2026[r][11] for r in po2026}
calc = {19: [], 20: [], 21: [], 22: []}
for m in range(1, 13):
    lifts_in = sum(lifts_go[v - 1] for v in range(1, 13) if m >= v)
    trays_in = sum(trays_go[v - 1] for v in range(1, 13) if m >= v)
    calc[19].append(runrate[19])
    calc[20].append(runrate[20] + lifts_in * LIFT / LL)
    calc[21].append(runrate[21] + trays_in * TRAY / TL)
    calc[22].append(runrate[22] + lifts_in * PERIPH / PL)
wbk = {r: [nv["PO"].cell(r, c).value for c in range(21, 33)] for r in (19, 20, 21, 22)}
maxdep = max(abs(calc[r][k] - wbk[r][k]) for r in calc for k in range(12))
res["dep"] = {"periph": PERIPH, "note3": note3, "lifts_go": lifts_go, "trays_go": trays_go,
              "runrate": runrate, "po2026": po2026, "calc": calc, "workbook": wbk, "max_abs_diff": maxdep,
              "fy_calc": {r: sum(calc[r]) for r in calc}, "fy_wb": {r: sum(wbk[r]) for r in wbk},
              "po_2026_N": {r: ov["PO"].cell(r, 14).value for r in calc},
              "unit_months_lift": sum(lifts_go[v - 1] * (13 - v) for v in range(1, 13)),
              "unit_months_tray": sum(trays_go[v - 1] * (13 - v) for v in range(1, 13)),
              "dec_units_lift": sum(lifts_go), "dec_units_tray": sum(trays_go),
              "new_fy": {"lift": sum(lifts_go[v - 1] * (13 - v) for v in range(1, 13)) * LIFT / LL,
                         "tray": sum(trays_go[v - 1] * (13 - v) for v in range(1, 13)) * TRAY / TL,
                         "periph": sum(lifts_go[v - 1] * (13 - v) for v in range(1, 13)) * PERIPH / PL},
              "dec27_monthly": {r: calc[r][11] for r in calc}}
if maxdep > 0.01: res["fail"].append("depreciation recompute")
# sampled month: Jun-27 (k = 5)
k = 5
lin = sum(lifts_go[:k + 1]); tin = sum(trays_go[:k + 1])
res["dep"]["sample"] = {"month": "Jun-27", "lifts_in_service": lin, "trays_in_service": tin,
                        "lines": {
                            19: [runrate[19], 0, runrate[19], wbk[19][k]],
                            20: [runrate[20], lin * LIFT / LL, runrate[20] + lin * LIFT / LL, wbk[20][k]],
                            21: [runrate[21], tin * TRAY / TL, runrate[21] + tin * TRAY / TL, wbk[21][k]],
                            22: [runrate[22], lin * PERIPH / PL, runrate[22] + lin * PERIPH / PL, wbk[22][k]]}}
# PO forecast step-ups
steps = {}
for r, unit in ((20, LIFT / LL), (21, TRAY / TL), (22, PERIPH / PL)):
    s = [po2026[r][k2] - po2026[r][k2 - 1] for k2 in (9, 10, 11)]
    steps[r] = {"steps": s, "units": [x / unit for x in s], "dec_minus_sep": po2026[r][11] - po2026[r][8]}
res["dep"]["steps"] = steps
res["dep"]["q4_deploysum"] = {"lifts": [ds_new.cell(19, c).value for c in (3, 4, 5)],
                              "trays": [ds_new.cell(20, c).value for c in (3, 4, 5)],
                              "sep_lifts": ds_new["B19"].value, "sep_trays": ds_new["B20"].value}
gap_l = sum(res["dep"]["q4_deploysum"]["lifts"]) - steps[20]["dec_minus_sep"] / (LIFT / LL)
gap_t = sum(res["dep"]["q4_deploysum"]["trays"]) - steps[21]["dec_minus_sep"] / (TRAY / TL)
res["dep"]["q4_gap"] = {"lifts": gap_l, "trays": gap_t, "monthly": gap_l * LIFT / LL + gap_t * TRAY / TL,
                        "annual": 12 * (gap_l * LIFT / LL + gap_t * TRAY / TL),
                        "periph_if_missing_annual": 12 * sum(res['dep']['q4_deploysum']['lifts']) * PERIPH / PL}
# schedule sheet checks
sch = nv["Depreciation Schedule"]
res["sched_check_row_max"] = max(abs(num(sch.cell(lay["check"], c).value)) for c in range(6, 19))
res["sched_bridge_diff"] = sch.cell(lay["r7"] + 5, 2).value
res["sched_recon_max_resid"] = sch.cell(lay["r6"][3], 7).value
if res["sched_check_row_max"] > 0.01 or abs(num(res["sched_bridge_diff"])) > 0.01:
    res["fail"].append("schedule internal checks")

# ------------------------------------------------------------------ 5. unit-cost reconciliation (independent)
recon = []
for c in range(2, 18):
    d = ds_new.cell(4, c).value
    l, t, cap = ds_new.cell(24, c).value, ds_new.cell(25, c).value, ds_new.cell(29, c).value
    recon.append({"month": d.strftime("%b-%y"), "lifts": l, "trays": t, "capex": cap,
                  "single_price": (cap / (l + t)) if (l + t) else None,
                  "calc": l * (LIFT + PERIPH) + t * TRAY, "resid": cap - (l * (LIFT + PERIPH) + t * TRAY),
                  "resid_rounded_note": cap - (l * (LIFT + 4138) + t * TRAY)})
res["recon"] = recon
res["recon_max_resid"] = max(abs(x["resid"]) for x in recon)
res["capex"] = {"R29": R29, "R35": ds_new["R35"].value, "D29": ds_new["D29"].value, "E29": ds_new["E29"].value,
                "dep_base": sum(lifts_go) * (LIFT + PERIPH) + sum(trays_go) * TRAY,
                "R24": R24, "R25": R25, "R33": ds_new["R33"].value, "R34": ds_new["R34"].value}

# ------------------------------------------------------------------ 6. errors
err = {}
new_errors = []; resolved = {}
for s in nf.sheetnames:
    if s == "Collation Notes":
        continue
    cnt3 = {}
    for row in nv[s].iter_rows():
        for c in row:
            if iserr(c.value):
                cnt3[c.value] = cnt3.get(c.value, 0) + 1
                if s in ov.sheetnames:
                    o = ov[s][c.coordinate].value
                    if not iserr(o):
                        new_errors.append(f"{s}!{c.coordinate}")
                else:
                    new_errors.append(f"{s}!{c.coordinate}")
    cnt2 = {}
    if s in ov.sheetnames:
        for row in ov[s].iter_rows():
            for c in row:
                if iserr(c.value):
                    cnt2[c.value] = cnt2.get(c.value, 0) + 1
                    if not iserr(nv[s][c.coordinate].value):
                        resolved[s] = resolved.get(s, 0) + 1
    err[s] = {"v2": cnt2, "v3": cnt3}
res["errors"] = err
res["new_errors"] = new_errors
res["resolved_errors"] = resolved
if new_errors: res["fail"].append("new errors")
# Prod Payroll: did anything that other sheets read change?
pp_reads = []
for s in nf.sheetnames:
    if s in ("Prod Payroll", "Collation Notes"): continue
    for row in nf[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and "Prod Payroll" in c.value:
                pp_reads.append(f"{s}!{c.coordinate}")
res["prod_payroll_readers"] = pp_reads
res["prod_payroll_value_changes"] = diff["Prod Payroll"]["value_diffs"]
pp_tot = {}
for r in (41, 62, 79, 81):
    pp_tot[r] = [nv["Prod Payroll"].cell(r, 1).value, ov["Prod Payroll"].cell(r, 19).value, nv["Prod Payroll"].cell(r, 19).value]
res["prod_payroll_totals"] = pp_tot

# ------------------------------------------------------------------ 7. package
z = zipfile.ZipFile(V3)
names = z.namelist()
wbxml = z.read("xl/workbook.xml").decode("utf8")
res["package"] = {"parts": len(names), "externalLink_parts": sum(1 for n in names if n.startswith("xl/externalLinks/")),
                  "externalReference_elements": wbxml.count("<externalReference "),
                  "defined_names": len(re.findall(r"<definedName[ >]", wbxml))}
if res["package"]["externalLink_parts"]: res["fail"].append("externalLink parts")

# ------------------------------------------------------------------ 8. DeploySum label map + tie-out
ds3 = nf["DeploySum"]; ds3v = nv["DeploySum"]
labmap = []
for r in range(1, 50):
    a = sf["DeploySum"].cell(r, 1).value
    a_v = sv["DeploySum"].cell(r, 1).value
    t = r if r <= 38 else r + 1
    b = ds3v.cell(t, 1).value
    ok = (a_v or "").replace("Dollars in $000s.", "Dollars in $ (relabelled in v3; the source said $000s).").replace("($000s)", "($)") == (b or "")
    labmap.append([r, t, a_v, b, ok])
res["labmap_bad"] = [x for x in labmap if not x[4]]
if res["labmap_bad"]: res["fail"].append("DeploySum label map")
# values: every mapped cell B:R equals the new file's cached value
vbad = []
for r in range(3, 50):
    for c in range(2, 19):
        a = sv["DeploySum"].cell(r, c).value
        b = ds3v.cell(r if r <= 38 else r + 1, c).value
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            if abs(a - b) > 1e-6: vbad.append(f"{get_column_letter(c)}{r}")
        elif norm(a) != norm(b):
            vbad.append(f"{get_column_letter(c)}{r}")
res["deploysum_value_mismatch"] = vbad
if vbad: res["fail"].append("DeploySum values vs new file")
tie = []
for r in range(70, 76):
    tie.append([ds3v.cell(r, 1).value] + [ds3v.cell(r, c).value for c in range(2, 19)])
res["tieout"] = tie
res["tieout_status"] = ds3v["A76"].value
tie_rows = {}
for lab, a, b in (("16+17 vs 54+57+60", (16, 17), (54, 57, 60)), ("18 vs 64", (18,), (64,)),
                  ("19 vs 55+58+61", (19,), (55, 58, 61)), ("20 vs 56+59+62", (20,), (56, 59, 62)),
                  ("21 vs 63", (21,), (63,))):
    tie_rows[lab] = {"FY_top": sum(num(ds3v.cell(x, 18).value) for x in a),
                     "FY_split": sum(num(ds3v.cell(x, 18).value) for x in b),
                     "SepDec26_top": sum(num(ds3v.cell(x, c).value) for x in a for c in range(2, 6)),
                     "SepDec26_split": sum(num(ds3v.cell(x, c).value) for x in b for c in range(2, 6)),
                     "max_abs_month": max(abs(sum(num(ds3v.cell(x, c).value) for x in a) -
                                              sum(num(ds3v.cell(x, c).value) for x in b)) for c in range(2, 19))}
res["tie_rows"] = tie_rows
if ds3v["A76"].value != "Tie-out rows 16-21 vs 54-64: OK": res["fail"].append("tie-out")
res["hidden_rows_v2"] = sorted(r for r, d in of["DeploySum"].row_dimensions.items() if d.hidden)
res["hidden_rows_v3"] = sorted(r for r, d in ds3.row_dimensions.items() if d.hidden)
res["ext_count"] = len(ext["ext"])
rows_ext = {}
for v3c, nc, f, v in ext["ext"]:
    rr = int(re.sub(r"[A-Z]+", "", v3c))
    rows_ext.setdefault(rr, []).append([v3c, nc, f, v])
res["ext_rows"] = {k: {"n": len(v), "cells": [x[0] for x in v], "new_cells": [x[1] for x in v],
                       "example": v[0][2], "label": ds3v.cell(k, 1).value if k > 4 else ("(row 3 tie-out text)" if k == 3 else "(month-end dates)")}
                   for k, v in sorted(rows_ext.items())}
res["relabel"] = ext["relabel"]
res["v2_frozen_comments_removed"] = [c for c in comments["removed"] if c.startswith("DeploySum!")]
res["deploysum_v2_errors"] = err["DeploySum"]["v2"]

# ------------------------------------------------------------------ 9. headline
def g(sh, a, wbv=nv): return num(wbv[sh][a].value)
H = {}
for lab, row in (("Revenue", 16), ("COGS", 63), ("Gross Profit", 64), ("Opex", 116), ("Net Ordinary Income", 117),
                 ("Net Other Income", 131), ("Net Income", 132)):
    H[lab] = {"row": row, "N": g("COO P&L", f"N{row}"), "AG": g("COO P&L", f"AG{row}"),
              "AG_v2": g("COO P&L", f"AG{row}", ov)}
dep27 = sum(g("COO P&L", f"AG{r}") for r in (19, 20, 21, 22))
dep26 = sum(g("COO P&L", f"N{r}") for r in (19, 20, 21, 22))
H["Depreciation (rows 19-22)"] = {"row": "19-22", "N": dep26, "AG": dep27, "AG_v2": 0.0}
H["Total cost (COGS+Opex)"] = {"row": "63+116", "N": H["COGS"]["N"] + H["Opex"]["N"],
                               "AG": H["COGS"]["AG"] + H["Opex"]["AG"], "AG_v2": H["COGS"]["AG_v2"] + H["Opex"]["AG_v2"]}
H["Cost excl. depreciation"] = {"row": "63+116-(19:22)", "N": H["Total cost (COGS+Opex)"]["N"] - dep26,
                                "AG": H["Total cost (COGS+Opex)"]["AG"] - dep27,
                                "AG_v2": H["Total cost (COGS+Opex)"]["AG_v2"]}
res["headline"] = H
res["po_2026_rev"] = {r: g("PO", f"N{r}") for r in (12, 13, 15)}

# ------------------------------------------------------------------ 10. variance scan (v3 values), as in v2
var = []
for s in ["FO", "PO", "CO", "FE", "PE", "COO P&L"]:
    for r in range(9, 142):
        lab = nv[s].cell(r, 1).value
        n_, ag = nv[s].cell(r, 14).value, nv[s].cell(r, 33).value
        if not isinstance(n_, (int, float)) or not isinstance(ag, (int, float)) or r in (65, 141):
            continue
        ch = ag - n_
        if abs(ch) > 50000 and (n_ == 0 or abs(ch) / abs(n_) > 0.5):
            var.append([s, r, lab, n_, ag, ch, (ch / abs(n_)) if n_ else None])
res["variance"] = var
res["coo_row65"] = {"AG": nv["COO P&L"]["AG65"].value, "N": nv["COO P&L"]["N65"].value if nv["COO P&L"]["N65"].value is not None else None}
res["sheets_v3"] = [[s, nf[s].sheet_state] for s in nf.sheetnames]
res["pp_S41"] = nv["Prod Payroll"]["S41"].value
res["po_5100_AG36"] = nv["PO"]["AG36"].value

json.dump(res, open(OUT, "w"), indent=1, default=str)
print("analyze_v3:", "FAIL " + "; ".join(res["fail"]) if res["fail"] else "all checks pass",
      "| rollup", res["rollup"], "| rev diff", round(res["revenue"]["diff"], 6), "| dep max diff", maxdep,
      "| new errors", len(new_errors), "| extLinks", res["package"]["externalLink_parts"],
      "| names", res["package"]["defined_names"])
sys.exit(1 if res["fail"] else 0)
