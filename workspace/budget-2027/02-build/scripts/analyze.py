"""Reconcile and scan the collated workbook against both inputs.

Usage:
  python3 -I analyze.py COO_BOOK SD_BOOK OUT_XLSX_RECALCULATED EXTREFS_JSON RESULT_JSON

Reads only. Writes RESULT_JSON with reconciliation, error scans and raw findings.
"""
import sys, re, json, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter, column_index_from_string

COO_PATH, SD_PATH, OUT_PATH, EXT_JSON, RES_JSON = sys.argv[1:6]

DEPTS = ["FO", "PO", "CO", "FE", "PE"]
SD_DEPTS = ["FO", "FE"]
KEY_ROWS = {16: "Total Income", 63: "Total COGS", 64: "Gross Profit", 116: "Total Expense", 132: "Net Income"}
MONTHS26 = [get_column_letter(i) for i in range(2, 14)]      # B..M
MONTHS27 = [get_column_letter(i) for i in range(21, 33)]     # U..AF
PL_ROWS = range(9, 142)
ERR_RE = re.compile(r"^(#(NULL!|DIV/0!|VALUE!|REF!|NAME\?|NUM!|N/A|ERROR!|CALC!|SPILL!)|Err:\d+)$")
PLACEHOLDER_RE = re.compile(r"PLACEHOLDER|TBD|AI suggested|assumed", re.I)


def is_err(v):
    return isinstance(v, str) and bool(ERR_RE.match(v.strip()))


def ftext(v):
    if isinstance(v, ArrayFormula):
        return v.text
    if isinstance(v, str) and v.startswith("="):
        return v
    return None


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def load(path):
    return openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True)


coo_f, coo_v = load(COO_PATH)
sd_f, sd_v = load(SD_PATH)
out_f, out_v = load(OUT_PATH)
ext = json.load(open(EXT_JSON))
ext_cells = {(e["sheet"], e["cell"]) for e in ext["external_cells"]}
R = {}

# ---------------------------------------------------------------- reconciliation
def val(wb, sh, cell):
    return wb[sh][cell].value

recon = []
for d in DEPTS:
    src = sd_v if d in SD_DEPTS else coo_v
    for r, lab in KEY_ROWS.items():
        for col in ("N", "AG"):
            s = num(val(src, d, f"{col}{r}")) or 0.0
            o = num(val(out_v, d, f"{col}{r}")) or 0.0
            recon.append({"dept": d, "src": "SD" if d in SD_DEPTS else "COO", "row": r, "line": lab,
                          "col": col, "source": s, "output": o, "delta": o - s})
R["recon"] = recon
R["replaced_coo_fofe"] = [{"dept": d, "row": r, "line": lab, "col": col,
                           "coo_book": num(val(coo_v, d, f"{col}{r}")) or 0.0,
                           "sd_book": num(val(sd_v, d, f"{col}{r}")) or 0.0}
                          for d in SD_DEPTS for r, lab in KEY_ROWS.items() for col in ("N", "AG")]

# COO P&L before vs after
pl = []
for r, lab in KEY_ROWS.items():
    for col in ("N", "AG"):
        b = num(val(coo_v, "COO P&L", f"{col}{r}")) or 0.0
        a = num(val(out_v, "COO P&L", f"{col}{r}")) or 0.0
        pl.append({"row": r, "line": lab, "col": col, "before": b, "after": a, "change": a - b})
R["coo_pl_before_after"] = pl

# COO P&L line = sum of depts (all numeric rows, B..N and U..AG)
rollup_fail, rollup_checked, worst = [], 0, 0.0
ADDITIVE_ROWS = [r for r in range(9, 133) if r != 65] + list(range(137, 141))   # 65 = GM% ratio on COO P&L, 133 = check row, 141 = ratio row
for r in ADDITIVE_ROWS:
    for col in MONTHS26 + ["N"] + MONTHS27 + ["AG"]:
        a = val(out_v, "COO P&L", f"{col}{r}")
        parts = [val(out_v, d, f"{col}{r}") for d in DEPTS]
        if num(a) is None and all(num(p) is None for p in parts):
            continue
        rollup_checked += 1
        sp = sum(num(p) or 0.0 for p in parts)
        diff = (num(a) or 0.0) - sp
        worst = max(worst, abs(diff))
        if abs(diff) > 1:
            rollup_fail.append({"cell": f"{col}{r}", "label": val(out_v, "COO P&L", f"A{r}"),
                                "coo_pl": num(a), "sum_depts": sp, "diff": diff})
R["rollup"] = {"checked": rollup_checked, "fail": rollup_fail, "max_abs_diff": worst}

# ---------------------------------------------------------------- cell-by-cell fidelity
def compare_sheet(src_v, src_name, out_name):
    s, o = src_v[src_name], out_v[out_name]
    maxr = max(s.max_row, o.max_row); maxc = max(s.max_column, o.max_column)
    res = {"cells": 0, "num_max_abs_diff": 0.0, "mismatch": [], "err_code_changed": [],
           "new_errors": [], "errors_cleared": []}
    for row in range(1, maxr + 1):
        for colno in range(1, maxc + 1):
            a = s.cell(row, colno).value; b = o.cell(row, colno).value
            if a in (None, "") and b in (None, ""):
                continue
            coord = f"{get_column_letter(colno)}{row}"
            res["cells"] += 1
            if is_err(a) or is_err(b):
                if is_err(a) and is_err(b):
                    if a != b:
                        res["err_code_changed"].append([coord, a, b])
                elif is_err(b):
                    res["new_errors"].append([coord, a, b])
                else:
                    res["errors_cleared"].append([coord, a, b])
                continue
            if isinstance(a, datetime.datetime) or isinstance(b, datetime.datetime):
                if a != b:
                    res["mismatch"].append([coord, str(a), str(b)])
                continue
            na, nb = num(a), num(b)
            if na is not None or nb is not None:
                na = na or 0.0; nb = nb or 0.0
                dd = abs(na - nb)
                res["num_max_abs_diff"] = max(res["num_max_abs_diff"], dd)
                if dd > max(0.01, 1e-9 * abs(na)):
                    res["mismatch"].append([coord, a, b])
                continue
            if a != b:
                res["mismatch"].append([coord, a, b])
    return res

fid = {}
for sh in ext["carried"]:
    fid[sh] = compare_sheet(sd_v, sh, sh)
for sh in ["COO P&L", "PO", "CO", "PE"]:
    if sh != "COO P&L":
        fid[sh] = compare_sheet(coo_v, sh, sh)
R["fidelity"] = {k: {**v, "mismatch_n": len(v["mismatch"]), "mismatch": v["mismatch"][:40],
                     "err_code_changed_n": len(v["err_code_changed"]),
                     "err_code_changed": v["err_code_changed"][:10],
                     "new_errors_n": len(v["new_errors"]), "errors_cleared_n": len(v["errors_cleared"]),
                     "errors_cleared": v["errors_cleared"][:20]} for k, v in fid.items()}

# formulas preserved?  count formulas source vs output for carried sheets
fcount = {}
for sh in ext["carried"]:
    sc = sum(1 for row in sd_f[sh].iter_rows() for c in row if ftext(c.value))
    oc = sum(1 for row in out_f[sh].iter_rows() for c in row if ftext(c.value))
    fcount[sh] = {"source_formulas": sc, "output_formulas": oc,
                  "external_replaced": sum(1 for e in ext["external_cells"] if e["sheet"] == sh)}
for sh in ["COO P&L", "PO", "CO", "PE"]:
    sc = sum(1 for row in coo_f[sh].iter_rows() for c in row if ftext(c.value))
    oc = sum(1 for row in out_f[sh].iter_rows() for c in row if ftext(c.value))
    fcount[sh] = {"source_formulas": sc, "output_formulas": oc, "external_replaced": 0}
R["formula_counts"] = fcount

# ---------------------------------------------------------------- error scan before/after
def err_counts(wbv, names, typed=False):
    # source books (Google exports) store cached errors as text; the LibreOffice output types them as errors.
    # typed=True counts only real error cells, so error codes quoted as text on the notes sheet are not counted.
    out = {}
    for sh in names:
        cnt = {}
        for row in wbv[sh].iter_rows():
            for c in row:
                if is_err(c.value) and (not typed or c.data_type == "e"):
                    cnt[c.value] = cnt.get(c.value, 0) + 1
        out[sh] = cnt
    return out
R["errors_before_coo"] = err_counts(coo_v, coo_v.sheetnames)
R["errors_before_sd"] = err_counts(sd_v, sd_v.sheetnames)
R["errors_after"] = err_counts(out_v, out_v.sheetnames, typed=True)

# ---------------------------------------------------------------- issue scans
F = {}

# (a) 2026 B:N differences FO/FE between books (cached values) + formula/constant differences
d26 = {}
for d in SD_DEPTS:
    rows = []
    for r in PL_ROWS:
        cells = []
        for col in MONTHS26 + ["N"]:
            a = num(val(coo_v, d, f"{col}{r}")) or 0.0
            b = num(val(sd_v, d, f"{col}{r}")) or 0.0
            if abs(a - b) > 0.005:
                cells.append([col, a, b])
        if cells:
            rows.append({"row": r, "label": val(sd_v, d, f"A{r}"), "cells": cells,
                         "N_coo": num(val(coo_v, d, f"N{r}")) or 0.0, "N_sd": num(val(sd_v, d, f"N{r}")) or 0.0})
    d26[d] = rows
F["diff_2026"] = d26

# formula text differences FO/FE (all columns), counted
fdiff = {}
for d in SD_DEPTS:
    n = 0; ex = []
    for r in range(1, 201):
        for ci in range(1, 37):
            a = coo_f[d].cell(r, ci).value; b = sd_f[d].cell(r, ci).value
            fa, fb = ftext(a), ftext(b)
            if (fa or fb) and fa != fb:
                n += 1
                if len(ex) < 400:
                    ex.append([f"{get_column_letter(ci)}{r}", str(a)[:120], str(b)[:160]])
    fdiff[d] = {"n": n, "examples": ex}
F["formula_diff_FO_FE"] = fdiff

# (b) revenue by dept
rev = {}
for d in DEPTS + ["COO P&L"]:
    rev[d] = {"N16": num(val(out_v, d, "N16")) or 0.0, "AG16": num(val(out_v, d, "AG16")) or 0.0,
              "R_method_row12": val(out_f, d, "R12"), "S12": val(out_f, d, "S12"), "T12": val(out_f, d, "T12")}
F["revenue"] = rev
# revenue lines with 2026 values: show method columns
rev_lines = []
for d in DEPTS:
    for r in (11, 12, 13, 15):
        n = num(val(out_v, d, f"N{r}")) or 0.0
        ag = num(val(out_v, d, f"AG{r}")) or 0.0
        if n or ag:
            rev_lines.append({"dept": d, "row": r, "label": val(out_v, d, f"A{r}"), "N": n, "AG": ag,
                              "R": val(out_f, d, f"R{r}"), "S": val(out_f, d, f"S{r}"), "T": val(out_f, d, f"T{r}"),
                              "U_formula": str(val(out_f, d, f"U{r}"))[:120]})
F["revenue_lines"] = rev_lines

# (d) hardcoded numbers inside formula rows
NUM_LIT = re.compile(r"(?<![A-Za-z0-9_$.:])(\d+(?:\.\d+)?)(?![\d.]*[A-Za-z(])")
def literals(f):
    t = re.sub(r'"[^"]*"', '""', f)
    t = re.sub(r"'(?:[^']|'')+'!", "", t)
    t = re.sub(r"\b[A-Za-z_][A-Za-z0-9_.]*!", "", t)
    t = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "R", t)
    return [float(x) for x in NUM_LIT.findall(t)]

hard = []
def is_total_label(lab):
    lab = (lab or "").strip()
    return lab.startswith("Total") or lab in ("Gross Profit", "Net Ordinary Income", "Net Other Income", "Net Income")

for d in DEPTS + ["COO P&L"]:
    ws = out_f[d]
    for r in PL_ROWS:
        lab = ws[f"A{r}"].value
        b_cols = MONTHS27
        kinds = {c: ("f" if ftext(ws[f"{c}{r}"].value) else ("n" if num(ws[f"{c}{r}"].value) is not None else "e")) for c in b_cols}
        consts = [c for c, k in kinds.items() if k == "n"]
        forms = [c for c, k in kinds.items() if k == "f"]
        if consts and forms:
            hard.append({"sheet": d, "row": r, "label": lab, "type": "mixed formula/constant in 2027 months",
                         "cells": [f"{c}{r}={ws[f'{c}{r}'].value}" for c in consts]})
        if is_total_label(lab) or (d == "COO P&L"):
            for c in MONTHS26 + ["N"] + MONTHS27 + ["AG"]:
                v = ws[f"{c}{r}"].value
                if num(v) is not None and num(v) != 0 and not ftext(v):
                    hard.append({"sheet": d, "row": r, "label": lab, "type": "constant in total/roll-up cell",
                                 "cells": [f"{c}{r}={v}"]})
        for c in ["N", "AG"] + MONTHS27:
            f = ftext(ws[f"{c}{r}"].value)
            if f:
                lits = [x for x in literals(f) if abs(x) >= 100]
                if lits:
                    hard.append({"sheet": d, "row": r, "label": lab, "type": "numeric literal inside formula",
                                 "cells": [f"{c}{r}: {f[:160]}"]})
# supporting sheets: literals >= 1000 inside formulas
for sh in ext["support"]:
    ws = out_f[sh]
    for row in ws.iter_rows():
        for c in row:
            f = ftext(c.value)
            if f and not (sh, c.coordinate) in ext_cells:
                lits = [x for x in literals(f) if abs(x) >= 1000 and x not in (1900, 2026, 2027, 2028, 2080)]
                if "SUMIF" in f.upper():   # GL account codes used as SUMIF criteria are not plugs
                    lits = [x for x in lits if not (x.is_integer() and 4000 <= x <= 9999)]
                if lits:
                    hard.append({"sheet": sh, "row": c.row, "label": ws.cell(c.row, 1).value,
                                 "type": "numeric literal inside formula", "cells": [f"{c.coordinate}: {f[:160]}"]})
# supporting sheets: rows mixing formulas and typed numbers across the monthly grid (possible overrides)
for sh in ext["support"]:
    ws = out_f[sh]
    for r in range(1, ws.max_row + 1):
        if (sh, f"B{r}") in ext_cells:
            continue
        kinds = {}
        for ci in range(2, 19):
            v = ws.cell(r, ci).value
            if v is None:
                continue
            kinds[get_column_letter(ci)] = "f" if ftext(v) else ("n" if num(v) is not None else "t")
        fs = [c for c, k in kinds.items() if k == "f"]
        # ignore GL-code (4000-9999) and row-pointer (1-200) integer keys used as lookup labels
        ns = [c for c, k in kinds.items() if k == "n"
              and not (float(ws[f"{c}{r}"].value).is_integer() and (1 <= ws[f"{c}{r}"].value <= 200 or 4000 <= ws[f"{c}{r}"].value <= 9999))]
        if len(fs) >= 3 and ns:
            hard.append({"sheet": sh, "row": r, "label": ws.cell(r, 1).value,
                         "type": "typed number in a formula row (monthly grid B:R)",
                         "cells": [f"{c}{r}={ws[f'{c}{r}'].value}" for c in ns]})
F["hardcodes"] = hard

# (e) placeholders, all sheets both books, values + comments
ph = []
for book, wbf in (("COO", coo_f), ("SD", sd_f)):
    for ws in wbf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                txt = v if isinstance(v, str) else ftext(v)
                if txt and PLACEHOLDER_RE.search(txt):
                    ph.append({"book": book, "sheet": ws.title, "cell": c.coordinate,
                               "match": PLACEHOLDER_RE.search(txt).group(0), "text": txt[:200]})
                if c.comment and PLACEHOLDER_RE.search(c.comment.text or ""):
                    ph.append({"book": book, "sheet": ws.title, "cell": c.coordinate + " (comment)",
                               "match": PLACEHOLDER_RE.search(c.comment.text).group(0), "text": c.comment.text[:200]})
F["placeholders"] = ph

# (f) wild variances 2027 vs 2026 (output values)
var = []
for d in DEPTS + ["COO P&L"]:
    for r in PL_ROWS:
        n = num(val(out_v, d, f"N{r}")); ag = num(val(out_v, d, f"AG{r}"))
        if n is None and ag is None:
            continue
        n = n or 0.0; ag = ag or 0.0
        diff = ag - n
        if abs(diff) > 50000 and (n == 0 or abs(diff) / abs(n) > 0.5):
            var.append({"sheet": d, "row": r, "label": val(out_v, d, f"A{r}"), "N": n, "AG": ag,
                        "diff": diff, "pct": (diff / abs(n)) if n else None})
F["variance"] = var

# (g1) COO P&L dept formulas omit a dept / wrong cell
DEPT_REF = re.compile(r"(?:'([^']+)'|([A-Za-z]+))!\$?([A-Z]{1,3})\$?(\d+)")
omit = []; pl_types = {}
chk = [num(out_v["COO P&L"][f"{c}133"].value) or 0.0 for c in MONTHS26 + MONTHS27]
F["check_row_133_max_abs"] = max(abs(x) for x in chk)
for r in PL_ROWS:
    ws = out_f["COO P&L"]
    if (ws[f"A{r}"].value or "").strip().lower() == "check":
        pl_types[r] = ["check"]; continue
    types = set()
    for c in MONTHS26 + MONTHS27:
        f = ftext(ws[f"{c}{r}"].value)
        if not f:
            v = ws[f"{c}{r}"].value
            types.add("const" if num(v) is not None else "blank")
            continue
        refs = [(m.group(1) or m.group(2), m.group(3), int(m.group(4))) for m in DEPT_REF.finditer(f)]
        if refs:
            types.add("dept")
            sheets = {x[0] for x in refs}
            bad = [x for x in refs if x[1] != c or x[2] != r]
            missing = [d for d in DEPTS if d not in sheets]
            if missing or bad:
                omit.append({"cell": f"{c}{r}", "label": ws[f"A{r}"].value, "formula": f,
                             "missing": missing, "misaligned": [f"{x[0]}!{x[1]}{x[2]}" for x in bad]})
        else:
            types.add("internal")
    pl_types[r] = sorted(types)
F["coo_pl_omits"] = omit
# leaf rows in COO P&L that are not linked to depts while depts carry values
unlinked = []
for r in PL_ROWS:
    if "dept" in pl_types[r]:
        continue
    lab = out_v["COO P&L"][f"A{r}"].value
    tot = 0.0
    for d in DEPTS:
        for c in MONTHS26 + MONTHS27:
            tot += abs(num(val(out_v, d, f"{c}{r}")) or 0.0)
    if tot > 0 and "internal" not in pl_types[r]:
        unlinked.append({"row": r, "label": lab, "types": pl_types[r], "abs_dept_values": tot})
F["coo_pl_unlinked"] = unlinked

# (g2) subtotal formulas skipping lines (hierarchy from labels)
def build_sections(ws):
    labels = {r: (ws[f"A{r}"].value or "").strip() for r in PL_ROWS}
    labels = {r: l for r, l in labels.items() if l and l != "\xa0"}
    rows = sorted(labels)
    sections = {}
    for t in rows:
        T = labels[t]
        if not T.startswith("Total - "):
            continue
        L = T[len("Total - "):]
        above = [r for r in rows if r < t and labels[r] == L]
        if not above:
            continue
        h = max(above)
        # a contiguous run of identical labels (e.g. NetSuite "(Summary)" header + own-posting line):
        # the topmost row of the run is the header, the others are leaf lines
        while (h - 1) in labels and labels[h - 1] == L:
            h -= 1
        sections[t] = h
    return labels, sections

REF_SAME = lambda col: re.compile(r"(?<![A-Z$])\$?" + col + r"\$?(\d+)(?::\$?" + col + r"\$?(\d+))?")

def refs_in(f, col):
    f2 = re.sub(r"(?:'[^']+'|[A-Za-z]+)!\$?[A-Z]{1,3}\$?\d+", "", f)
    out = set()
    for m in REF_SAME(col).finditer(f2):
        a = int(m.group(1)); b = int(m.group(2)) if m.group(2) else a
        out.update(range(a, b + 1))
    return out

def has_values(r, col_list, sheet):
    t = 0.0
    for c in col_list:
        t += abs(num(val(out_v, sheet, f"{c}{r}")) or 0.0)
    return t

subtot = []
for d in ["COO P&L"] + DEPTS:
    ws = out_f[d]
    labels, sections = build_sections(ws)
    headers = set(sections.values())
    for tot, hdr in sections.items():
        inner = [r for r in labels if hdr < r < tot]
        # direct children: skip rows inside nested sections
        nested = [(h, t) for t, h in sections.items() if hdr < h and t < tot]
        child = []
        for r in inner:
            if any(h <= r < t for h, t in nested):
                continue
            if r in headers:
                continue
            child.append(r)
        for col in ("B", "U", "AF"):
            f = ftext(ws[f"{col}{tot}"].value)
            if not f:
                continue
            used = refs_in(f, col)
            missing = [r for r in child if r not in used]
            extra = [r for r in used if r not in child and r not in headers]
            if missing or extra:
                subtot.append({"sheet": d, "total_row": tot, "label": labels[tot], "col": col, "formula": f,
                               "expected_children": child, "missing": missing, "extra": sorted(extra),
                               "missing_values_abs": {r: has_values(r, MONTHS26 + MONTHS27, d) for r in missing},
                               "missing_labels": {r: labels[r] for r in missing}})
        # values parked on header rows (excluded from totals)
        hv = has_values(hdr, MONTHS26 + MONTHS27, d)
        used_any = set().union(*[refs_in(ftext(ws[f"{c}{tot}"].value) or "", c) for c in ("B", "U", "AF")])
        if hv and hdr not in used_any:
            subtot.append({"sheet": d, "total_row": tot, "label": labels[tot], "col": "-", "formula": "",
                           "expected_children": child, "missing": [hdr], "extra": [],
                           "missing_values_abs": {hdr: hv}, "missing_labels": {hdr: labels[hdr] + " (header row carries values)"}})
F["subtotals"] = subtot

# special calc rows
spec = []
EXPECT = {64: ("16", "63", "-"), 117: ("64", "116", "-"), 131: ("122", "130", "-"), 132: ("117", "131", "+")}
for d in ["COO P&L"] + DEPTS:
    ws = out_f[d]
    for r, (a, b, op) in EXPECT.items():
        for col in ("B", "U", "AF"):
            f = ftext(ws[f"{col}{r}"].value) or ""
            if f.replace("$", "") != f"={col}{a}{op}{col}{b}":
                spec.append({"sheet": d, "cell": f"{col}{r}", "formula": f, "expected": f"={col}{a}{op}{col}{b}"})
F["special_rows"] = spec

# (h) hidden sheets / rows / cols
hid = []
for book, wbf in (("COO", coo_f), ("SD", sd_f), ("OUTPUT", out_f)):
    for ws in wbf.worksheets:
        hr = [k for k, d_ in ws.row_dimensions.items() if d_.hidden]
        hc = [k for k, d_ in ws.column_dimensions.items() if d_.hidden]
        if ws.sheet_state != "visible" or hr or hc:
            hid.append({"book": book, "sheet": ws.title, "state": ws.sheet_state,
                        "hidden_rows": len(hr), "hidden_cols": hc[:20]})
F["hidden"] = hid

# (i) T&E reclass
te = sd_v["Fld Maint T&E Data"]
recl = []
for r in range(5, te.max_row + 1):
    if str(te[f"M{r}"].value).strip().lower() == "yes":
        recl.append({"row": r, "date": str(te[f"A{r}"].value)[:10], "vendor": te[f"B{r}"].value,
                     "amount": num(te[f"C{r}"].value), "booked_gl": te[f"D{r}"].value, "budget_gl": te[f"E{r}"].value,
                     "event": te[f"G{r}"].value, "tag_basis": te[f"J{r}"].value, "status": te[f"K{r}"].value})
F["reclass"] = recl

# (j) dependence on Old Draft
od = []
for book, wbf in (("SD", sd_f), ("OUTPUT", out_f)):
    for ws in wbf.worksheets:
        n = 0; ex = []
        for row in ws.iter_rows():
            for c in row:
                f = ftext(c.value)
                if f and "Old Draft" in f:
                    n += 1
                    if len(ex) < 5: ex.append(f"{c.coordinate}: {f[:120]}")
        if n:
            od.append({"book": book, "sheet": ws.title, "refs": n, "examples": ex})
F["old_draft"] = od
# sheet dependency map SD
SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_\.\[\]]+))!")
depmap = {}
for ws in sd_f.worksheets:
    ds = {}
    for row in ws.iter_rows():
        for c in row:
            f = ftext(c.value)
            if f:
                for m in SHEET_REF.finditer(f):
                    s = (m.group(1) or m.group(2)).replace("''", "'")
                    ds[s] = ds.get(s, 0) + 1
    depmap[ws.title] = ds
F["sd_depmap"] = depmap

# DeploySum driver rows: hardcoded vs formula
ds = sd_f["DeploySum"]
drv = {"formula": 0, "const": 0, "rows_const": []}
for r in range(53, 67):
    k = [("f" if ftext(ds.cell(r, ci).value) else "n") for ci in range(2, 19) if ds.cell(r, ci).value is not None]
    if k and all(x == "n" for x in k):
        drv["rows_const"].append(r)
F["deploysum_drivers"] = drv

# Prod Payroll status
F["prod_payroll"] = {"errors_before": R["errors_before_sd"].get("Prod Payroll"),
                     "errors_after": R["errors_after"].get("Prod Payroll"),
                     "refs_into_prod_payroll_from_output": [
                         f"{ws.title}!{c.coordinate}: {ftext(c.value)[:80]}"
                         for ws in out_f.worksheets if ws.title != "Prod Payroll"
                         for row in ws.iter_rows() for c in row
                         if ftext(c.value) and "Prod Payroll" in ftext(c.value)]}

# defined names in COO book
dn = coo_f.defined_names
F["defined_names"] = {"coo_total": len(dn), "coo_ref_errors": sum(1 for _, d_ in dn.items() if "#REF" in (d_.attr_text or "")),
                      "output_total": len(out_f.defined_names)}

# are any COO-book defined names used by formulas in the output?
NAMES = {n.upper() for n in coo_f.defined_names.keys()}
name_use = 0
for ws in out_f.worksheets:
    if ws.title == "Collation Notes":
        continue
    for row in ws.iter_rows():
        for c in row:
            f = ftext(c.value)
            if not f:
                continue
            f2 = re.sub(r'"[^"]*"', "", f); f2 = re.sub(r"#(N/A|REF!|NAME\?|VALUE!|DIV/0!|NUM!|NULL!)", "", f2)
            f2 = re.sub(r"'[^']*'!", "", f2); f2 = re.sub(r"\b[A-Za-z]+!", "", f2)
            f2 = re.sub(r"\$?[A-Z]{1,3}\$?\d+", "", f2); f2 = re.sub(r"\$?[A-Z]{1,3}:\$?[A-Z]{1,3}", "", f2)
            for m_ in re.finditer(r"\b([A-Za-z_][A-Za-z0-9_.]*)(?![A-Za-z0-9_.(])", f2):
                if m_.group(1).upper() in NAMES and m_.group(1).upper() not in ("TRUE", "FALSE"):
                    name_use += 1
F["name_usage"] = name_use

# YoY sign convention check (AJ)
yoy = []
for d in DEPTS + ["COO P&L"]:
    ws = out_f[d]
    for r in PL_ROWS:
        f = ftext(ws[f"AJ{r}"].value)
        if f:
            fs = f.replace("$", "")
            kind = "AG-N" if fs == f"=AG{r}-N{r}" else ("N-AG" if fs == f"=N{r}-AG{r}" else "other")
            yoy.append((d, r, kind))
cost_rows_ag_n = [(d, r) for d, r, k in yoy if k == "AG-N" and 17 < r < 64 and r not in (64,)]
F["yoy_sign"] = {"cost_rows_using_AG_minus_N": cost_rows_ag_n[:50],
                 "other": [(d, r) for d, r, k in yoy if k == "other"][:50]}

R["findings"] = F
json.dump(R, open(RES_JSON, "w"), indent=1, default=str)

# console summary
print("rollup checked", rollup_checked, "fails", len(rollup_fail), "max diff", worst)
for k, v in R["fidelity"].items():
    print("fidelity", k, "cells", v["cells"], "mismatch", v["mismatch_n"], "new_err", v["new_errors_n"],
          "err_changed", v["err_code_changed_n"], "cleared", v["errors_cleared_n"], "maxdiff", round(v["num_max_abs_diff"], 6))
print("formula counts", fcount)
for x in pl: print("PL", x)
print("diff2026 rows", {d: len(v) for d, v in d26.items()}, "formula diffs", {d: v["n"] for d, v in fdiff.items()})
print("hardcodes", len(hard), "placeholders", len(ph), "variance", len(var), "omits", len(omit),
      "unlinked", len(unlinked), "subtot", len(subtot), "special", len(spec), "reclass", len(recl))
