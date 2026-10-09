"""Collation Notes sheet: render (same generator as build_v4.py / build_v5.py) and extract (inverse), for v6.

render(old_xml, notes, styles_xml) -> (xml, last_row, n_formulas)
extract(sheet_xml) -> notes dict {"title", "col_widths", "sections": [...]}

extract() reads the Collation Notes part of a delivered build back into the JSON structure that render() takes.
Usage as a script (round-trip check):
  python3 -I cnotes_v6.py BUILD_XLSX OUT_JSON
It extracts the notes from BUILD_XLSX, renders them again with render(), and stops unless the result is byte-identical
to the part in BUILD_XLSX. So the v6 notes start from exactly what v5 shipped.
"""
import sys, re, json, zipfile
from xml.sax.saxutils import escape, unescape
from openpyxl.utils import get_column_letter, column_index_from_string

ST = {"title": 1, "heading": 2, "header": 3, "text": 4, "num": 5, "line": 0}


def render(old, notes, styles_xml):
    xfs = re.findall(r"<xf .*?(?:/>|</xf>)", re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles_xml, re.S).group(1), re.S)
    assert 'numFmtId="165"' in xfs[5] and 'wrapText="true"' in xfs[4] and 'fillId="2"' in xfs[3], "style table changed"
    rows = {}

    def put(r, c, val, style):
        ref = f"{get_column_letter(c)}{r}"
        if val is None:
            return
        if isinstance(val, dict):
            f = escape(val["f"])
            if val.get("t") == "str":
                x = f'<c r="{ref}" s="{style}" t="str"><f>{f}</f><v>{escape(str(val["v"]))}</v></c>'
            else:
                x = f'<c r="{ref}" s="{style}"><f>{f}</f><v>{repr(float(val["v"]))}</v></c>'
        elif isinstance(val, bool):
            x = f'<c r="{ref}" s="{style}" t="b"><v>{int(val)}</v></c>'
        elif isinstance(val, (int, float)):
            x = f'<c r="{ref}" s="{style}"><v>{repr(val) if isinstance(val, float) else val}</v></c>'
        else:
            s = str(val)
            if s.startswith("="):
                s = s[1:]
            s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)
            x = f'<c r="{ref}" s="{style}" t="inlineStr"><is><t xml:space="preserve">{escape(s)}</t></is></c>'
        rows.setdefault(r, []).append(x)

    r = 1
    put(r, 1, notes["title"], ST["title"]); r += 2
    maxc = 1; nform = 0
    for sec in notes["sections"]:
        put(r, 1, sec["heading"], ST["header"] if sec.get("style") == "highlight" else ST["heading"]); r += 1
        for line in sec.get("lines", []):
            put(r, 1, line, ST["line"]); r += 1
        t = sec.get("table")
        if t:
            for j, h in enumerate(t["header"], 1):
                put(r, j, h, ST["header"])
            maxc = max(maxc, len(t["header"])); r += 1
            for row in t["rows"]:
                for j, v in enumerate(row, 1):
                    if isinstance(v, dict):
                        nform += 1
                        put(r, j, v, ST["text"] if v.get("t") == "str" else ST["num"])
                    elif isinstance(v, (int, float)) and not isinstance(v, bool):
                        put(r, j, v, ST["num"])
                    else:
                        put(r, j, v, ST["text"])
                r += 1
            for line in t.get("after", []) + sec.get("after", []):
                put(r, 1, line, ST["header"]); r += 1
        r += 1
    last = max(rows)
    widths = notes["col_widths"]
    cols = "".join(f'<col min="{j}" max="{j}" width="{w}" customWidth="1"/>' for j, w in enumerate(widths, 1))
    data = "".join(f'<row r="{k}">{"".join(rows[k])}</row>' for k in sorted(rows))
    head = old[:old.index("<sheetPr")]
    view = re.search(r"<sheetViews>.*?</sheetViews>", old, re.S).group(0)
    tail = old[old.index("</sheetData>") + len("</sheetData>"):]
    xml = (head + '<sheetPr filterMode="false"><pageSetUpPr fitToPage="false"/></sheetPr>'
           + f'<dimension ref="A1:{get_column_letter(max(maxc, len(widths)))}{last}"/>' + view
           + '<sheetFormatPr defaultRowHeight="15"/>' + f"<cols>{cols}</cols><sheetData>{data}</sheetData>" + tail)
    return xml, last, nform


def _num(s):
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    return float(s)


def extract(xml):
    widths = []
    for w in re.findall(r'<col min="\d+" max="\d+" width="([^"]+)" customWidth="1"/>', re.search(r"<cols>(.*?)</cols>", xml).group(1)):
        widths.append(_num(w))
    rows = {}
    for rnum, body in re.findall(r'<row r="(\d+)">(.*?)</row>', re.search(r"<sheetData>(.*?)</sheetData>", xml, re.S).group(1), re.S):
        cells = {}
        for ref, s, rest in re.findall(r'<c r="([A-Z]+)\d+" s="(\d+)"(.*?)</c>', body, re.S):
            col = column_index_from_string(ref); s = int(s)
            if rest.startswith(' t="inlineStr"><is><t xml:space="preserve">'):
                v = unescape(re.match(r' t="inlineStr"><is><t xml:space="preserve">(.*)</t></is>$', rest, re.S).group(1))
            elif rest.startswith(' t="str"><f>'):
                m = re.match(r' t="str"><f>(.*?)</f><v>(.*?)</v>$', rest, re.S)
                v = {"f": unescape(m.group(1)), "t": "str", "v": unescape(m.group(2))}
            elif rest.startswith("><f>"):
                m = re.match(r"><f>(.*?)</f><v>(.*?)</v>$", rest, re.S)
                v = {"f": unescape(m.group(1)), "v": float(m.group(2))}
            elif rest.startswith(' t="b"><v>'):
                v = bool(int(re.match(r' t="b"><v>(\d)</v>$', rest).group(1)))
            else:
                v = _num(re.match(r"><v>(.*?)</v>$", rest).group(1))
            cells[col] = (s, v)
        rows[int(rnum)] = cells
    keys = sorted(rows)
    assert keys[0] == 1 and rows[1][1][0] == ST["title"]
    notes = {"title": rows[1][1][1], "col_widths": widths, "sections": []}
    # blocks of consecutive rows, from row 3
    blocks, cur = [], []
    for k in keys[1:]:
        if cur and k != cur[-1] + 1:
            blocks.append(cur); cur = []
        cur.append(k)
    if cur:
        blocks.append(cur)
    for b in blocks:
        h = rows[b[0]][1]
        sec = {"heading": h[1]}
        if h[0] == ST["header"]:
            sec["style"] = "highlight"
        lines, table, i = [], None, 1
        while i < len(b) and set(rows[b[i]]) == {1} and rows[b[i]][1][0] == ST["line"]:
            lines.append(rows[b[i]][1][1]); i += 1
        if i < len(b):
            hdr = rows[b[i]]
            assert all(s == ST["header"] for s, _ in hdr.values())
            n = max(hdr)
            table = {"header": [hdr[j][1] for j in range(1, n + 1)], "rows": []}
            i += 1
            after = []
            while i < len(b):
                rr = rows[b[i]]
                if set(rr) == {1} and rr[1][0] == ST["header"]:
                    after.append(rr[1][1])
                else:
                    assert not after, "data row after an 'after' line"
                    table["rows"].append([rr[j][1] if j in rr else None for j in range(1, max(n, max(rr)) + 1)])
                i += 1
            table["after"] = after
        if lines:
            sec["lines"] = lines
        if table:
            sec["table"] = table
        notes["sections"].append(sec)
    return notes


if __name__ == "__main__":
    BUILD, OUT = sys.argv[1:3]
    z = zipfile.ZipFile(BUILD)
    wb = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(r'<sheet name="Collation Notes"[^>]*r:id="(rId\d+)"', wb).group(1)
    part = "xl/" + re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)
    xml = z.read(part).decode()
    notes = extract(xml)
    again, last, nform = render(xml, notes, z.read("xl/styles.xml").decode())
    assert again == xml, "Collation Notes round trip is not byte-identical"
    json.dump(notes, open(OUT, "w"), indent=1, ensure_ascii=False)
    print(f"Collation Notes extracted from {BUILD}: {len(notes['sections'])} sections, {last} rows, {nform} formulas; "
          f"round trip byte-identical")
