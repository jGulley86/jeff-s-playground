"""Cell-by-cell comparison of two builds (formulas and cached values) on the budget sheets.

Usage:
  python3 -I diff_versions.py OLD_XLSX NEW_XLSX OUT_JSON

Compares every cell in the used range (union of both) of: COO P&L, FO, PO, CO, FE, PE and every driver tab
present in both files. Formula text is compared exactly; cached values are compared exactly for text/errors/dates
and with tolerance 0 for numbers (any difference counts). Comments are compared separately (they are not values).
Exit code 1 if any formula or value differs.
"""
import sys, json, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

OLD, NEW, OUT = sys.argv[1:4]
CORE = ["COO P&L", "FO", "PO", "CO", "FE", "PE"]


def norm(v):
    if isinstance(v, ArrayFormula):
        return ("array", v.ref, v.text)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


of = openpyxl.load_workbook(OLD); ov = openpyxl.load_workbook(OLD, data_only=True)
nf = openpyxl.load_workbook(NEW); nv = openpyxl.load_workbook(NEW, data_only=True)
drivers = [s for s in of.sheetnames if s in nf.sheetnames and s not in CORE and s != "Collation Notes"]
res = {"old": OLD, "new": NEW, "sheets": {}, "comments_added": [], "comments_removed": [], "comments_changed": []}
tot_f = tot_v = cells = 0
for s in CORE + drivers:
    a_f, b_f, a_v, b_v = of[s], nf[s], ov[s], nv[s]
    mr = max(a_f.max_row, b_f.max_row); mc = max(a_f.max_column, b_f.max_column)
    df = dv = n = 0; ex = []
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            fa, fb = norm(a_f.cell(r, c).value), norm(b_f.cell(r, c).value)
            va, vb = norm(a_v.cell(r, c).value), norm(b_v.cell(r, c).value)
            if fa is None and fb is None and va is None and vb is None:
                continue
            n += 1
            coord = f"{get_column_letter(c)}{r}"
            if fa != fb:
                df += 1
                if len(ex) < 10: ex.append(["formula", coord, str(fa)[:80], str(fb)[:80]])
            if va != vb:
                dv += 1
                if len(ex) < 10: ex.append(["value", coord, str(va)[:80], str(vb)[:80]])
            ca, cb = a_f.cell(r, c).comment, b_f.cell(r, c).comment
            ta, tb = (ca.text if ca else None), (cb.text if cb else None)
            if ta != tb:
                key = f"{s}!{coord}"
                if ta is None: res["comments_added"].append([key, tb])
                elif tb is None: res["comments_removed"].append([key, ta])
                else: res["comments_changed"].append([key, ta[:80], tb[:80]])
    res["sheets"][s] = {"cells": n, "formula_diffs": df, "value_diffs": dv, "examples": ex}
    tot_f += df; tot_v += dv; cells += n
res["total"] = {"sheets": len(res["sheets"]), "cells": cells, "formula_diffs": tot_f, "value_diffs": tot_v}
res["sheets_only_in_old"] = [s for s in of.sheetnames if s not in nf.sheetnames]
res["sheets_only_in_new"] = [s for s in nf.sheetnames if s not in of.sheetnames]
json.dump(res, open(OUT, "w"), indent=1, default=str)
print("diff:", res["total"], "| comments added", len(res["comments_added"]), "removed", len(res["comments_removed"]),
      "changed", len(res["comments_changed"]))
sys.exit(1 if (tot_f or tot_v) else 0)
