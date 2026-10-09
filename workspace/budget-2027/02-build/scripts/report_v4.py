"""Write 02-build_notes_v4.md and the v4 Collation Notes content (JSON). Notes-and-presentation fix of v3.

Usage:
  python3 -I report_v4.py V3_MD V3_XLSX FIG_JSON LO_VERSION_TXT COO SD NEW OUT_NOTES_JSON OUT_MD [ANALYZE_JSON]

- Every new number comes from FIG_JSON (figures_v4.py: workbook cached values + inputs). No figure is typed here.
- v3 text is reused section by section from V3_MD; each edit is a targeted replacement that must hit exactly once
  (the script stops otherwise), so nothing is changed silently.
- The assumptions register, issues list, '$000s' relabel table and external-formula list are read from the v3
  Collation Notes sheet (structured cells) and checked against the v3 md before editing.
- Without ANALYZE_JSON the md has placeholders for the diff / package sections (pass 1); the Collation Notes JSON
  does not depend on ANALYZE_JSON, so it is identical in both passes.
"""
import sys, json, re, hashlib, collections
import openpyxl

(V3_MD, V3_XLSX, FIG, LOV, COO, SD, NEW, OUT_JSON, OUT_MD) = sys.argv[1:10]
AN = json.load(open(sys.argv[10])) if len(sys.argv) > 10 else None
F = json.load(open(FIG))
lov = open(LOV).read().strip()
v3md = open(V3_MD).read()
H = F["headline"]; C = F["contribution"]


def sha(p):
    return hashlib.sha256(open(p, "rb").read()).hexdigest()


def f0(x):
    if x is None: return ""
    x = float(x)
    if abs(x) < 0.5: return "0"
    return f"({abs(x):,.0f})" if x < 0 else f"{x:,.0f}"


def f2(x):
    x = float(x)
    return f"({abs(x):,.2f})" if x < 0 else f"{x:,.2f}"


def eff(x):
    """Effect on the 2027 COO contribution: + raises it, ( ) lowers it."""
    x = float(x)
    if abs(x) < 0.5: return "0"
    return f"+{x:,.0f}" if x > 0 else f"({abs(x):,.0f})"


def pct_nm(new, base):
    """% change; 'n/m' when the base is zero or negative, or the sign flips (positive base, negative new)."""
    if base is None or base <= 0 or new < 0:
        return "n/m"
    return f"{(new - base) / base * 100:.0f}%"


def rep(text, old, new, count=1):
    n = text.count(old)
    assert n == count, (n, old[:90])
    return text.replace(old, new)


# ------------------------------------------------------------------ v3 md sections
parts = re.split(r"(?m)^## ", v3md)
head_txt = parts[0]
SEC = collections.OrderedDict()
for p in parts[1:]:
    title, _, body = p.partition("\n")
    SEC[title.strip()] = body
expected = ["Headline", "Changes from v2", "Completion criteria", "Depreciation schedule summary",
            "Unit-cost reconciliation", "Revenue", "Assumptions register", "DeploySum refresh and tie-out",
            "v2 -> v3 diff", "Error scan", "Package", "Sheets in the output", "Severity scale", "Issues (v3)",
            "Carried from v2 (still valid: FO, FE, CO, PE and PO payroll unchanged)", "Decisions", "Not done / limits"]
assert list(SEC) == expected, list(SEC)

# ------------------------------------------------------------------ v3 Collation Notes sheet (structured)
wb3 = openpyxl.load_workbook(V3_XLSX)
cn3 = wb3["Collation Notes"]
sheet_secs = collections.OrderedDict(); cur = None; mode = None
for r in range(1, cn3.max_row + 1):
    vals = [cn3.cell(r, c).value for c in range(1, cn3.max_column + 1)]
    a = cn3.cell(r, 1)
    if all(x is None for x in vals):
        mode = None; continue
    if a.font is not None and a.font.b and a.font.sz == 12 and all(x is None for x in vals[1:]):
        cur = a.value; sheet_secs[cur] = {"lines": [], "header": None, "rows": []}; mode = "start"; continue
    if cur is None:
        continue
    if mode == "start" and a.font is not None and a.font.b:
        while vals and vals[-1] is None: vals.pop()
        sheet_secs[cur]["header"] = vals; mode = "rows"; continue
    if mode == "rows":
        n = len(sheet_secs[cur]["header"])
        sheet_secs[cur]["rows"].append(vals[:n])
    else:
        sheet_secs[cur]["lines"].append(a.value)


def sec_by_prefix(pfx):
    k = [x for x in sheet_secs if x.startswith(pfx)]
    assert len(k) == 1, (pfx, list(sheet_secs))
    return sheet_secs[k[0]]


reg = [list(x) for x in sec_by_prefix("Assumptions register")["rows"]]
iss_rows = [list(x) for x in sec_by_prefix("Issues (v3)")["rows"]]
relabel = sec_by_prefix("'$000s' relabel")
extf = sec_by_prefix("DeploySum external-link formulas")
mdrow = lambda row: "| " + " | ".join("" if x is None else str(x) for x in row) + " |"
# consistency: the sheet rows render to exactly the v3 md rows
for row in reg:
    assert mdrow(row) in SEC["Assumptions register"], row[0]
for row in iss_rows:
    assert mdrow(row) in SEC["Issues (v3)"], row[0]
iss = collections.OrderedDict((row[0], row) for row in iss_rows)
IHDR = ["ID", "Sev", "Location (sheet!cell)", "Finding", "Evidence", "2027 impact", "Basis", "Recommended action"]

# ------------------------------------------------------------------ derived figures (all from F)
q4 = F["q4"]; di = F["dep_inputs"]; i32 = F["i32_5010"]; i02 = F["i02"]
lifts_po, trays_po = sum(q4["po_lifts"]), sum(q4["po_trays"])
lifts_ds, trays_ds = sum(q4["ds_lifts"]), sum(q4["ds_trays"])
dec_l, dec_t = F["i31_dec_units"]
jl, jt = F["jan27_units"]
RR = F["runrate"]
rev = H["Revenue"]["AG"]; cost27 = H["COGS"]["AG"] + H["Opex"]["AG"]
flip = F["flip_alone"]; after = F["after_alone"]

# ------------------------------------------------------------------ 1. Decisions pending
DEC_HDR = ["#", "Decision (user)", "Issue", "Current treatment in the budget",
           "Alternatives: effect on 2027 COO contribution (+ raises, ( ) lowers)", "Who"]
DEC = [
    ["1", "SlipBot fleet status (5010 depreciation)", "I-32",
     f"Carried flat at the PO Dec-26 run rate {f2(RR['19'])}/month = {f0(F['i32_full'])} in 2027 (builder assumption, A-17).",
     f"Fully depreciated or retired from Jan-27: {eff(F['i32_full'])}. Continue the 2026 decline ({f0(i32['slope_jan_sep'])}/month, Jan-Sep 26): {eff(F['i32_trend'])} "
     f"(actuals Jan-Aug only, {f0(i32['slope_jan_aug'])}/month: {eff(F['i32_trend_actuals'])}). Ongoing at the run rate: 0. "
     f"Related: the 5020 base of {f0(F['i32_5020_jan26'])}/month (Jan-26, before any SlipLift depreciation) may end with the SlipBots: up to {eff(F['i32_5020_base_x12'])}.",
     "User, with PO/Finance (fixed-asset register)"],
    ["2", "Q4-26 go-lives: PO run rate or DeploySum", "I-31",
     f"Inside the PO Dec-26 run rate: PO steps up for {lifts_po} lifts / {trays_po} trays (DeploySum: {f0(lifts_ds)} / {f0(trays_ds)}).",
     f"Sep-26 run rate + DeploySum Oct-Dec 26 vintages: {eff(F['i31_short'])}. If PO's Dec-26 step-up ({f0(dec_l)} lifts, {f0(dec_t)} trays) is Jan-27 units already in the Jan-27 vintage and PO's Oct/Nov counts are right: {eff(F['i31_double'])}. Keep as is: 0.",
     "User, after the PO owner says what the Q4-26 step-ups are"],
    ["3", "SlipLift peripherals: GL and useful life", "I-33",
     f"5020 Robot Peripherals, {f0(di['life_periph'])} months: {f0(F['i33_new_periph'])} of 2027 depreciation (builder assumption, A-10/A-11).",
     f"GL 5011 instead of 5020: 0 (moves {f0(F['i33_new_periph'])} from PO row 22 to row 20). Life 60 months: {eff(F['i33_life60'])}. Life {f0(di['life_tray'])} months: {eff(F['i33_life_tray'])}.",
     "User"],
    ["4", "Revenue split across 4100 / 4500 / 4900", "I-01 (A-02, A-03)",
     f"All {f0(F['rev_fy27'])} of plan revenue on 4100; 4500 and 4900 = 0 (orchestrator assumption).",
     f"Split the plan revenue across the three lines: 0 (classification only; gross margin unchanged). Budget implementation / one-time revenue in addition to the plan: not determinable (2026 4500 + 4900: {f0(F['rev_4500_4900_2026'])}).",
     "User / Finance"],
]

# ------------------------------------------------------------------ 2. Headline (renamed)
HL = [("Revenue", "Revenue (company recognized revenue, 9/30/26 sales plan)", "16"),
      ("COGS", "COGS (COO departments)", "63"),
      ("Depreciation", "  of which depreciation (in COGS)", "19-22"),
      ("Gross Profit", "Gross Profit", "64"),
      ("Opex", "Opex (COO departments)", "116"),
      ("Net Ordinary Income", "COO contribution before other income (workbook label 'Net Ordinary Income')", "117"),
      ("Net Other Income", "Net Other Income", "131"),
      ("Net Income", "**COO contribution** (workbook label 'Net Income')", "132")]
hl_md = ["| Line | COO P&L row | 2026 (N) | 2027 v4 (AG) | Change | % |", "|---|---|---:|---:|---:|---:|"]
for k, lab, row in HL:
    n, ag = H[k]["N"], H[k]["AG"]
    hl_md.append(f"| {lab} | {row} | {f0(n)} | {f0(ag)} | {f0(ag - n)} | {pct_nm(ag, n)} |")
nm_rows = [lab for k, lab, row in HL if pct_nm(H[k]["AG"], H[k]["N"]) == "n/m"]

# ------------------------------------------------------------------ 3. Caveat block
CAV_HDR = ["Item", "Direction", "What could happen", "Effect on 2027 COO contribution", "Contribution after this item alone", "Source cells"]
CAV = [
    ["I-17", "downside", "Logistics & warehouse payroll (SD model) is not in PO 5100; if PO's typed production payroll excludes the warehouse team, PO is understated.",
     -F["i17"], C - F["i17"], f"SD Logistics&WH Payroll!S{F['i17_row']} (not carried)"],
    ["I-05", "downside", f"Production payroll model 'Total lift payroll' {f0(F['i05_model'])} vs PO 5100 typed {f0(F['i05_po'])}; gap if the model is right (lift payroll only; tray and supervisor rows do not compute in LibreOffice). May overlap I-17.",
     -F["i05"], C - F["i05"], "Prod Payroll!S41; PO!AG36"],
    ["I-10", "downside", "FO/FE lines with 2026 spend budgeted at 0 for 2027 (5920 software, 6610 SG&A salaries, 7200 contractors); effect if they continue at the 2026 level.",
     -F["i10"], C - F["i10"], "FO!N58, N67, N91; FE!N91 (AG = 0)"],
    ["I-26", "downside", "Head of Service Delivery has no salary entered (current employee).",
     None, None, "Fld Eng Budget!B21 (blank)"],
    ["I-31", "two-way", f"Shortfall if DeploySum's Q4-26 go-lives are right ({f0(lifts_ds - lifts_po)} lifts, {f0(trays_ds - trays_po)} trays more than PO's run rate).",
     F["i31_short"], C + F["i31_short"], "PO!J20:M21; DeploySum!C19:E20"],
    ["I-31", "two-way", f"Double count if PO's Dec-26 step-up ({f0(dec_l)} lifts, {f0(dec_t)} trays; DeploySum Dec-26 go-lives 0) is Jan-27 units ({f0(jl)} lifts, {f0(jt)} trays go live Jan-27, built Nov-26). Excludes the 5020 Dec step ({eff(F['i31_double_5020'])} more if it is the same units).",
     F["i31_double"], C + F["i31_double"], "PO!L20:M21; DeploySum!E19:F20"],
    ["I-32", "upside", "SlipBot fleet fully depreciated or retired from Jan-27 (5010 = 0).",
     F["i32_full"], C + F["i32_full"], "PO!M19; 'Depreciation Schedule'!B21"],
    ["I-32", "upside", f"Trend alternative: 5010 keeps falling at the 2026 rate ({f0(i32['jan'])} Jan-26 -> {f0(i32['sep'])} Sep-26, {f0(i32['slope_jan_sep'])}/month).",
     F["i32_trend"], C + F["i32_trend"], "PO!B19, J19, M19"],
    ["I-33", "two-way", "Peripherals life 60 months instead of 84 (a longer life, e.g. 120 months, gives the opposite sign).",
     F["i33_life60"], C + F["i33_life60"], "'Depreciation Schedule'!B11"],
    ["A-14", "upside", "Depreciation convention: mid-month in the go-live month instead of a full month.",
     -F["sens_mid_month"], C - F["sens_mid_month"], "'Depreciation Schedule'!B13"],
    ["A-14", "upside", "Depreciation convention: start the month after go-live.",
     -F["sens_next_month"], C - F["sens_next_month"], "'Depreciation Schedule'!B13 = 1"],
    ["A-22", "downside", f"Revenue rests on unsigned / forecast bookings (Sep-Dec 26 includes {f0(F['unsigned_sep_dec26'])} unsigned at 100%). Revenue effect not determinable from the workbook.",
     None, None, "DeploySum!A49 (note 5), row 40"],
]
cav_md = ["| " + " | ".join(CAV_HDR) + " |", "|---|---|---|---:|---:|---|"]
for it, d, what, e, a, src in CAV:
    cav_md.append(f"| {it} | {d} | {what} | {'n/d' if e is None else eff(e)} | {'n/d' if a is None else f0(a)} | {src} |")

# sign-flip statement, built from the numbers
alone_flip = [k for k in ("I-05", "I-17", "I-10", "I-31") if flip[k]]
alone_keep = [k for k in ("I-05", "I-17", "I-10", "I-31") if not flip[k]]
assert C > 0
s1 = (f"The 2027 COO contribution of {f0(C)} changes sign on " +
      " or ".join(f"{k} alone ({f0(C)} - {f0(C - after[k])} = {f0(after[k])})" for k in alone_flip) + "." if alone_flip
      else f"No single quantified downside changes the sign of the {f0(C)} contribution.")
s2 = (f"Together, I-17 + I-05 take it to {f0(F['after_i17_i05'])} (if they do not overlap; the PO owner must say). "
      f"I-10 alone leaves {f0(after['I-10'])}; I-10 with the I-31 shortfall gives {f0(F['after_i10_i31'])}. "
      f"All quantified downsides together (I-17, I-05, I-10, I-31 shortfall): {f0(F['after_all_down'])}.")
s3 = (f"Upside: I-32 can add up to {f0(F['i32_full'])} (contribution {f0(F['upside_full'])}); the trend alternative adds {f0(F['i32_trend'])} ({f0(F['upside_trend'])}). "
      f"With costs held fixed, revenue {f0(C)} ({F['rev_breakeven_pct'] * 100:.1f}%) below plan also takes the contribution to zero.")
s4 = (f"Conclusion: the sign of the 2027 COO contribution is not established by this budget. Treat {f0(C)} as a point estimate inside a range from "
      f"{f0(F['after_all_down'])} to {f0(F['upside_full'])} until I-17, I-05 and I-32 are answered.")
assert len(alone_flip) >= 1 and F["after_i17_i05"] < 0 and F["upside_full"] > 0   # the conclusion above depends on these
SIGN = [s1, s2, s3, s4]

CAV1 = (f"(1) What the figure is. The bottom line, {f0(C)}, is a **COO contribution**: the company's recognized revenue ({f0(rev)}, all of it) "
        f"less the costs of the five COO departments only (COGS + Opex of FO, PO, CO, FE, PE: {f0(cost27)}). Costs outside COO (for example sales, G&A, R&D) "
        f"are not in this workbook, so it is **not company Net Income**. The workbook row keeps its GL label 'Net Income' (COO P&L!A132) because v4 does not change the P&L tabs; "
        f"the notes and Collation Notes call it 'COO contribution'.")
CAV2 = (f"(2) What the revenue is. 2027 revenue is the 9/30/26 sales plan's recognized revenue (DeploySum row 40, static cached values, A-24), not signed contracts. "
        f"It includes unsigned and forecast bookings: Sep-Dec 26 bookings of {f0(F['bookings_sep_dec26'])} include {f0(F['unsigned_sep_dec26'])} of unsigned forecast at 100% "
        f"(DeploySum note 5, A49: bookings are gross ACV), and all {f0(F['bookings_fy27'])} of FY27 bookings are forecast (new logos {f0(F['bookings_fy27_newlogo'])}, expansion {f0(F['bookings_fy27_expansion'])}). "
        f"{f0(F['rev_above_exit'])} ({F['rev_above_exit_pct'] * 100:.0f}%) of 2027 revenue is above the Dec-26 exit rate x 12 ({f0(F['rev_exit_runrate_x12'])}), so it depends on go-lives that have not happened yet.")
CAV3 = "(3) Open exposures, each on its own against the headline. 'n/d' = not determinable from the workbook. Items can overlap (I-05 and I-17 both concern production payroll), so do not add them without the owners' answers."

# ------------------------------------------------------------------ 4. issue edits (O1, O2, O3, O6, B1)
def set_issue(i, **kw):
    row = iss[i]
    for k, v in kw.items():
        row[IHDR.index(k)] = v

# O6: I-02 restated in cost terms (v2 evidence quoted Net Income before revenue was budgeted)
drivers = iss["I-02"][4].split("Drivers:")[1].strip()
set_issue("I-02",
          **{"Location (sheet!cell)": "FO!AG63; FO!AG116; FE!AG63; FE!AG116; COO P&L!AG132",
             "Finding": f"Replacing FO and FE with the SD versions adds {f0(i02['delta'])} to 2027 COO cost (COGS + Opex). The SD build is much heavier than the COO book's FO/FE. (v4: restated in cost terms; the earlier evidence quoted 2027 Net Income from before revenue was budgeted.)",
             "Evidence": f"FO cost (AG63 + AG116): COO book {f0(i02['FO_coo'])} -> SD {f0(i02['FO_sd'])}; FE: COO book {f0(i02['FE_coo'])} -> SD {f0(i02['FE_sd'])}; together {f0(i02['FO_coo'] + i02['FE_coo'])} -> {f0(i02['FO_sd'] + i02['FE_sd'])}. With the COO book's FO/FE, the 2027 COO contribution would be {f0(i02['contribution_with_coo_book'])} instead of {f0(C)}. Drivers: {drivers}",
             "2027 impact": f"{f0(-i02['delta'])} on the contribution",
             "Basis": f"impact {f0(-i02['delta'])}"})
# B1 / I-05: quantify the gap
iss["I-05"][4] = rep(iss["I-05"][4], "Cells read from Prod Payroll by other sheets",
                     f"Gap: {f0(F['i05_model'])} - {f0(F['i05_po'])} = {f0(F['i05'])} (lift payroll only). If the model is right and PO 5100 is the whole production payroll, the 2027 COO contribution falls by at least this gap, which on its own turns it negative ({f0(after['I-05'])}). May overlap I-17. Cells read from Prod Payroll by other sheets")
iss["I-05"][5] = f"not determinable (exposure {f0(F['i05'])}: lift payroll model {f0(F['i05_model'])} vs PO 5100 {f0(F['i05_po'])})"
# O1: I-31 two-way
set_issue("I-31",
          **{"Finding": (f"Oct-Dec 26 go-lives are carried in the PO Dec-26 run rate, not added explicitly. The risk runs both ways. "
                         f"(a) Shortfall: PO steps up for {lifts_po} lifts and {trays_po} trays in Q4-26, DeploySum has {f0(lifts_ds)} and {f0(trays_ds)}; if DeploySum is right the budget is short {f0(-F['i31_short'])}. "
                         f"(b) Possible double count: PO's Dec-26 step-up ({f0(dec_l)} lifts, {f0(dec_t)} trays) has no DeploySum Dec-26 go-live. If these are units DeploySum places in Jan-27 ({f0(jl)} lifts, {f0(jt)} trays, built Nov-26), they sit in both the run rate and the Jan-27 vintage: {f0(F['i31_double'])} too much. "
                         f"If (b) holds and DeploySum's Q4 counts are right, the gross Q4 shortfall is {f0(-F['i31_gross_short_if_double'])} and the net is still {f0(F['i31_short'])}. The 5020 step-ups cannot be reproduced from DeploySum's peripheral cost."),
             "2027 impact": f"{eff(F['i31_short'])} to {eff(F['i31_double'])} (5020 Dec step: up to {eff(F['i31_double_5020'])} more)",
             "Basis": f"impact {f0(max(-F['i31_short'], F['i31_double']))} (< $50k)",
             "Recommended action": "PO owner to say what the Oct/Nov/Dec-26 step-ups on PO rows 20-22 are (which go-lives, which month). Then the user chooses: keep the PO run rate (current), or Sep-26 run rate + DeploySum Oct-Dec 26 vintages. Do not do both."})
iss["I-31"][4] = rep(iss["I-31"][4], "Gap not budgeted:", f"Dec-26 step-ups with no DeploySum Dec-26 go-live: {f0(dec_l)} lifts x {f2(di['m_lift'])} + {f0(dec_t)} trays x {f2(di['m_tray'])} = {f2(F['i31_double'] / 12)}/month, {f0(F['i31_double'])} for 2027. Gap not budgeted:")
# O2, O3: I-32 attribution, trend alternative, 5020 base
iss["I-32"][3] = rep(iss["I-32"][3], "for all 12 months of 2027, as the handoff instructed.",
                     "for all 12 months of 2027. The flat carry is a builder assumption (A-17): the brief asked for the PO 2026 monthly depreciation as the existing-fleet input, not for a flat carry.")
iss["I-32"][3] += (f" Trend alternative: 5010 fell through 2026 ({f0(i32['jan'])} Jan -> {f0(i32['sep'])} Sep, {f0(i32['slope_jan_sep'])}/month); continuing that decline through 2027 lowers depreciation by {f0(F['i32_trend'])} "
                   f"({f0(F['i32_trend_actuals'])} on the Jan-Aug actuals trend, {f0(i32['slope_jan_aug'])}/month). The 5020 base ({f0(F['i32_5020_jan26'])} in Jan-26, when 5011 was {f0(F['i32_5011_jan26'])}) predates SlipLift depreciation, so it is probably SlipBot-related and may share the SlipBot end-of-life risk ({f0(F['i32_5020_base_x12'])} for 2027 at that rate).")
iss["I-32"][5] = f"up to {eff(F['i32_full'])} (5010 = 0); trend {eff(F['i32_trend'])}; 5020 base up to {eff(F['i32_5020_base_x12'])}"
iss["I-32"][7] = rep(iss["I-32"][7], "USER DECISION: confirm SlipBot fleet status (fully depreciated / retired / ongoing)",
                     "USER DECISION: confirm SlipBot fleet status (fully depreciated / retired / ongoing / declining as in 2026)")
# I-07: attribution of the flat carry
iss["I-07"][3] = rep(iss["I-07"][3], "Existing fleet = PO Dec-26 run rate, carried flat.", "Existing fleet = PO Dec-26 run rate, carried flat (builder assumption, A-17).")
sev_count = collections.Counter(r[1] for r in iss.values())

# ------------------------------------------------------------------ 5. register edits (O3, O4, O5) and new lines
regd = collections.OrderedDict((r[0], r) for r in reg)
regd["A-14"][2] = rep(regd["A-14"][2], "sensitivity: 1-month delay lowers 2027 by 274,833",
                      f"sensitivity: 1-month delay lowers 2027 depreciation by {f0(-F['sens_next_month'])}; mid-month convention lowers it by {f0(-F['sens_mid_month'])}")
regd["A-17"][2] = f"{regd['A-17'][2]}; trend alternative on 5010 {eff(F['i32_trend'])} to the contribution"
regd["A-18"][1] = "Oct-Dec 26 go-lives covered by the PO run rate, not added explicitly; the risk runs both ways"
regd["A-18"][2] = f"{eff(F['i31_short'])} if DeploySum's Q4 counts are right; {eff(F['i31_double'])} if PO's Dec-26 step-ups are Jan-27 units"
regd["A-18"][4] = "builder decision between the brief's two options; PO owner to confirm what the step-ups are (I-31)"
regd["A-19"][4] = rep(regd["A-19"][4], "orchestrator instruction to carry;", "builder assumption (flat carry, A-17);")
regd["A-22"] = ["A-22", "Revenue includes unsigned and forecast bookings, taken at 100%",
                f"Sep-Dec 26 bookings {f0(F['bookings_sep_dec26'])} incl. {f0(F['unsigned_sep_dec26'])} unsigned; FY27 bookings {f0(F['bookings_fy27'])} all forecast",
                "DeploySum!A49 (note 5); DeploySum rows 7-13, 40", "source: 9/30/26 sales plan; no probability weighting applied (builder: taken as given)"]
regd["A-23"] = ["A-23", "Existing-fleet base is the PO Dec-26 forecast, not an actual",
                f"{f2(sum(RR.values()))}/month (PO!M19:M22; Sep-Dec 26 are forecast columns J:M; actuals run to Aug-26)",
                "'Depreciation Schedule'!B21:B24", "builder note (PO 2026 file); refresh when Q4-26 actuals close"]
regd["A-24"] = ["A-24", "DeploySum F40:Q40 (2027 revenue) are static cached values from the 9/30/26 file, not live links",
                f"{F['ds40_static_values']} of 12 months static; PO!U12:AF12 link to them (formulas)",
                "DeploySum!F40:Q40 -> PO!U12:AF12", "build v3 (external workbook [1] not supplied); a plan change needs a new resolved export"]
regd["A-25"] = ["A-25", "Revenue start timing vs go-live: revenue follows the plan's own recognition timing; depreciation starts in the go-live month (A-14)",
                "not determinable: the workbook cannot show whether revenue starts in the go-live month",
                "DeploySum row 40 vs rows 19/20; 'Depreciation Schedule'!B13", "builder note; plan owner to confirm (Revenue Roll-up not supplied)"]
reg = list(regd.values())
RHDR = ["ID", "Assumption", "Value / effect", "Where", "Source"]

# ------------------------------------------------------------------ 6. variance scan with n/m (O8)
var_body = SEC["Issues (v3)"].split("### Variance scan")[1]
wbv = openpyxl.load_workbook(V3_XLSX, data_only=True)
var_md = ["| Sheet | Row | Line | 2026 N | 2027 AG | Change | % |", "|---|---:|---|---:|---:|---:|---:|"]
nvar = nnm = 0
for ln in var_body.splitlines():
    m = re.match(r"\| (COO P&L|FO|PO|CO|FE|PE) \| (\d+) \| (.*?) \| (.*?) \| (.*?) \| (.*?) \| (.*?) \|$", ln)
    if not m: continue
    s, r = m.group(1), int(m.group(2))
    n, ag = wbv[s][f"N{r}"].value or 0, wbv[s][f"AG{r}"].value or 0
    assert (f0(n), f0(ag), f0(ag - n)) == (m.group(4), m.group(5), m.group(6)), ln
    p = pct_nm(ag, n); nvar += 1; nnm += p == "n/m"
    var_md.append(f"| {s} | {r} | {m.group(3)} | {f0(n)} | {f0(ag)} | {f0(ag - n)} | {p} |")

# ------------------------------------------------------------------ 7. md assembly
L = []; A = L.append
A("# 02-build notes v4 - 2027 COO budget collation")
A("")
A("## Decisions pending")
A("")
A("Four user decisions are open. The builder has not resolved them; the budget carries the 'current treatment' column. Effects are on the 2027 COO contribution, each on its own.")
A("")
A("| " + " | ".join(DEC_HDR) + " |"); A("|---|---|---|---|---|---|")
for d in DEC: A(mdrow(d))
A("")
A(f"v4 is a notes-and-presentation fix of v3 (critic v3: B1 and O1-O8). **Budget values are unchanged**: only the Collation Notes sheet and these notes changed. "
  + (f"Diff vs v3: {AN['diff']['total']['formula_diffs']} formula and {AN['diff']['total']['value_diffs']} value diffs on {AN['diff']['total']['sheets']} sheets ({AN['diff']['total']['cells']:,} cells); see 'v3 -> v4 diff'."
     if AN else "Diff vs v3: see 'v3 -> v4 diff'."))
A("")
A("## Headline: 2027 COO contribution")
A("")
L.extend(hl_md)
A("")
A(f"- **2027 COO contribution (COO P&L AG132): {f0(C)}** vs 2026 {f0(H['Net Income']['N'])}. Change {f0(C - H['Net Income']['N'])}; % n/m (2026 base negative). It includes 2027 plan revenue of {f0(rev)} (all to 4100, orchestrator assumption) and depreciation of {f0(H['Depreciation']['AG'])}.")
A(f"- Cost-only view: total COO cost (COGS + Opex, AG63 + AG116) {f0(cost27)} vs 2026 {f0(H['COGS']['N'] + H['Opex']['N'])} (change {f0(cost27 - H['COGS']['N'] - H['Opex']['N'])}, {pct_nm(cost27, H['COGS']['N'] + H['Opex']['N'])}). Excluding depreciation: {f0(cost27 - H['Depreciation']['AG'])} vs 2026 {f0(H['COGS']['N'] + H['Opex']['N'] - H['Depreciation']['N'])}.")
A(f"- 2027 gross margin {H['Gross Profit']['AG'] / rev * 100:.1f}%. '%' is 'n/m' where the 2026 base is zero or negative or the sign flips: {'; '.join(x.replace('**', '') for x in nm_rows)}.")
A(f"- Issues: {sev_count['HIGH']} HIGH, {sev_count['MED']} MED, {sev_count['LOW']} LOW (no severity changes in v4).")
A("")
A("### Caveats on the headline (read before using the figure)")
A("")
A(CAV1); A(""); A(CAV2); A(""); A(CAV3); A("")
L.extend(cav_md)
A("")
A("**Sign-flip statement (computed from the table):** " + " ".join(SIGN))
A("")
A("## Changes from v3")
A("")
CHG = [
    ["B1", "Headline caveats", f"Headline relabelled 'COO contribution' (md and Collation Notes); caveat block under the headline in both: (1) contribution, not company NI; (2) revenue = 9/30/26 sales plan incl. {f0(F['unsigned_sep_dec26'])} unsigned bookings; (3) exposure table (I-17, I-05, I-10, I-26, I-31 both ways, I-32 incl. trend, I-33, A-14, A-22) and a computed sign-flip statement. 'Decisions pending' added at the top.", "Decisions pending; Headline; Caveats; Collation Notes rows 3 onward"],
    ["O1", "I-31 / A-18 two-way", f"Reworded: shortfall {eff(F['i31_short'])} vs possible double count {eff(F['i31_double'])}; PO owner to confirm the step-ups. The 'so there is no double count' sentence in the depreciation section was replaced.", "I-31; A-18; Depreciation schedule summary"],
    ["O2", "I-32 trend and 5020", f"Trend alternative {eff(F['i32_trend'])} (actuals-only {eff(F['i32_trend_actuals'])}); 5020 Jan-26 base {f0(F['i32_5020_jan26'])}/month flagged as likely SlipBot-related ({eff(F['i32_5020_base_x12'])}).", "I-32; Decisions pending #1; caveat table"],
    ["O3", "Flat carry attribution", "Flat carry is a builder assumption everywhere in the notes: A-19 source, I-07, I-32 and the SlipBot paragraph changed (A-17 already said so). The Depreciation Schedule tab text was not edited (v4 freezes every sheet except Collation Notes).", "A-19; I-07; I-32; Depreciation schedule summary"],
    ["O4", "Revenue timing; mid-month", f"A-25 added (revenue start vs go-live: not determinable). A-14 now shows the mid-month sensitivity (2027 depreciation lower by {f0(-F['sens_mid_month'])}) next to the next-month start (lower by {f0(-F['sens_next_month'])}).", "A-14; A-25"],
    ["O5", "Missing assumptions", "A-22 unsigned/forecast bookings; A-23 existing base is a Dec-26 forecast; A-24 DeploySum F40:Q40 static cached values.", "Assumptions register"],
    ["O6", "Stale I-02 evidence", f"I-02 restated in cost terms from the inputs and the workbook: +{f0(i02['delta'])} COO cost; contribution with the COO book's FO/FE {f0(i02['contribution_with_coo_book'])}.", "I-02"],
    ["O7", "Stale Prod Payroll trace", f"Paragraph rewritten with current counts ({F['prodpay']['errors']} error cells of {F['prodpay']['formulas']} formulas, all #NAME?); the 'still valid' heading removed; v1-review labels (O2, O3, O10) dropped from sub-headings.", "Cross-department, payroll and T&E checks"],
    ["O8", "% on negative / sign-flip bases", f"'n/m' in the headline (md and live Collation Notes formula) and in the variance scan ({nnm} of {nvar} rows).", "Headline; Variance scan; Collation Notes"],
]
A("| Item | Critic point | What v4 did | Where |"); A("|---|---|---|---|")
for c in CHG: A(mdrow(c))
A("")
A("No critic item was declined. One partial: O3's attribution is fixed in the notes and on Collation Notes, but the explanatory text on the Depreciation Schedule tab (B17, D21, row 19 heading) is unchanged because the brief allows no edits outside Collation Notes.")
A("")
A("## Build and sources")
A("")
A("Output: `02-build/02-build_COO_2027_Budget_Collated_v4.xlsx` (v1-v3 files unchanged)  ")
A("Scripts: `02-build/scripts/` - v4 entry point `run_build_v4.sh` (figures_v4.py, report_v4.py, build_v4.py, diff_versions.py, recalc.sh as a check, analyze_v4.py). It starts from the delivered v3 workbook and rewrites only the Collation Notes sheet part of the package; every other part is copied byte for byte. `run_build_v3.sh`, `run_build_v2.sh` and `run_build.sh` still reproduce v3, v2 and v1.  ")
A(f"Inputs (read-only): `COO_Dept_PLs_2027_Budget_2.xlsx` sha256 `{sha(COO)}`; `Service_Delivery_PL_worksheets.xlsx` sha256 `{sha(SD)}`; `2027_Sales_Deploy_Prod_to_COO_as_of_093026.xlsx` sha256 `{sha(NEW)}`  ")
A("Briefs: `00-brief/brief_v2.md`, `00-brief/brief_v3.md`; critic review `03-review/03-review_critic_v3.md`; verifier `04-verify/04-verify_report_v3.md` (v3 PASS; v4 values identical).  ")
A(f"Recalculation engine (check only, not used for the deliverable): {lov}.")
A("")
A("## Completion criteria (v4)")
A("")
if AN:
    d = AN["diff"]["total"]
    A("| Criterion | Result |"); A("|---|---|")
    A(f"| 0 diffs vs v3 on every sheet except Collation Notes | {d['formula_diffs']} formula / {d['value_diffs']} value diffs on {d['sheets']} sheets, {d['cells']:,} cells (diff_versions.py, exact, no tolerance). Package: {AN['package']['identical_parts']} of {AN['package']['parts_v3']} parts byte-identical to v3; changed: {', '.join(AN['package']['changed_parts'])}. |")
    A(f"| Figures recompute from v4 | figures_v4.py on v4 = on v3: {AN['figures_identical']}. |")
    A(f"| Collation Notes live formulas | {AN['cn']['formulas']} formulas; LibreOffice recalc of v4 vs stored values: max diff {AN['cn']['max_num_diff']:.1e}, text mismatches {AN['cn']['text_mismatch']}. |")
    A("| Every critic item addressed or declined with a reason | B1, O1-O8: table 'Changes from v3' (none declined; O3 partial on the frozen tab, reason given). |")
    A(f"| No stale v2 text | Checked phrases absent: {', '.join(repr(x) for x in AN.get('stale_checked', []))}. |" if AN.get("stale_checked") else "| No stale v2 text | see report_v4.py checks |")
    A(f"| % shows n/m on negative or sign-flipping bases | Headline: {len(nm_rows)} rows n/m; variance scan {nnm} of {nvar} rows n/m; Collation Notes % formula uses the same rule. |")
else:
    A("(pass 1 placeholder)")
A("")
# carried sections
dep = SEC["Depreciation schedule summary"]
dep = rep(dep, re.search(r"\*\*Decision: rely on the run rate\.\*\*.*\n", dep).group(0),
          (f"**Treatment: rely on the run rate (builder decision; two-way risk, I-31).** Oct-Dec 26 go-lives are not added explicitly. Why: the run rate is PO's own forecast and already contains Q4-26 step-ups priced like the schedule; adding DeploySum's Q4 units on top would count the {lifts_po} lifts and {trays_po} trays already in it twice. "
           f"The risk runs both ways. If DeploySum's Q4 counts are right, the run rate is short {f0(lifts_ds - lifts_po)} lifts / {f0(trays_ds - trays_po)} trays ({f0(-F['i31_short'])} for 2027; memo in schedule section 5). "
           f"If PO's Dec-26 step-up ({f0(dec_l)} lifts, {f0(dec_t)} trays; DeploySum has no Dec-26 go-lives) is units that DeploySum places in Jan-27, they are also in the Jan-27 vintage: {f0(F['i31_double'])} counted twice. PO owner to confirm; user decision (Decisions pending #2).\n"))
dep = rep(dep, "This carries the run rate through 2027 as instructed; it is not evidence that the fleet is still depreciating.",
          f"The flat carry is a builder assumption (A-17), not evidence that the fleet is still depreciating. 5010 fell through 2026 ({f0(i32['slope_jan_sep'])}/month Jan-Sep); continuing that trend would lower 2027 by {f0(F['i32_trend'])}.")
A("## Depreciation schedule summary"); A(dep.rstrip()); A("")
A("## Unit-cost reconciliation"); A(SEC["Unit-cost reconciliation"].rstrip()); A("")
A("## Revenue"); A(SEC["Revenue"].rstrip()); A("")
A(f"Revenue basis (A-22, A-24, A-25): plan figures, not contracts. Sep-Dec 26 bookings {f0(F['bookings_sep_dec26'])} include {f0(F['unsigned_sep_dec26'])} unsigned at 100%; FY27 bookings {f0(F['bookings_fy27'])} are forecast. "
  f"DeploySum!F40:Q40 hold static values ({F['ds40_static_values']} of 12). Dec-26 exit rate {f0(F['rev_dec26_monthly'])}/month x 12 = {f0(F['rev_exit_runrate_x12'])}; the other {f0(F['rev_above_exit'])} of 2027 revenue depends on 2027 go-lives.")
A("")
A("## Assumptions register"); A("")
A("| " + " | ".join(RHDR) + " |"); A("|---|---|---|---|---|")
for r in reg: A(mdrow(r))
A("")
A("A-22..A-25 are new in v4; A-14, A-17, A-18 and A-19 were edited (critic O1, O3, O4).")
A("")
A("## DeploySum refresh and tie-out (v3 build; unchanged in v4)"); A(SEC["DeploySum refresh and tie-out"].rstrip()); A("")
A("## v3 -> v4 diff"); A("")
if AN:
    A("`diff_versions.py` (exact: formulas as text, cached values with no tolerance) on every sheet present in both files except Collation Notes:")
    A("")
    A("| Sheet | Cells | Formula diffs | Value diffs |"); A("|---|---:|---:|---:|")
    for s, x in AN["diff"]["sheets"].items():
        A(f"| {s} | {x['cells']:,} | {x['formula_diffs']} | {x['value_diffs']} |")
    A(f"| **Total** | {AN['diff']['total']['cells']:,} | {AN['diff']['total']['formula_diffs']} | {AN['diff']['total']['value_diffs']} |")
    A("")
    A(f"Comments added / removed / changed: {AN['diff']['comments']}. Sheets only in v3 / only in v4: {AN['diff']['only']}. "
      f"Package: {AN['package']['identical_parts']} of {AN['package']['parts_v3']} parts byte-identical; changed: {', '.join(AN['package']['changed_parts'])}; added: {AN['package']['added_parts'] or 'none'}; removed: {AN['package']['removed_parts'] or 'none'}. "
      f"Why not a LibreOffice round trip: re-saving v4 through LibreOffice changes {AN['lo_roundtrip']['value_diffs']} cached values vs v3 (max abs {AN['lo_roundtrip']['max_abs']:.1e}; sheets {AN['lo_roundtrip']['sheets']}), so v4 patches the Collation Notes part instead and uses LibreOffice only to check its formulas.")
A("")
A("## Error scan"); A("")
if AN:
    A("| Sheet | v3 | v4 |"); A("|---|---|---|")
    for s, x in AN["errors"].items(): A(f"| {s} | {x['v3'] or 'none'} | {x['v4'] or 'none'} |")
    A(""); A(f"New errors in v4: **{len(AN['new_errors'])}**. Prod Payroll's #NAME? cells are the XLOOKUP rows LibreOffice cannot evaluate (I-05, I-25).")
A("")
A("## Package"); A("")
if AN:
    A(f"Parts {AN['package']['parts_v4']}; `xl/externalLinks/*` **{AN['package']['externalLink_parts']}**; defined names {AN['package']['defined_names_v4']:,} (v3: {AN['package']['defined_names_v3']:,}).")
A("")
A("## Sheets in the output"); A("")
if AN:
    A("| # | Sheet | State |"); A("|---|---|---|")
    for i, (s, st) in enumerate(AN["sheets"], 1): A(f"| {i} | {s} | {st} |")
A("")
A("## Severity scale"); A(SEC["Severity scale"].rstrip()); A("")
A("## Issues (v4)"); A("")
A("Severity per the scale above. v4 changed the text of I-02 (O6), I-05 (B1), I-07 and I-32 (O2, O3) and I-31 (O1); no severity changed. All other rows are as in v3.")
A(""); A("| " + " | ".join(IHDR) + " |"); A("|---|---|---|---|---|---|---|---|")
for r in iss.values(): A(mdrow(r))
A("")
A("### Variance scan: rows with |AG-N| > $50k and > 50% (values unchanged from v3)")
A("")
A("'%' = (AG - N) / N; 'n/m' where N is zero or negative, or AG is negative while N is positive.")
A("")
L.extend(var_md)
A("")
car = SEC["Carried from v2 (still valid: FO, FE, CO, PE and PO payroll unchanged)"]
car = rep(car, "### Cross-department overlap (from v2)", "### Cross-department overlap")
car = rep(car, "### Raise timing consistency (O10)", "### Raise timing consistency")
car = rep(car, "### Logistics & Warehouse payroll (O2) and MeLi (O3)", "### Logistics & Warehouse payroll and MeLi")
old_trace = re.search(r"Rates and timing only, no headcount and no cost:.*\n", car).group(0)
pp = F["prodpay"]
car = rep(car, old_trace,
          (f"Rates and timing only: other sheets read only these typed raise inputs from Prod Payroll, never a headcount or cost cell. "
           f"Prod Payroll now partly computes (v3 resolved DeploySum): 'Total lift payroll' S41 = {f0(F['i05_model'])}, but {pp['errors']} cells are errors ({', '.join(f'{k} {v}' for k, v in pp['codes'].items())}; XLOOKUP rows that LibreOffice cannot evaluate, see I-05), out of {pp['formulas']} formulas. "
           f"PO 5110-5140 2027 (PO!U33:AF35) are typed values with no link to any tab, so PO AG33 {f0(F['po_ag33'])} and FO 5510 share no headcount source: no double count through Prod Payroll. "
           f"The open question is the reverse: the model's lift payroll exceeds PO 5100 ({f0(F['i05_po'])}) by {f0(F['i05'])} (I-05).\n"))
A("## Cross-department, payroll and T&E checks (first built in v2; FO, FE, CO, PE and PO payroll values unchanged since)")
A(car.rstrip()); A("")
dec = SEC["Decisions"].rstrip()
A("## Decisions"); A(dec)
for x in [
    "D1 (v4): Patch, do not round-trip. v4 copies the v3 package and rewrites only xl/worksheets/sheet1.xml (Collation Notes). A LibreOffice re-save changes cached values in the last floating-point digit"
    + (f" ({AN['lo_roundtrip']['value_diffs']} cells, max {AN['lo_roundtrip']['max_abs']:.1e}, in this run)" if AN else "")
    + ", which would break the '0 value diffs' requirement; LibreOffice is used only to check the new sheet's formulas.",
    "D2 (v4): Collation Notes strings are written as inline strings, and existing v3 cell styles are reused, so sharedStrings.xml and styles.xml stay byte-identical (the v3 notes strings remain in sharedStrings.xml, unused).",
    "D3 (v4): Formula cells on Collation Notes store values computed by figures_v4.py; analyze_v4.py checks them against a LibreOffice recalculation of v4.",
    "D4 (v4): '% change' = n/m when the base is zero or negative or the new value is negative with a positive base; the live formula returns text ('n/m' or e.g. '110%') so no new number format is needed.",
    "D5 (v4): The workbook's 'Net Income' label (COO P&L!A132 and dept tabs) is not renamed, because the brief allows no change outside Collation Notes; the notes and Collation Notes use 'COO contribution'.",
    "D6 (v4): Exposures are shown one at a time against the headline and not netted, because I-05 and I-17 may overlap; combined figures are shown with that caveat.",
    "D7 (v4): The I-32 trend alternative uses the Jan->Sep-26 slope (critic's basis, Sep is the first forecast month) and also shows the Jan->Aug actuals-only slope; each is floored so 5010 cannot go negative.",
    "D8 (v4): Decisions pending are presented, not resolved; the budget keeps the v3 treatment.",
]:
    A(f"- {x}")
A("")
A("## Not done / limits")
A(SEC["Not done / limits"].rstrip())
A("- v4: revenue effect of unsigned/forecast bookings and the revenue-vs-go-live timing (A-22, A-25) cannot be quantified from the workbook; the Revenue Roll-up and 2026 Exist New tabs of the plan were not supplied.")
A("- v4: I-26 (missing salary) has no amount in the workbook, so it is n/d in the caveat table.")
A("- v4: the explanatory text on the Depreciation Schedule tab still describes the flat carry neutrally and the Q4-26 choice as 'Never both'; it was not edited (sheet frozen in v4).")
md = "\n".join(L) + "\n"

# stale-text checks (completion criterion 'no stale v2 text')
STALE = ["as the handoff instructed", "as instructed;", "before (5,277,004)", "796 of Prod Payroll", "still valid",
         "not added explicitly, so there is no double count", "orchestrator instruction to carry", "(O10)", "(O2)", "(O3)", "2027 v2 (AG)",
         "Changes from v2", "Issues (v3)"]
# the change log and the completion table quote the removed phrases; check everything else
md_chk = re.sub(r"(?s)## Changes from v3\n.*?\n## ", "## ", md)
md_chk = re.sub(r"(?s)## Completion criteria \(v4\)\n.*?\n## ", "## ", md_chk)
assert "## Changes from v3" not in md_chk and "## Completion criteria" not in md_chk
for s in STALE:
    assert s not in md_chk, s
# % column: n/m on every negative base in the headline
for k, lab, row in HL:
    if H[k]["N"] <= 0 or H[k]["AG"] < 0:
        assert pct_nm(H[k]["AG"], H[k]["N"]) == "n/m"
open(OUT_MD, "w").write(md)

# ------------------------------------------------------------------ 8. Collation Notes JSON (independent of ANALYZE_JSON)
CP = "'COO P&L'!"
def fnum(expr, value):
    return {"f": expr, "v": value}
hl_rows = []
for k, lab, row in HL:
    n, ag = H[k]["N"], H[k]["AG"]
    if k == "Depreciation":
        nf, af = f"SUM({CP}N19:N22)", f"SUM({CP}AG19:AG22)"
    else:
        nf, af = f"{CP}N{row}", f"{CP}AG{row}"
    pf = f'IF(OR({nf}<=0,{af}<0),"n/m",TEXT(({af}-{nf})/{nf},"0%"))'
    hl_rows.append([lab.replace("**", ""), row, fnum(nf, n), fnum(af, ag), fnum(f"{af}-({nf})", ag - n),
                    {"f": pf, "v": pct_nm(ag, n), "t": "str"}])
cav_rows = [[it, d, what, ("n/d" if e is None else round(e, 2)), ("n/d" if a is None else round(a, 2)), src]
            for it, d, what, e, a, src in CAV]
notes = {"title": "2027 COO Budget - Collation Notes (build v4, draft for review)", "stale_checked": STALE,
         "col_widths": [22, 12, 44, 60, 60, 30, 30, 50],
         "sections": [
             {"heading": "Decisions pending (user) - not resolved by the builder", "style": "highlight",
              "table": {"header": DEC_HDR, "rows": DEC}},
             {"heading": "Headline: 2027 COO contribution (live formulas on COO P&L)",
              "table": {"header": ["Line", "COO P&L row", "2026 (N)", "2027 (AG)", "Change", "% (n/m if base <= 0 or sign flips)"], "rows": hl_rows}},
             {"heading": "Caveats on the headline (read before using the figure)", "style": "highlight",
              "lines": [CAV1.replace("**", ""), CAV2.replace("**", ""), CAV3],
              "table": {"header": CAV_HDR, "rows": cav_rows},
              "after": ["Sign-flip statement (computed): " + SIGN[0], SIGN[1], SIGN[2], SIGN[3]]},
             {"heading": "Source and build",
              "lines": [f"Base: v3 collated workbook; only this sheet changed in v4 (0 formula / 0 value diffs on every other sheet). Inputs unchanged; plan file sha256 {sha(NEW)[:16]}...",
                        "Full notes: 02-build/02-build_notes_v4.md. Numbers on this sheet: live formulas (headline) or values computed by figures_v4.py from this workbook and the inputs."]},
             {"heading": "Changes from v3 (critic v3 items)", "table": {"header": ["Item", "Critic point", "What v4 did", "Where"], "rows": CHG}},
             {"heading": "Assumptions register", "table": {"header": RHDR, "rows": reg}},
             {"heading": "Issues (v4) - severity: HIGH = blocks sign-off or > $250k; MED = $50k-250k or owner confirmation; LOW = < $50k or hygiene",
              "table": {"header": ["ID", "Sev", "Location", "Finding", "Evidence", "2027 impact", "Basis", "Recommended action"], "rows": list(iss.values())}},
             {"heading": "'$000s' relabel (DeploySum)", "table": {"header": relabel["header"], "rows": relabel["rows"]}},
             {"heading": "DeploySum external-link formulas in the 9/30/26 file, stored as cached values",
              "table": {"header": extf["header"], "rows": extf["rows"]}},
         ]}
json.dump(notes, open(OUT_JSON, "w"), indent=1, default=str, sort_keys=True)
print(f"report_v4: wrote {OUT_MD} ({len(md):,} chars, analyze={'yes' if AN else 'no'}) and {OUT_JSON}; issues {dict(sev_count)}; "
      f"register {len(reg)}; variance n/m {nnm}/{nvar}")
