"""v7 scenario helper (V1 / critic O4): the I-28 'remove from FO' effect with every knock-on, by recalculation.

Usage:
  python3 -I scen_v7.py make V6_XLSX OUT_SCRATCH_XLSX
      Scratch copy of the v6 workbook with Fld Ops Payroll!B11 (current techs) reduced by one. Only that one <v> changes;
      every other zip part is copied byte for byte. The scratch file is analysis input only and is never shipped.
  python3 -I scen_v7.py compare V6_XLSX BASE_RECALC_XLSX SCEN_RECALC_XLSX OUT_JSON
      Compares the LibreOffice recalculation of v6 (base) with the recalculation of the scratch copy (scenario):
      COO contribution (COO P&L!AG132), every COO P&L row whose AG value moves, and the driver cells that move on
      Fld Ops Payroll and Fld Maint Budget. Checks that the base recalculation equals v6's cached AG132.
"""
import sys, re, json, zipfile
import openpyxl

SHEET = "Fld Ops Payroll"


def part_of(z, name):
    wb = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(rf'<sheet name="{re.escape(name)}"[^>]*r:id="(rId\d+)"', wb).group(1)
    return "xl/" + re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)


def make(v6, out):
    zin = zipfile.ZipFile(v6)
    part = part_of(zin, SHEET)
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    old = None
    for info in zin.infolist():
        data = zin.read(info.filename)
        if info.filename == part:
            x = data.decode("utf-8")
            m = re.search(r'(<c r="B11"[^>]*>)<v>([0-9.]+)</v></c>', x)
            assert m and "<f" not in m.group(1), "Fld Ops Payroll!B11 must be a typed number"
            old = float(m.group(2))
            new = old - 1
            x = x[:m.start()] + f"{m.group(1)}<v>{new:g}</v></c>" + x[m.end():]
            data = x.encode("utf-8")
        zout.writestr(info, data)
    zout.close()
    print(f"scratch {out}: {SHEET}!B11 {old:g} -> {old - 1:g}")


def compare(v6, base, scen, out):
    wv = openpyxl.load_workbook(v6, data_only=True)
    wb = openpyxl.load_workbook(base, data_only=True)
    ws = openpyxl.load_workbook(scen, data_only=True)
    num = lambda v: float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
    c_v6 = num(wv["COO P&L"]["AG132"].value)
    c_b = num(wb["COO P&L"]["AG132"].value); c_s = num(ws["COO P&L"]["AG132"].value)
    assert abs(c_b - c_v6) < 1e-6, ("base recalculation differs from v6 cached AG132", c_b, c_v6)
    rows = []
    sh = wb["COO P&L"]; ss = ws["COO P&L"]
    for r in range(1, sh.max_row + 1):
        a, b = num(sh.cell(r, 33).value), num(ss.cell(r, 33).value)
        if abs(a - b) > 1e-6:
            rows.append({"row": r, "label": sh.cell(r, 1).value, "base": a, "scen": b, "diff": b - a})
    drivers = {}
    for name in ("Fld Ops Payroll", "Fld Maint Budget"):
        d = []
        for rr in wb[name].iter_rows():
            for c in rr:
                b = ws[name][c.coordinate].value
                if isinstance(c.value, (int, float)) and isinstance(b, (int, float)) and abs(c.value - b) > 1e-6:
                    d.append(c.coordinate)
        drivers[name] = {"n_cells": len(d), "first": d[:12]}
    # Fld Maint Budget: which COO P&L GL rows carry the knock-on (rows other than FO 5510-5540)
    fo_pay = [x for x in rows if str(x["label"]).split(" - ")[0] in ("5510", "5530", "5540")]
    gl_rows = [x for x in rows if x["row"] < 117 and re.match(r"\d{4} - ", str(x["label"]))]   # GL lines only (no totals, no memo rows)
    other = [x for x in gl_rows if x not in fo_pay]
    res = {"contribution_v6": c_v6, "contribution_base_recalc": c_b, "contribution_scen": c_s, "effect": c_s - c_b,
           "B11_base": num(wb[SHEET]["B11"].value), "B11_scen": num(ws[SHEET]["B11"].value),
           "fo_payroll_effect": -sum(x["diff"] for x in fo_pay), "knock_on_effect": -sum(x["diff"] for x in other),
           "knock_on_rows": [{"row": x["row"], "gl": str(x["label"]), "effect": -x["diff"]} for x in other],
           "coo_pl_rows_moved": rows, "driver_cells_moved": drivers,
           "sup_hc_dec27_base": num(wb[SHEET]["Q66"].value), "sup_hc_dec27_scen": num(ws[SHEET]["Q66"].value),
           "sup_hc_by_month_base": [num(wb[SHEET].cell(66, c).value) for c in range(6, 18)],
           "sup_hc_by_month_scen": [num(ws[SHEET].cell(66, c).value) for c in range(6, 18)]}
    assert abs(res["fo_payroll_effect"] + res["knock_on_effect"] - res["effect"]) < 1e-6, "effect does not split into rows"
    json.dump(res, open(out, "w"), indent=1, default=str)
    print(f"scen_v7: B11 {res['B11_base']:g} -> {res['B11_scen']:g}: contribution {c_b:,.2f} -> {c_s:,.2f} "
          f"(effect {res['effect']:+,.2f} = FO payroll {res['fo_payroll_effect']:+,.2f} + knock-on {res['knock_on_effect']:+,.2f} "
          f"on {[x['gl'] for x in res['knock_on_rows']]})")


if __name__ == "__main__":
    if sys.argv[1] == "make":
        make(*sys.argv[2:4])
    elif sys.argv[1] == "compare":
        compare(*sys.argv[2:6])
    else:
        raise SystemExit("usage: scen_v7.py make|compare ...")
