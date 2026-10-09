import sys, openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
wb = openpyxl.load_workbook(sys.argv[1]); wv = openpyxl.load_workbook(sys.argv[1], data_only=True)
for ws in wb.worksheets:
    n=0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value,(ArrayFormula,DataTableFormula)):
                n+=1
                if n<=6: print(ws.title, c.coordinate, type(c.value).__name__, getattr(c.value,'ref',None), str(getattr(c.value,'text',''))[:200], '=>', wv[ws.title][c.coordinate].value)
    if n: print(ws.title, "total arrays", n)
