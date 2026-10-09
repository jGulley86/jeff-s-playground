"""v4 build: notes-and-presentation fix only. Copy the delivered v3 package and rewrite ONLY the Collation Notes part.

Usage:
  python3 -I build_v4.py V3_XLSX NOTES_JSON OUT_XLSX

Why a package patch: re-saving the workbook through openpyxl drops cached values, and a LibreOffice re-save changes
some cached values in the last floating-point digit. Either would break '0 value diffs vs v3'. So:
  - every zip part except xl/worksheets/sheet1.xml (the 'Collation Notes' sheet, checked via workbook.xml + rels)
    is copied byte for byte, in the same order;
  - sheet1.xml is regenerated from NOTES_JSON with inline strings (sharedStrings.xml untouched) and the v3 cell
    styles already used on that sheet (styles.xml untouched): 1 title, 2 heading, 3 bold header/highlight,
    4 wrapped text, 5 wrapped number '#,##0;(#,##0)'.
  - formula cells are written with the value computed by report_v4.py ({"f", "v", "t"}); analyze_v4.py checks them
    against a LibreOffice recalculation.
"""
import sys, json, re, zipfile
from xml.sax.saxutils import escape
from openpyxl.utils import get_column_letter

V3, NOTES, OUT = sys.argv[1:4]
zin = zipfile.ZipFile(V3)
wbxml = zin.read("xl/workbook.xml").decode()
rels = zin.read("xl/_rels/workbook.xml.rels").decode()
rid = re.search(r'<sheet name="Collation Notes"[^>]*r:id="(rId\d+)"', wbxml).group(1)
target = re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)
PART = "xl/" + target
assert PART == "xl/worksheets/sheet1.xml", PART
old = zin.read(PART).decode()
assert "<legacyDrawing" not in old and "_rels/sheet1.xml.rels" not in " ".join(zin.namelist())
ST = {"title": 1, "heading": 2, "header": 3, "text": 4, "num": 5, "line": 0}
# the styles must still mean what they meant on the v3 sheet
st = zin.read("xl/styles.xml").decode()
xfs = re.findall(r"<xf .*?(?:/>|</xf>)", re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", st, re.S).group(1), re.S)
assert 'numFmtId="165"' in xfs[5] and 'wrapText="true"' in xfs[4] and 'fillId="2"' in xfs[3], "style table changed"

notes = json.load(open(NOTES))
rows = {}          # r -> list of cell xml


def put(r, c, val, style):
    ref = f"{get_column_letter(c)}{r}"
    if val is None:
        return
    if isinstance(val, dict):                       # formula with stored value
        f = escape(val["f"])
        if val.get("t") == "str":
            x = f'<c r="{ref}" s="{style}" t="str"><f>{f}</f><v>{escape(str(val["v"]))}</v></c>'
        else:
            x = f'<c r="{ref}" s="{style}"><f>{f}</f><v>{repr(float(val["v"]))}</v></c>'
    elif isinstance(val, bool):
        x = f'<c r="{ref}" s="{style}" t="b"><v>{int(val)}</v></c>'
    elif isinstance(val, (int, float)):
        x = f'<c r="{ref}" s="{style}"><v>{repr(val) if isinstance(val, float) else val}</v></c>'
    else:
        s = str(val)
        if s.startswith("="):
            s = s[1:]                                   # same rule as v3: text never becomes a formula
        s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)
        x = f'<c r="{ref}" s="{style}" t="inlineStr"><is><t xml:space="preserve">{escape(s)}</t></is></c>'
    rows.setdefault(r, []).append(x)


r = 1
put(r, 1, notes["title"], ST["title"]); r += 2
maxc = 1
nform = 0
for sec in notes["sections"]:
    put(r, 1, sec["heading"], ST["header"] if sec.get("style") == "highlight" else ST["heading"]); r += 1
    for line in sec.get("lines", []):
        put(r, 1, line, ST["line"]); r += 1
    t = sec.get("table")
    if t:
        for j, h in enumerate(t["header"], 1):
            put(r, j, h, ST["header"])
        maxc = max(maxc, len(t["header"])); r += 1
        for row in t["rows"]:
            for j, v in enumerate(row, 1):
                if isinstance(v, dict):
                    nform += 1
                    put(r, j, v, ST["text"] if v.get("t") == "str" else ST["num"])
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    put(r, j, v, ST["num"])
                else:
                    put(r, j, v, ST["text"])
            r += 1
        for line in t.get("after", []) + sec.get("after", []):
            put(r, 1, line, ST["header"]); r += 1
    r += 1
last = max(rows)
widths = notes["col_widths"]
cols = "".join(f'<col min="{j}" max="{j}" width="{w}" customWidth="1"/>' for j, w in enumerate(widths, 1))
data = "".join(f'<row r="{k}">{"".join(rows[k])}</row>' for k in sorted(rows))
head = old[:old.index("<sheetPr")]                  # xml decl + <worksheet ...> with the v3 namespaces
view = re.search(r"<sheetViews>.*?</sheetViews>", old, re.S).group(0)
tail = old[old.index("</sheetData>") + len("</sheetData>"):]   # printOptions, margins, page setup, header/footer
xml = (head + '<sheetPr filterMode="false"><pageSetUpPr fitToPage="false"/></sheetPr>'
       + f'<dimension ref="A1:{get_column_letter(max(maxc, len(widths)))}{last}"/>' + view
       + '<sheetFormatPr defaultRowHeight="15"/>' + f"<cols>{cols}</cols><sheetData>{data}</sheetData>" + tail)

zout = zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED)
for info in zin.infolist():
    if info.filename == PART:
        zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
        zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = info.external_attr
        zout.writestr(zi, xml.encode("utf-8"))
    else:
        zout.writestr(info, zin.read(info.filename))     # byte-identical content, same order
zout.close()
print(f"built {OUT}: replaced {PART} ({last} rows, {nform} formulas); {len(zin.infolist()) - 1} other parts copied unchanged")
