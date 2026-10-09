"""v5 build: package patch of the delivered v4 workbook.

Usage:
  python3 -I build_v5.py V4_XLSX PATCH_JSON OUT_XLSX [NOTES_JSON]

PATCH_JSON: {"input": [sheet, coord, number], "values": [[sheet, coord, raw_v_text], ...]}
  - "input": the one typed input that changes (Fld Eng Budget!B21). The cell must exist in v4 with no value and no
    formula (v4: '<c r="B21" s="274"/>'); it is written as a number with the same style.
  - "values": cached values of formula cells, written as the exact <v> text LibreOffice produced when it recalculated
    the workbook with the new input (patches_v5.py). Only the <v> text of an existing numeric formula cell is
    replaced; the formula, style and type are kept. Any other shape stops the script.
NOTES_JSON (optional): the Collation Notes sheet part is regenerated from it exactly as build_v4.py does
  (inline strings, v4 styles; formula cells carry the values computed by report_v5.py). Without it the v4 Collation
  Notes part is copied unchanged (stage 1 / stage 2 builds).
Every other zip part is copied byte for byte, in the same order. sharedStrings.xml and styles.xml are untouched.
"""
import sys, json, re, zipfile
from xml.sax.saxutils import escape
from openpyxl.utils import get_column_letter

V4, PATCH, OUT = sys.argv[1:4]
NOTES = sys.argv[4] if len(sys.argv) > 4 else None
zin = zipfile.ZipFile(V4)
wbxml = zin.read("xl/workbook.xml").decode()
rels = zin.read("xl/_rels/workbook.xml.rels").decode()


def part_of(sheet):
    nm = escape(sheet, {'"': "&quot;"})
    rid = re.search(rf'<sheet name="{re.escape(nm)}"[^>]*r:id="(rId\d+)"', wbxml).group(1)
    return "xl/" + re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)


P = json.load(open(PATCH))
by_part = {}
s_in, c_in, v_in = P["input"]
by_part.setdefault(part_of(s_in), []).append(("input", c_in, v_in))
for s, c, raw in P["values"]:
    by_part.setdefault(part_of(s), []).append(("value", c, raw))
CN_PART = part_of("Collation Notes")
assert CN_PART == "xl/worksheets/sheet1.xml" and CN_PART not in by_part, "Collation Notes is never value-patched"

stats = {"input": 0, "values": 0}


def patch_sheet(xml, items):
    for kind, coord, val in items:
        if kind == "input":
            pat = re.compile(rf'<c r="{coord}"( s="\d+")?/>')
            m = pat.findall(xml)
            assert len(pat.findall(xml)) == 1, (coord, "input cell must be an empty styled cell in v4")
            assert isinstance(val, (int, float)) and not isinstance(val, bool)
            xml = pat.sub(lambda mm: f'<c r="{coord}"{mm.group(1) or ""} t="n"><v>{val}</v></c>', xml, count=1)
            stats["input"] += 1
        else:
            pat = re.compile(rf'<c r="{coord}"((?: s="\d+")?(?: t="n")?)><f([^>]*)>([^<]*)</f><v>([^<]*)</v></c>')
            hits = pat.findall(xml)
            assert len(hits) == 1, (coord, len(hits), "expected one numeric formula cell")
            float(val)                                         # the new cached value must be a number
            xml = pat.sub(lambda mm: f'<c r="{coord}"{mm.group(1)}><f{mm.group(2)}>{mm.group(3)}</f><v>{val}</v></c>', xml, count=1)
            stats["values"] += 1
    return xml


# ------------------------------------------------------------------ Collation Notes (same generator as build_v4.py)
def notes_xml(old, notes):
    st = zin.read("xl/styles.xml").decode()
    xfs = re.findall(r"<xf .*?(?:/>|</xf>)", re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", st, re.S).group(1), re.S)
    assert 'numFmtId="165"' in xfs[5] and 'wrapText="true"' in xfs[4] and 'fillId="2"' in xfs[3], "style table changed"
    ST = {"title": 1, "heading": 2, "header": 3, "text": 4, "num": 5, "line": 0}
    rows = {}

    def put(r, c, val, style):
        ref = f"{get_column_letter(c)}{r}"
        if val is None:
            return
        if isinstance(val, dict):
            f = escape(val["f"])
            if val.get("t") == "str":
                x = f'<c r="{ref}" s="{style}" t="str"><f>{f}</f><v>{escape(str(val["v"]))}</v></c>'
            else:
                x = f'<c r="{ref}" s="{style}"><f>{f}</f><v>{repr(float(val["v"]))}</v></c>'
        elif isinstance(val, bool):
            x = f'<c r="{ref}" s="{style}" t="b"><v>{int(val)}</v></c>'
        elif isinstance(val, (int, float)):
            x = f'<c r="{ref}" s="{style}"><v>{repr(val) if isinstance(val, float) else val}</v></c>'
        else:
            s = str(val)
            if s.startswith("="):
                s = s[1:]
            s = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", s)
            x = f'<c r="{ref}" s="{style}" t="inlineStr"><is><t xml:space="preserve">{escape(s)}</t></is></c>'
        rows.setdefault(r, []).append(x)

    r = 1
    put(r, 1, notes["title"], ST["title"]); r += 2
    maxc = 1; nform = 0
    for sec in notes["sections"]:
        put(r, 1, sec["heading"], ST["header"] if sec.get("style") == "highlight" else ST["heading"]); r += 1
        for line in sec.get("lines", []):
            put(r, 1, line, ST["line"]); r += 1
        t = sec.get("table")
        if t:
            for j, h in enumerate(t["header"], 1):
                put(r, j, h, ST["header"])
            maxc = max(maxc, len(t["header"])); r += 1
            for row in t["rows"]:
                for j, v in enumerate(row, 1):
                    if isinstance(v, dict):
                        nform += 1
                        put(r, j, v, ST["text"] if v.get("t") == "str" else ST["num"])
                    elif isinstance(v, (int, float)) and not isinstance(v, bool):
                        put(r, j, v, ST["num"])
                    else:
                        put(r, j, v, ST["text"])
                r += 1
            for line in t.get("after", []) + sec.get("after", []):
                put(r, 1, line, ST["header"]); r += 1
        r += 1
    last = max(rows)
    widths = notes["col_widths"]
    cols = "".join(f'<col min="{j}" max="{j}" width="{w}" customWidth="1"/>' for j, w in enumerate(widths, 1))
    data = "".join(f'<row r="{k}">{"".join(rows[k])}</row>' for k in sorted(rows))
    head = old[:old.index("<sheetPr")]
    view = re.search(r"<sheetViews>.*?</sheetViews>", old, re.S).group(0)
    tail = old[old.index("</sheetData>") + len("</sheetData>"):]
    xml = (head + '<sheetPr filterMode="false"><pageSetUpPr fitToPage="false"/></sheetPr>'
           + f'<dimension ref="A1:{get_column_letter(max(maxc, len(widths)))}{last}"/>' + view
           + '<sheetFormatPr defaultRowHeight="15"/>' + f"<cols>{cols}</cols><sheetData>{data}</sheetData>" + tail)
    return xml, last, nform


zout = zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED)
changed = []
cn_info = ""
for info in zin.infolist():
    data = zin.read(info.filename)
    new = None
    if info.filename in by_part:
        new = patch_sheet(data.decode("utf-8"), by_part[info.filename]).encode("utf-8")
    elif info.filename == CN_PART and NOTES:
        x, last, nform = notes_xml(data.decode("utf-8"), json.load(open(NOTES)))
        new = x.encode("utf-8"); cn_info = f"; Collation Notes regenerated ({last} rows, {nform} formulas)"
    if new is None:
        zout.writestr(info, data)                              # byte-identical, same order
    else:
        zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
        zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = info.external_attr
        zout.writestr(zi, new); changed.append(info.filename)
zout.close()
assert stats["input"] == 1 and stats["values"] == len(P["values"]), stats
print(f"built {OUT}: input {s_in}!{c_in} = {v_in}; {stats['values']} cached values patched; parts changed {changed}{cn_info}")
