"""v5 checks. Reads only.

Usage:
  python3 -I analyze_v5.py V4_XLSX V5_XLSX V5_LO_XLSX DIFF_JSON DEPS_JSON NOISE_JSON PATCH_JSON FIG_STAGE2_JSON
                           FIG_V5_JSON NOTES_JSON NOTES_MD FIG_V4_JSON OUT_JSON

- Input diff: every cell whose value differs v4 -> v5 (all sheets but Collation Notes) is classified: input (no formula
  in v4 or v5), downstream (in the static trace from Fld Eng Budget!B21), or outside. Pass: input = [B21] only,
  outside = 0, 0 formula diffs, 0 comment diffs (diff_versions.py totals must agree).
- Package: only Collation Notes and the sheets holding patched cells change; no part added or removed.
- Cached values: LibreOffice recalculation of v5 vs v5's stored values, every sheet but Collation Notes, |diff| <= 1e-6.
- COO P&L = FO + PO + CO + FE + PE on every numeric non-ratio cell, rows 12+, columns B:N and U:AG.
- Collation Notes: formulas vs LibreOffice; required text (I-26 RESOLVED, N1, N2 label, N3 case).
- figures_v5.py on the final file = on stage 2. New errors vs v4. Stale v4 numbers in the notes.
Exit code 1 if a completion criterion fails.
"""
import sys, os, json, re, zipfile
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

(V4, V5, V5LO, DIFF, DEPS, NOISE, PATCH, FIGS2, FIG5, NOTES, MD, FIG4, OUT) = sys.argv[1:14]
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from stale_v5 import scan
ERR = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#ERROR!"}
TOL = 1e-6
res = {"fail": []}
isnum = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
P = json.load(open(PATCH)); deps = json.load(open(DEPS)); noise = json.load(open(NOISE))
down = {tuple(x) for x in deps["cells"]}
SEED = tuple(P["input"][:2])

f4 = openpyxl.load_workbook(V4); f5 = openpyxl.load_workbook(V5)
v4 = openpyxl.load_workbook(V4, data_only=True); v5 = openpyxl.load_workbook(V5, data_only=True)
lo = openpyxl.load_workbook(V5LO, data_only=True)


def isf(x):
    return isinstance(x, ArrayFormula) or (isinstance(x, str) and x.startswith("="))


# ------------------------------------------------------------------ diff classification
d = json.load(open(DIFF))
inputs, dn, outside = [], [], []
fdiff = []        # formula text changes (a cell that is a formula in v4 or v5 and whose formula differs)
by_sheet = {}
for s in v4.sheetnames:
    if s == "Collation Notes": continue
    a, b = v4[s], v5[s]
    for r in range(1, max(a.max_row, b.max_row) + 1):
        for c in range(1, max(a.max_column, b.max_column) + 1):
            x, y = a.cell(r, c).value, b.cell(r, c).value
            coord = a.cell(r, c).coordinate
            fx, fy = f4[s][coord].value, f5[s][coord].value
            nx = (fx.ref, fx.text) if isinstance(fx, ArrayFormula) else fx
            ny = (fy.ref, fy.text) if isinstance(fy, ArrayFormula) else fy
            if (isf(fx) or isf(fy)) and nx != ny: fdiff.append(f"{s}!{coord}")
            if x == y: continue
            if not isf(fx) and not isf(fy):
                inputs.append(f"{s}!{coord}")
            elif (s, coord) in down:
                dn.append((s, coord)); by_sheet[s] = by_sheet.get(s, 0) + 1
            else:
                outside.append(f"{s}!{coord}")
res["diff"] = {"total": d["total"], "sheets": {s: {k: x[k] for k in ("cells", "formula_diffs", "value_diffs")} for s, x in d["sheets"].items()},
               "comments": f"{len(d['comments_added'])} / {len(d['comments_removed'])} / {len(d['comments_changed'])}",
               "only": f"{d['sheets_only_in_old'] or 'none'} / {d['sheets_only_in_new'] or 'none'}",
               "class": {"input_cells": len(inputs), "input_list": inputs, "downstream_value_cells": len(dn), "by_sheet": by_sheet,
                         "outside": len(outside), "outside_list": outside[:20], "formula_text_diffs": len(fdiff),
                         "dv_formula_diffs_are_inputs": d["total"]["formula_diffs"] == len(inputs)}}
if inputs != [f"{SEED[0]}!{SEED[1]}"]: res["fail"].append(f"input diffs {inputs}")
if v5[SEED[0]][SEED[1]].value != P["input"][2]: res["fail"].append("B21 value")
if outside: res["fail"].append("value changes outside the downstream set")
# diff_versions.py counts a typed input as a 'formula diff' (it compares cell contents); real formula changes: fdiff
if fdiff or d["total"]["formula_diffs"] != len(inputs): res["fail"].append(f"formula diffs {fdiff[:5]}")
if d["comments_added"] or d["comments_removed"] or d["comments_changed"]: res["fail"].append("comment diffs")
if d["total"]["value_diffs"] != len(inputs) + len(dn) + len(outside): res["fail"].append("diff_versions total disagrees")
if len(dn) != len(P["values"]): res["fail"].append("downstream diffs != patched cells")
if d["sheets_only_in_old"] or d["sheets_only_in_new"]: res["fail"].append("sheet list")
res["noise"] = noise
res["deps"] = deps["info"]

# ------------------------------------------------------------------ package
z4, z5 = zipfile.ZipFile(V4), zipfile.ZipFile(V5)
n4, n5 = z4.namelist(), z5.namelist()
changed = [n for n in n4 if n in n5 and z4.read(n) != z5.read(n)]
wbx = z5.read("xl/workbook.xml").decode(); rel = z5.read("xl/_rels/workbook.xml.rels").decode()


def part(sheet):
    nm = sheet.replace("&", "&amp;")
    rid = re.search(rf'<sheet name="{re.escape(nm)}"[^>]*r:id="(rId\d+)"', wbx).group(1)
    return "xl/" + re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rel).group(1)


allowed = {part("Collation Notes")} | {part(s) for s, _, _ in P["values"]} | {part(SEED[0])}
res["package"] = {"parts_v4": len(n4), "parts_v5": len(n5), "identical_parts": len(n4) - len(changed), "changed_parts": changed,
                  "allowed_parts": sorted(allowed),
                  "added_parts": [n for n in n5 if n not in n4], "removed_parts": [n for n in n4 if n not in n5], "order_same": n4 == n5,
                  "externalLink_parts": sum(1 for n in n5 if n.startswith("xl/externalLinks/")),
                  "defined_names_v4": len(re.findall(r"<definedName[ >]", z4.read("xl/workbook.xml").decode())),
                  "defined_names_v5": len(re.findall(r"<definedName[ >]", wbx))}
if set(changed) - allowed or res["package"]["added_parts"] or res["package"]["removed_parts"] or not res["package"]["order_same"]:
    res["fail"].append("package parts")
res["sheets"] = [(s, f5[s].sheet_state) for s in f5.sheetnames]
if f5.sheetnames != f4.sheetnames: res["fail"].append("sheet list/order")

# ------------------------------------------------------------------ LibreOffice recalculation of v5 vs stored values
cnt = 0; mx = 0.0; tm = []
for s in v5.sheetnames:
    if s == "Collation Notes": continue
    a, b = v5[s], lo[s]
    for r in range(1, max(a.max_row, b.max_row) + 1):
        for c in range(1, max(a.max_column, b.max_column) + 1):
            x, y = a.cell(r, c).value, b.cell(r, c).value
            if x == y: continue
            if isnum(x) and isnum(y):
                cnt += 1; mx = max(mx, abs(x - y))
            else:
                tm.append([s, a.cell(r, c).coordinate, str(x)[:40], str(y)[:40]])
res["lo_check"] = {"cells_diff": cnt, "max_abs": mx, "text_mismatch": len(tm), "examples": tm[:10]}
if mx > TOL or tm: res["fail"].append("stored values differ from LibreOffice recalculation")

# ------------------------------------------------------------------ COO P&L = sum of departments
D = ("FO", "PO", "CO", "FE", "PE")
c_ = v5["COO P&L"]; cf = f5["COO P&L"]; n = 0; mxs = 0.0; bad = []
for r in range(12, c_.max_row + 1):
    for col in list(range(2, 15)) + list(range(21, 34)):           # B:N and U:AG (O and AH are ratio columns)
        x = c_.cell(r, col).value
        fx = cf.cell(r, col).value
        if not isnum(x) or (isinstance(fx, str) and "/" in fx): continue
        ys = [v5[s].cell(r, col).value for s in D]
        if any(y is not None and not isnum(y) for y in ys): continue
        n += 1; e = abs(x - sum(y or 0 for y in ys)); mxs = max(mxs, e)
        if e > TOL: bad.append([r, col, x])
res["coo_sum"] = {"cells": n, "max_abs": mxs, "fail": len(bad), "examples": bad[:5]}
if bad: res["fail"].append("COO P&L != sum of departments")

# ------------------------------------------------------------------ Collation Notes
cn = f5["Collation Notes"]; cnt = 0; mx = 0.0; tm = []
for row in cn.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            cnt += 1
            a, b = v5["Collation Notes"][c.coordinate].value, lo["Collation Notes"][c.coordinate].value
            if isnum(a) and isnum(b): mx = max(mx, abs(a - b))
            elif a != b: tm.append([c.coordinate, a, b])
txt = " ".join(str(c.value) for row in cn.iter_rows() for c in row if isinstance(c.value, str))
F5 = json.load(open(FIG5))
need = ["Decisions pending", "COO contribution", "Caveats on the headline", "Sign-flip statement", "not company Net Income",
        "I-17", "I-05", "I-10", "I-26", "I-31", "I-32", "n/m", "Changes from v4"]
missing = [x for x in need if x not in txt]
i26_rows = [(row[1].value, str(row[3].value)) for row in cn.iter_rows(min_col=1, max_col=4) if row[0].value == "I-26"]
# caveat table row (direction RESOLVED) and issues table row (Sev RESOLVED), both with a 'RESOLVED (user-confirmed' text
resolved_ok = (sum(1 for b, _ in i26_rows if b == "RESOLVED") == 2 and not any(b in ("HIGH", "MED", "LOW", "downside", "upside") for b, _ in i26_rows)
               and any(t.startswith("RESOLVED (user-confirmed") for _, t in i26_rows))
n1_ok = all(x in txt for x in ("Depreciation Schedule'!B17", "D21", "row 19 heading", "A-17, A-18, A-19 and I-31"))
n2_ok = "illustrative extremes of quantified items; excludes I-02, A-22; lower end assumes no I-05/I-17 overlap" in txt
n3_ok = f"{F5['i32_trend_sep']:,.0f}" in txt and "continued from Sep-26" in txt
res["cn"] = {"formulas": cnt, "max_num_diff": mx, "text_mismatch": len(tm), "mismatches": tm, "missing_text": missing,
             "resolved_ok": resolved_ok, "n1_ok": n1_ok, "n2_ok": n2_ok, "n3_ok": n3_ok}
if mx > TOL or tm or cnt == 0: res["fail"].append("Collation Notes formula values")
if missing or not (resolved_ok and n1_ok and n2_ok and n3_ok): res["fail"].append(f"Collation Notes text {missing} {resolved_ok} {n1_ok} {n2_ok} {n3_ok}")

# ------------------------------------------------------------------ figures, errors
res["figures_identical"] = json.load(open(FIGS2)) == F5
if not res["figures_identical"]: res["fail"].append("figures differ stage 2 vs final")
errs = {}; new_errors = []
for s in v4.sheetnames:
    if s == "Collation Notes": continue
    e4, e5 = {}, {}
    for row in v4[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value in ERR: e4[c.value] = e4.get(c.value, 0) + 1
    for row in v5[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value in ERR:
                e5[c.value] = e5.get(c.value, 0) + 1
                x = v4[s][c.coordinate].value
                if not (isinstance(x, str) and x in ERR): new_errors.append(f"{s}!{c.coordinate}")
    errs[s] = {"v4": e4, "v5": e5}
res["errors"] = errs; res["new_errors"] = new_errors
if new_errors: res["fail"].append("new errors")

# ------------------------------------------------------------------ stale v4 numbers (pass-1 md + Collation Notes)
notes = json.load(open(NOTES))


def flat(o):
    if isinstance(o, dict): return " ".join(flat(v) for v in o.values())
    if isinstance(o, list): return " ".join(flat(v) for v in o)
    return str(o)


cn_secs = [x for x in notes["sections"] if not x["heading"].startswith(("Changes from v4", "Head of Service Delivery salary"))]
st = scan(open(MD).read(), flat(cn_secs), V4, P, json.load(open(FIG4)), F5, V5)
res["stale_numbers"] = st
if st["hits"]: res["fail"].append("stale v4 numbers in notes")
res["stale_checked"] = notes.get("stale_checked", [])

json.dump(res, open(OUT, "w"), indent=1, default=str)
k = res["diff"]["class"]
print("analyze_v5: input", k["input_list"], "| downstream value diffs", k["downstream_value_cells"], k["by_sheet"], "| outside", k["outside"],
      "| formula text diffs", len(fdiff), "| parts changed", changed, "| LO check max", f"{res['lo_check']['max_abs']:.1e}",
      "| COO sum max", f"{mxs:.1e}", "| CN formulas", res["cn"]["formulas"], "| new errors", len(new_errors),
      "| stale", st["hits"], "/", st["checked"], "| FAIL" if res["fail"] else "| all checks pass", res["fail"])
sys.exit(1 if res["fail"] else 0)
