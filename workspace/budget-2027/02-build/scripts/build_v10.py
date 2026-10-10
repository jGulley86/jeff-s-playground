"""v10 build: Logistics & Warehouse payroll booked in Central Operations, the material handler moved PO -> CO, and the new CO
employee (brief v6). Package patch of v9.

Usage:
  python3 -I build_v10.py stage1 V9_XLSX SD_BOOK NAME_FILE OUT_STAGE1_XLSX OUT_SPEC_JSON
      v9 package with the v10 edits only (cached values of edited / new formula cells are placeholders until patched):
        1. New sheet 'Logistics&WH Payroll', appended after HC-COO, copied cell by cell from the SD book (values, formulas incl.
           shared formulas, cached values, styles, column widths, row heights, merged cells, cell notes). Shared strings are
           inlined (sharedStrings.xml untouched); the SD cell formats it uses are APPENDED to styles.xml (existing entries keep
           their index, so no existing cell changes format); theme colours resolved to RGB from the SD theme. Its formulas read
           DeploySum, Fld Ops Payroll and Prod Payroll, which are all in the workbook (checked). The SD sheet's drawing part is
           an empty Google Sheets drawing (no shapes) and is not copied.
        2. HC-CO row 18 (the empty current-employee row): B18 name (read from NAME_FILE, outside the repo), C18 title
           placeholder, D18 110,000, E18 Jan-27. F18 (Dec-27), G18 (3%) and H18 (= D8, Oct-27) are the row's existing defaults.
        3. CO U67:AF69 (6610 / 6630 / 6640, Jan-Dec 27): typed values -> formulas = HC-CO rows 67-69 (the roster total, now
           incl. row 18; equal to the v9 typed values for the 5 existing people) + the Logistics&WH model less seat LOG-01
           (row 34, burdened at the model's rates B5 / B6); new-hire cost (row 46) -> 6610. CO AI67:AI69 notes.
        4. PO U33:AF35 (5110-5140, the B131 base): v7 amount less PO HC r20 (Material Handler = seat MH-01) at PO HC rates.
      OUT_SPEC_JSON lists every edited / new cell with the old and new content.
  python3 -I build_v10.py final STAGE1_XLSX PATCH_JSON OUT_XLSX [NOTES_JSON]
      = build_v8.final (cached values from patches_v10.py; optional Collation Notes from NOTES_JSON).
Every other zip part is copied byte for byte.
"""
import sys, re, json, zipfile, os, colorsys, datetime
from xml.sax.saxutils import escape, unescape
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_v8
from build_v8 import part_of, cell_xml

LW = "Logistics&WH Payroll"
LWQ = "'Logistics&WH Payroll'"
Y27_LW = [get_column_letter(c) for c in range(6, 18)]      # F..Q = Jan..Dec 27 on the SD models
P27 = [get_column_letter(c) for c in range(21, 33)]        # U..AF on the P&L tabs
H27 = [get_column_letter(c) for c in range(13, 25)]        # M..X on the HC tabs
NEW_ROW = 18
NEW_SALARY = 110000
NEW_START = datetime.date(2027, 1, 1)
NEW_TITLE = "(title not provided; new CO employee, user-confirmed 2026-10-10)"
CELL = lambda coord: re.compile(rf'<c r="{coord}"(?: [^>]*?)?(?:/>|>.*?</c>)', re.S)
UNESC = {"&apos;": "'", "&quot;": '"'}


def serial(d):
    return (d - datetime.date(1899, 12, 30)).days


def replace_cell(x, coord, new_xml):
    m = list(CELL(coord).finditer(x))
    assert len(m) == 1, (coord, len(m))
    return x[:m[0].start()] + new_xml + x[m[0].end():], m[0].group(0)


def old_content(cxml):
    f = re.search(r"<f[^>]*>([^<]*)</f>", cxml)
    if f: return "=" + unescape(f.group(1), UNESC)
    t = re.search(r"<t[^>]*>([^<]*)</t>", cxml)
    if t: return unescape(t.group(1), UNESC)
    v = re.search(r"<v>([^<]*)</v>", cxml)
    return v.group(1) if v else ""


def style_of(cxml):
    return int(re.search(r' s="(\d+)"', cxml).group(1))


def fcell(coord, f, style):
    return f'<c r="{coord}" s="{style}" t="n"><f aca="false">{escape(f, {chr(39): "&apos;"})}</f><v>0</v></c>'


def part_map(z):
    wb = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    out = {}
    for m in re.finditer(r'<sheet ([^>]*)/>', wb):
        a = m.group(1)
        nm = unescape(re.search(r'name="([^"]+)"', a).group(1), UNESC)
        rid = re.search(r'r:id="(rId\d+)"', a).group(1)
        tgt = re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rel_norm(rels)).group(1)
        out[nm] = tgt if tgt.startswith("xl/") else "xl/" + tgt.lstrip("/")
    return out


def rel_norm(rels):
    """Relationship elements with Id first (the SD book writes Type/Target before Id)."""
    def fix(m):
        a = m.group(1)
        i = re.search(r'Id="[^"]+"', a).group(0)
        return "<Relationship " + i + " " + re.sub(r'\s*Id="[^"]+"', "", a).strip() + "/>"
    return re.sub(r"<Relationship ([^>]*?)/>", fix, rels)


# ------------------------------------------------------------------ 1. copy the Logistics&WH sheet (build_v6 method)
STD_PALETTE = ("000000 FFFFFF FF0000 00FF00 0000FF FFFF00 FF00FF 00FFFF 000000 FFFFFF FF0000 00FF00 0000FF FFFF00 FF00FF "
               "00FFFF 800000 008000 000080 808000 800080 008080 C0C0C0 808080 9999FF 993366 FFFFCC CCFFFF 660066 FF8080 "
               "0066CC CCCCFF 000080 FF00FF FFFF00 00FFFF 800080 800000 008080 0000FF 00CCFF CCFFFF CCFFCC FFFF99 99CCFF "
               "FF99CC CC99FF FFCC99 3366FF 33CCCC 99CC00 FFCC00 FF9900 FF6600 666699 969696 003366 339966 003300 333300 "
               "993300 993366 333399 333333").split()
THEME_IDX = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]


def items(xml, tag, child):
    blk = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", xml, re.S)
    return re.findall(rf"<{child}(?:\s[^>]*)?/>|<{child}(?:\s[^>]*)?>.*?</{child}>", blk.group(1), re.S) if blk else []


def append_block(xml, tag, new_items):
    if not new_items:
        return xml
    m = re.search(rf'<{tag} count="(\d+)"', xml)
    n = int(m.group(1))
    xml = xml[:m.start(1)] + str(n + len(new_items)) + xml[m.end(1):]
    i = xml.index(f"</{tag}>")
    return xml[:i] + "".join(new_items) + xml[i:]


def copy_sheet(zv, zs, vst):
    """Returns (sheet_xml, comments_xml, vml_bytes, new_styles_xml, stats)."""
    sparts = part_map(zs)
    sst_x = zs.read("xl/styles.xml").decode()
    assert "<indexedColors>" not in sst_x
    theme = zs.read("xl/theme/theme1.xml").decode()
    cs = re.search(r"<a:clrScheme.*?</a:clrScheme>", theme, re.S).group(0)
    scheme = {}
    for k in THEME_IDX:
        blk = re.search(rf"<a:{k}>(.*?)</a:{k}>", cs, re.S).group(1)
        m = re.search(r'srgbClr val="([0-9A-Fa-f]{6})"', blk) or re.search(r'lastClr="([0-9A-Fa-f]{6})"', blk)
        scheme[k] = m.group(1).upper()

    def tint_rgb(hexrgb, tint):
        r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
        r, g, b = colorsys.hls_to_rgb(h, l, s)
        return "".join(f"{round(x * 255):02X}" for x in (r, g, b))

    cstat = {"theme": 0, "indexed": 0}

    def fix_colours(x):
        def rep(m):
            tag, attrs = m.group(1), m.group(2)
            th = re.search(r'theme="(\d+)"', attrs); ix = re.search(r'indexed="(\d+)"', attrs)
            tn = re.search(r'tint="([-0-9.Ee]+)"', attrs)
            if th:
                rgb = scheme[THEME_IDX[int(th.group(1))]]
                if tn:
                    rgb = tint_rgb(rgb, float(tn.group(1)))
                cstat["theme"] += 1
                return f'<{tag} rgb="FF{rgb}"/>'
            if ix and int(ix.group(1)) < 64:
                cstat["indexed"] += 1
                return f'<{tag} rgb="FF{STD_PALETTE[int(ix.group(1))]}"/>'
            return m.group(0)
        return re.sub(r"<(color|fgColor|bgColor)(\s[^>]*?)/>", rep, x)

    w_fonts, w_fills, w_borders = items(sst_x, "fonts", "font"), items(sst_x, "fills", "fill"), items(sst_x, "borders", "border")
    w_xfs = items(sst_x, "cellXfs", "xf")
    w_nf = {int(a): b for a, b in re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"/>', sst_x)}
    v_fonts, v_fills, v_borders = items(vst, "fonts", "font"), items(vst, "fills", "fill"), items(vst, "borders", "border")
    v_xfs = items(vst, "cellXfs", "xf")
    v_nf = {int(a): b for a, b in re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"/>', vst)}
    assert len(v_xfs) == int(re.search(r'<cellXfs count="(\d+)"', vst).group(1))
    new_fonts, new_fills, new_borders, new_nf, new_xfs = [], [], [], {}, []
    fmap, flmap, bmap, nmap, xmap = {}, {}, {}, {}, {}

    def map_xf(i):
        if i in xmap:
            return xmap[i]
        xf = w_xfs[i]
        at = dict(re.findall(r'(\w+)="([^"]*)"', re.match(r"<xf([^>]*?)/?>", xf).group(1)))
        nf = int(at.get("numFmtId", 0))
        if nf >= 164:
            if nf not in nmap:
                code = w_nf[nf]
                same = [k for k, v in v_nf.items() if v == code]
                nmap[nf] = same[0] if same else max(list(v_nf) + list(new_nf) + [163]) + 1
                if not same:
                    new_nf[nmap[nf]] = code
            nf = nmap[nf]
        fo = int(at.get("fontId", 0)); fi = int(at.get("fillId", 0)); bo = int(at.get("borderId", 0))
        if fo not in fmap:
            fmap[fo] = len(v_fonts) + len(new_fonts); new_fonts.append(fix_colours(w_fonts[fo]))
        if fi not in flmap:
            flmap[fi] = len(v_fills) + len(new_fills); new_fills.append(fix_colours(w_fills[fi]))
        if bo not in bmap:
            bmap[bo] = len(v_borders) + len(new_borders); new_borders.append(fix_colours(w_borders[bo]))
        body = re.search(r"<xf[^>]*?>(.*)</xf>", xf, re.S)
        body = body.group(1) if body else ""
        attrs = (f'numFmtId="{nf}" fontId="{fmap[fo]}" fillId="{flmap[fi]}" borderId="{bmap[bo]}" xfId="0" '
                 f'applyNumberFormat="1" applyFont="1" applyFill="1" applyBorder="1"')
        if "<alignment" in body:
            attrs += ' applyAlignment="1"'
        if "<protection" in body:
            attrs += ' applyProtection="1"'
        xmap[i] = len(v_xfs) + len(new_xfs); new_xfs.append(f"<xf {attrs}>{body}</xf>" if body else f"<xf {attrs}/>")
        return xmap[i]

    SI = re.findall(r"<si>(.*?)</si>", zs.read("xl/sharedStrings.xml").decode(), re.S)
    vsheets = set(part_map(zv))
    src = sparts[LW]
    xml = zs.read(src).decode()
    refs = set()

    def cell(m):
        attrs, inner = m.group(1), m.group(2)
        attrs = re.sub(r' s="(\d+)"', lambda mm: f' s="{map_xf(int(mm.group(1)))}"', attrs)
        if inner is None:
            return f"<c{attrs}/>"
        if ' t="s"' in attrs:
            idx = int(re.search(r"<v>(\d+)</v>", inner).group(1))
            return f'<c{attrs.replace(" t=\"s\"", " t=\"inlineStr\"")}><is>{SI[idx]}</is></c>'
        fm = re.search(r"<f[^>]*>(.*?)</f>", inner, re.S)
        if fm:
            body = unescape(fm.group(1), UNESC)
            assert not re.search(r"\[\d+\]", body), "external reference"
            for s in re.findall(r"'([^']+)'!", body) + re.findall(r"(?<!['\w])([A-Za-z][\w ]*)!", body):
                refs.add(s)
        return f"<c{attrs}>{inner}</c>"

    xml = re.sub(r"<c(\s[^>]*?)(?:/>|>(.*?)</c>)", cell, xml, flags=re.S)
    xml = re.sub(r'(<row\s[^>]*?) s="(\d+)"', lambda m: f'{m.group(1)} s="{map_xf(int(m.group(2)))}"', xml)
    xml = re.sub(r'(<col\s[^>]*?) style="(\d+)"', lambda m: f'{m.group(1)} style="{map_xf(int(m.group(2)))}"', xml)
    xml = xml.replace(' tabSelected="1"', "")
    for s in refs:
        assert s in vsheets, f"Logistics&WH formula reads {s!r}, not in the workbook"
    # relationships: keep comments + VML (cell notes), drop the empty drawing
    srel = rel_norm(zs.read(src.replace("worksheets/", "worksheets/_rels/") + ".rels").decode())
    rels = re.findall(r'<Relationship Id="(rId\d+)" Type="([^"]+)" Target="([^"]+)"/>', srel)
    comments = vml = None; drop = []
    for rid, typ, tgt in rels:
        data = zs.read("xl/" + tgt.replace("../", ""))
        if typ.endswith("/comments"):
            comments = data.decode()
        elif typ.endswith("/vmlDrawing"):
            vml = data
        elif typ.endswith("/drawing"):
            assert "<xdr:twoCellAnchor" not in data.decode() and "<xdr:oneCellAnchor" not in data.decode() \
                and "<xdr:absoluteAnchor" not in data.decode(), "drawing is not empty"
            drop.append(rid)
        else:
            raise SystemExit(f"unexpected relationship {typ}")
    for rid in drop:
        xml, n = re.subn(rf'<drawing r:id="{rid}"/>', "", xml); assert n == 1
    st = vst
    st = append_block(st, "numFmts", [f'<numFmt numFmtId="{k}" formatCode="{v}"/>' for k, v in sorted(new_nf.items())])
    st = append_block(st, "fonts", new_fonts)
    st = append_block(st, "fills", new_fills)
    st = append_block(st, "borders", new_borders)
    st = append_block(st, "cellXfs", new_xfs)
    for tag, child, old in (("fonts", "font", v_fonts), ("fills", "fill", v_fills), ("borders", "border", v_borders),
                            ("cellXfs", "xf", v_xfs)):
        assert items(st, tag, child)[:len(old)] == old, f"existing {tag} changed"
    stats = {"from": src, "cells": len(re.findall(r"<c\s", xml)), "formulas": xml.count("<f"), "reads": sorted(refs),
             "styles_appended": {"cellXfs": len(new_xfs), "fonts": len(new_fonts), "fills": len(new_fills),
                                 "borders": len(new_borders), "numFmts": len(new_nf)},
             "colours_resolved": cstat, "dropped": "empty drawing (Google Sheets, no shapes)" if drop else ""}
    return xml, comments, vml, st, stats


# ------------------------------------------------------------------ stage 1
def stage1(v9, sd, name_file, out, spec_out):
    zv = zipfile.ZipFile(v9); zs = zipfile.ZipFile(sd)
    name = open(name_file, encoding="utf-8").read().strip()
    assert name and "\n" not in name
    vparts = part_map(zv)
    assert LW not in vparts
    spec = {"lw": None, "hcco": [], "co": [], "po": [], "parts": {}}
    vst = zv.read("xl/styles.xml").decode()
    lw_xml, lw_comments, lw_vml, new_styles, stats = copy_sheet(zv, zs, vst)
    spec["lw"] = stats

    # 2. HC-CO row 18
    p_hc = part_of(zv, "HC-CO"); x = zv.read(p_hc).decode("utf-8")
    for coord in ("B18", "C18", "D18"):
        assert re.search(rf'<c r="{coord}" s="\d+"/>', x), f"HC-CO {coord} must be empty"
    m = CELL("E18").search(x); assert old_content(m.group(0)) == "46266"
    for coord, kind, val in (("B18", "text", name), ("C18", "text", NEW_TITLE), ("D18", "num", NEW_SALARY),
                             ("E18", "num", serial(NEW_START))):
        m = CELL(coord).search(x)
        x, old = replace_cell(x, coord, cell_xml(coord, kind, val, style_of(m.group(0))))
        spec["hcco"].append({"cell": coord, "kind": kind, "old": old_content(old),
                             "new": "<name: NAME_FILE>" if coord == "B18" else val})
    new_hc = x

    # 3. CO U67:AF69
    p_co = part_of(zv, "CO"); x = zv.read(p_co).decode("utf-8")
    for k, col in enumerate(P27):
        lc, hc = Y27_LW[k], H27[k]
        fs = {67: f"'HC-CO'!{hc}67+{LWQ}!{lc}43-{LWQ}!{lc}34+{LWQ}!{lc}46",
              68: f"'HC-CO'!{hc}68+{LWQ}!{lc}44-{LWQ}!{lc}34*{LWQ}!$B$5",
              69: f"'HC-CO'!{hc}69+{LWQ}!{lc}45-{LWQ}!{lc}34*{LWQ}!$B$6"}
        for row, f in fs.items():
            coord = f"{col}{row}"
            m = CELL(coord).search(x); oc = old_content(m.group(0))
            assert re.fullmatch(r"[-0-9.E]+", oc), (coord, oc)
            x, _ = replace_cell(x, coord, fcell(coord, f, style_of(m.group(0))))
            spec["co"].append({"cell": coord, "old": oc, "new": "=" + f})
    notes = {
        67: ("v10: 6610 2027 = 'HC-CO' row 67 (CO HC roster incl. the new CO employee, row 18) + 'Logistics&WH Payroll' salaries less seat "
             "LOG-01 (row 43 - row 34; LOG-01 is CO HC r17, already in the roster) + the model's new-hire cost (row 46, mapped to 6610 as "
             "in FO 5510). Material handler MH-01 (PO HC r20) is here; PO 5110-5140 base excludes it. User decision 2026-10-10 (brief v6)."),
        68: ("v10: 6630 2027 = 'HC-CO' row 68 (CO HC rate D6) + 'Logistics&WH Payroll' payroll taxes less LOG-01 (row 44 - row 34 x B5; "
             "model rate 6.8%, not the CO HC rate)."),
        69: ("v10: 6640 2027 = 'HC-CO' row 69 (CO HC rate D7) + 'Logistics&WH Payroll' benefits less LOG-01 (row 45 - row 34 x B6; "
             "model rate 16.4%, not the CO HC rate)."),
    }
    for row, txt in notes.items():
        coord = f"AI{row}"
        m = CELL(coord).search(x); assert m and old_content(m.group(0)) == "", coord
        x, _ = replace_cell(x, coord, cell_xml(coord, "text", txt, style_of(m.group(0))))
        spec["co"].append({"cell": coord, "old": "", "new": txt})
    new_co = x

    # 4. PO U33:AF35: the B131 base less PO HC r20 (Material Handler = MH-01), at PO HC rates
    p_po = part_of(zv, "PO"); x = zv.read(p_po).decode("utf-8")
    sub = {33: "'HC-PO'!{h}20", 34: "'HC-PO'!{h}20*'HC-PO'!$D$6", 35: "'HC-PO'!{h}20*'HC-PO'!$D$7"}
    for row in (33, 34, 35):
        for k, col in enumerate(P27):
            coord = f"{col}{row}"
            m = CELL(coord).search(x); oc = old_content(m.group(0))
            mm = re.fullmatch(r"=\(1-'Depreciation Schedule'!\$B\$131\)\*([-0-9.E]+)\+('Depreciation Schedule'!\w+)", oc)
            assert mm, (coord, oc)
            nf = f"(1-'Depreciation Schedule'!$B$131)*({mm.group(1)}-{sub[row].format(h=H27[k])})+{mm.group(2)}"
            vv = re.search(r"<v>([^<]*)</v>", m.group(0)).group(1)
            new = f'<c r="{coord}" s="{style_of(m.group(0))}" t="n"><f aca="false">{escape(nf, {chr(39): "&apos;"})}</f><v>{vv}</v></c>'
            x, _ = replace_cell(x, coord, new)
            spec["po"].append({"cell": coord, "old": oc, "new": "=" + nf})
    new_po = x

    # workbook, rels, content types
    wbx = zv.read("xl/workbook.xml").decode(); rels = zv.read("xl/_rels/workbook.xml.rels").decode()
    ct = zv.read("[Content_Types].xml").decode()
    assert re.findall(r'<sheet name="([^"]+)"', wbx)[-1] == "HC-COO"
    num = max(int(re.search(r"sheet(\d+)\.xml", p).group(1)) for p in vparts.values()) + 1
    cnum = max(int(n) for n in re.findall(r"xl/comments(\d+)\.xml", " ".join(zv.namelist()))) + 1
    vnum = max(int(n) for n in re.findall(r"xl/drawings/vmlDrawing(\d+)\.vml", " ".join(zv.namelist()))) + 1
    sid = max(int(v) for v in re.findall(r'<sheet name="[^"]+" sheetId="(\d+)"', wbx)) + 1
    rid = max(int(v) for v in re.findall(r'Id="rId(\d+)"', rels)) + 1
    sheet_p = f"xl/worksheets/sheet{num}.xml"; rel_p = f"xl/worksheets/_rels/sheet{num}.xml.rels"
    com_p = f"xl/comments{cnum}.xml"; vml_p = f"xl/drawings/vmlDrawing{vnum}.vml"
    srel = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    lw_rels = re.findall(r'<legacyDrawing r:id="(rId\d+)"/>', lw_xml)
    assert len(lw_rels) == 1
    # renumber: comments rId1, vml = the legacyDrawing id
    srel += (f'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="../comments{cnum}.xml"/>'
             f'<Relationship Id="{lw_rels[0]}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" '
             f'Target="../drawings/vmlDrawing{vnum}.vml"/></Relationships>')
    assert lw_rels[0] != "rId1"
    add_sheet = f'<sheet name="{escape(LW)}" sheetId="{sid}" state="visible" r:id="rId{rid}"/>'
    wbx = wbx.replace("</sheets>", add_sheet + "</sheets>", 1)
    rels = rels.replace("</Relationships>", f'<Relationship Id="rId{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                        f'relationships/worksheet" Target="worksheets/sheet{num}.xml"/></Relationships>', 1)
    ct = ct.replace("</Types>", f'<Override PartName="/{sheet_p}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                    f'<Override PartName="/{com_p}" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"/>'
                    f'<Override PartName="/{vml_p}" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
                    "</Types>", 1)
    new_parts = {sheet_p: lw_xml.encode(), rel_p: srel.encode(), com_p: lw_comments.encode(), vml_p: lw_vml}
    assert not (set(new_parts) & set(zv.namelist()))
    repl = {p_hc: new_hc, p_co: new_co, p_po: new_po, "xl/workbook.xml": wbx, "xl/_rels/workbook.xml.rels": rels,
            "[Content_Types].xml": ct, "xl/styles.xml": new_styles}
    zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in zv.infolist():
        data = zv.read(info.filename)
        if info.filename in repl:
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = info.external_attr
            zo.writestr(zi, repl[info.filename].encode("utf-8"))
        else:
            zo.writestr(info, data)
    for n in sorted(new_parts):
        zo.writestr(n, new_parts[n])
    zo.close()
    spec["parts"] = {"HC-CO": p_hc, "CO": p_co, "PO": p_po, LW: sheet_p, "added": sorted(new_parts), "changed": sorted(repl)}
    spec["lw_cells"] = sorted(set(re.findall(r'<c r="([A-Z]+\d+)"', lw_xml)), key=lambda c: (int(re.sub(r"\D", "", c)), c))
    spec["lw_formula_cells"] = [m.group(1) for m in re.finditer(r'<c r="([A-Z]+\d+)"([^>]*?)(/>|>(.*?)</c>)', lw_xml, re.S)
                                if m.group(4) and "<f" in m.group(4)]
    assert len(spec["lw_formula_cells"]) == stats["formulas"]
    json.dump(spec, open(spec_out, "w"), indent=1)
    print(f"stage1 {out}: new sheet {LW} ({stats['cells']} cells, {stats['formulas']} formulas, reads {stats['reads']}, styles "
          f"appended {stats['styles_appended']}); HC-CO cells {len(spec['hcco'])}; CO cells {len(spec['co'])}; PO cells {len(spec['po'])}")


if __name__ == "__main__":
    if sys.argv[1] == "stage1":
        stage1(*sys.argv[2:7])
    elif sys.argv[1] == "final":
        build_v8.final(*sys.argv[2:6])
    else:
        raise SystemExit(__doc__)
