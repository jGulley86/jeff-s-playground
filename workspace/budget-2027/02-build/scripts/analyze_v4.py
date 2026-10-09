"""v4 checks. Reads only.

Usage:
  python3 -I analyze_v4.py V3_XLSX V4_XLSX V4_LO_XLSX DIFF_JSON FIG_V3_JSON FIG_V4_JSON NOTES_JSON OUT_JSON

- DIFF_JSON: diff_versions.py v3 -> v4 (exact). Must be 0 formula / 0 value diffs.
- Package: every zip part byte-identical to v3 except the Collation Notes sheet part.
- figures_v4.py on v4 must equal figures_v4.py on v3.
- Collation Notes formulas: stored values vs LibreOffice's recalculation of v4 (V4_LO_XLSX).
- LibreOffice round trip (V4_LO_XLSX vs v3, other sheets): counted to document why v4 is a package patch.
- Error scan, defined names, externalLink parts, sheet list and states.
Exit code 1 if a completion criterion fails.
"""
import sys, json, re, zipfile, filecmp
import openpyxl

V3, V4, V4LO, DIFF, FIG3, FIG4, NOTES, OUT = sys.argv[1:9]
ERR = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#ERROR!"}
res = {"fail": []}

d = json.load(open(DIFF))
res["diff"] = {"total": d["total"], "sheets": {s: {k: x[k] for k in ("cells", "formula_diffs", "value_diffs")} for s, x in d["sheets"].items()},
               "comments": f"{len(d['comments_added'])} / {len(d['comments_removed'])} / {len(d['comments_changed'])}",
               "only": f"{d['sheets_only_in_old'] or 'none'} / {d['sheets_only_in_new'] or 'none'}"}
if d["total"]["formula_diffs"] or d["total"]["value_diffs"] or d["comments_added"] or d["comments_removed"] or d["comments_changed"]:
    res["fail"].append("diff vs v3 not 0")

# ------------------------------------------------------------------ package
z3, z4 = zipfile.ZipFile(V3), zipfile.ZipFile(V4)
n3, n4 = z3.namelist(), z4.namelist()
same = [n for n in n3 if n in n4 and z3.read(n) == z4.read(n)]
changed = [n for n in n3 if n in n4 and z3.read(n) != z4.read(n)]
wb4 = z4.read("xl/workbook.xml").decode(); wb3 = z3.read("xl/workbook.xml").decode()
res["package"] = {"parts_v3": len(n3), "parts_v4": len(n4), "identical_parts": len(same), "changed_parts": changed,
                  "added_parts": [n for n in n4 if n not in n3], "removed_parts": [n for n in n3 if n not in n4],
                  "order_same": n3 == n4,
                  "externalLink_parts": sum(1 for n in n4 if n.startswith("xl/externalLinks/")),
                  "defined_names_v3": len(re.findall(r"<definedName[ >]", wb3)),
                  "defined_names_v4": len(re.findall(r"<definedName[ >]", wb4))}
if changed != ["xl/worksheets/sheet1.xml"] or res["package"]["added_parts"] or res["package"]["removed_parts"]:
    res["fail"].append("package parts other than Collation Notes changed")

# ------------------------------------------------------------------ figures
res["figures_identical"] = json.load(open(FIG3)) == json.load(open(FIG4))
if not res["figures_identical"]: res["fail"].append("figures differ v3 vs v4")

# ------------------------------------------------------------------ workbooks
f4 = openpyxl.load_workbook(V4); v4 = openpyxl.load_workbook(V4, data_only=True)
v3 = openpyxl.load_workbook(V3, data_only=True); lo = openpyxl.load_workbook(V4LO, data_only=True)
res["sheets"] = [(s, f4[s].sheet_state) for s in f4.sheetnames]
if [s for s, _ in res["sheets"]] != v3.sheetnames: res["fail"].append("sheet list/order changed")

# Collation Notes formulas: stored vs LibreOffice
cn = f4["Collation Notes"]; cnt = 0; mx = 0.0; tm = []
for row in cn.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            cnt += 1
            a, b = v4["Collation Notes"][c.coordinate].value, lo["Collation Notes"][c.coordinate].value
            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                mx = max(mx, abs(a - b))
            elif a != b:
                tm.append([c.coordinate, a, b])
res["cn"] = {"formulas": cnt, "max_num_diff": mx, "text_mismatch": len(tm), "mismatches": tm}
if mx > 1e-6 or tm or cnt == 0: res["fail"].append("Collation Notes formula values")
txt = " ".join(str(c.value) for row in cn.iter_rows() for c in row if isinstance(c.value, str))
for need in ("Decisions pending", "COO contribution", "Caveats on the headline", "Sign-flip statement", "not company Net Income",
             "I-17", "I-05", "I-10", "I-26", "I-31", "I-32", "n/m"):
    if need not in txt: res["fail"].append(f"Collation Notes missing '{need}'")

# error scan
errs = {}; new_errors = []
for s in v3.sheetnames:
    if s == "Collation Notes": continue
    e3, e4 = {}, {}
    for row in v3[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value in ERR: e3[c.value] = e3.get(c.value, 0) + 1
    for row in v4[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value in ERR:
                e4[c.value] = e4.get(c.value, 0) + 1
                x = v3[s][c.coordinate].value
                if not (isinstance(x, str) and x in ERR): new_errors.append(f"{s}!{c.coordinate}")
    errs[s] = {"v3": e3, "v4": e4}
res["errors"] = errs; res["new_errors"] = new_errors
if new_errors: res["fail"].append("new errors")

# LibreOffice round trip vs v3 (documentation only)
n = 0; mxa = 0.0; sh = {}
for s in v3.sheetnames:
    if s == "Collation Notes": continue
    a_ws, b_ws = v3[s], lo[s]
    for row in a_ws.iter_rows():
        for c in row:
            b = b_ws[c.coordinate].value
            if c.value != b:
                n += 1; sh[s] = sh.get(s, 0) + 1
                if isinstance(c.value, (int, float)) and isinstance(b, (int, float)):
                    mxa = max(mxa, abs(c.value - b))
res["lo_roundtrip"] = {"value_diffs": n, "max_abs": mxa, "sheets": sh}
res["stale_checked"] = json.load(open(NOTES)).get("stale_checked", [])

json.dump(res, open(OUT, "w"), indent=1, default=str)
print("analyze_v4:", "diff", d["total"], "| parts changed", changed, "| figures identical", res["figures_identical"],
      "| CN formulas", cnt, "max diff", mx, "text mismatches", len(tm), "| new errors", len(new_errors),
      "| LO round-trip diffs", n, "| FAIL" if res["fail"] else "| all checks pass", res["fail"])
sys.exit(1 if res["fail"] else 0)
