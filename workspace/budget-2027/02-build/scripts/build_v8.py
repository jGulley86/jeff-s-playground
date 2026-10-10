"""v8 build: production labour capitalised (brief v5). Package patch of the delivered v7 workbook.

Usage:
  python3 -I build_v8.py stage1 V7_XLSX OUT_STAGE1_XLSX OUT_SPEC_JSON
      v7 package with the v8 formula edits only (cached values of edited cells are placeholders):
        1. Prod Payroll!B48:Q48: the 16 XLOOKUP cells -> an INDEX/MATCH equivalent that LibreOffice evaluates
           ('next tier up, else max'); no other Prod Payroll cell is edited.
        2. Depreciation Schedule: new sections 8-10 appended below row 126 (rows 128+): capitalised production labour
           by build month (Prod Payroll rows 41/62/79), labour per unit built, labour by go-live vintage, depreciation
           by GL, memo. Rows 1-126 are not touched.
        3. PO!U20:AF20 / U21:AF21 (5011 / 5012): + the capitalised-labour depreciation (Depreciation Schedule row
           193 / 194). PO!U33:AF35 (5110-5140): =(1-'Depreciation Schedule'!$B$131)*<v7 amount> = 0 (capitalised;
           the switch B131 = 0 gives the sensitivity with the current staff expensed). PO!AI20:AI21, AI33:AI35: notes.
      OUT_SPEC_JSON lists every edited / new cell with the old and new content, and the row map of the new block.
  python3 -I build_v8.py final STAGE1_XLSX PATCH_JSON OUT_XLSX [NOTES_JSON]
      stage 1 + the cached values in PATCH_JSON (patches_v8.py: LibreOffice values of the edited cells and of every
      formula cell downstream of them) + (optionally) the Collation Notes part rendered from NOTES_JSON.
Every other zip part is copied byte for byte, in the same order. sharedStrings.xml and styles.xml are untouched (new text
is written as inline strings; new cells reuse existing Depreciation Schedule styles).
"""
import sys, re, json, zipfile, os
from xml.sax.saxutils import escape
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

DS = "Depreciation Schedule"
# styles already used on the Depreciation Schedule (v3): text 0, title 1, section 2, header 3, input $ 140, input int 141,
# date header 144, units 145, bold label 146, money 147, GL 148
S_TXT, S_SEC, S_HDR, S_IN, S_DATE, S_UNIT, S_BOLD, S_MON, S_GL = 0, 2, 3, 141, 144, 145, 146, 147, 148
MCOL = [get_column_letter(c) for c in range(2, 18)]          # B..Q = Sep-26 .. Dec-27 (same columns as Prod Payroll)
Y26 = MCOL[:4]                                                # Sep-Dec 26
Y27 = MCOL[4:]                                                # F..Q = Jan-Dec 27 (same columns as the schedule's months)
PP = "'Prod Payroll'"
XL_RANGE = "$B$20:$O$20"


def part_of(z, sheet):
    wbxml = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    nm = escape(sheet, {'"': "&quot;"})
    rid = re.search(rf'<sheet name="{re.escape(nm)}"[^>]*r:id="(rId\d+)"', wbxml).group(1)
    return "xl/" + re.search(rf'<Relationship Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)


def xlookup_equiv(x):
    """XLOOKUP(x, r, r, MAX(r), 1) on an ascending, unique tier row r: the smallest tier >= x, or MAX(r) above every tier."""
    r = XL_RANGE
    return (f"IF({x}>MAX({r}),MAX({r}),INDEX({r},IFERROR(MATCH({x},{r},1),0)+ISNA(MATCH({x},{r},0))))")


# ------------------------------------------------------------------ the new Depreciation Schedule block (rows 128+)
def block():
    """Returns (rows, rowmap). rows: {r: [(col, kind, content, style)]}, kind in text|num|f."""
    rows = {}; R = {}

    def put(r, col, kind, content, style):
        rows.setdefault(r, []).append((col, kind, content, style))

    def months(r, fn, style=S_MON, cols=MCOL):
        for c in cols:
            put(r, c, "f", fn(c), style)

    r = 128
    put(r, "A", "text", "8. CAPITALISED PRODUCTION LABOUR (added in v8; user-confirmed 2026-10-10: production labour is capitalised)", S_SEC)
    put(129, "A", "text", "Prod Payroll monthly cost (rows 41, 62, 79; model unchanged) is capitalised into the units built that month: split lift / tray "
        "(supervisors pro rata to the lift : tray labour), pooled by build month (a month with no units built rolls into the next build month), "
        "divided by the units built, added to the go-live vintage (B133 months after the build) and depreciated with the section 1 lives, salvage "
        "and start rule. Section 9 feeds PO rows 20/21 (U:AF). PO 5110-5140 2027 = (1 - B131) x the v7 amount = 0 (capitalised).", S_TXT)
    put(130, "A", "text", "Input", S_HDR); put(130, "B", "text", "Value", S_HDR); put(130, "C", "text", "Unit", S_HDR)
    put(131, "A", "text", "Current Prod Payroll staff (Prod Payroll B9:C9) capitalised: 1 = yes (booked); 0 = expensed in PO 5110-5140 (sensitivity)", S_TXT)
    put(131, "B", "num", 1, S_IN); put(131, "C", "text", "switch", S_TXT)
    put(132, "A", "text", "Sep-Dec 26 model labour capitalised into the 2027 go-lives (Jan-Feb 27 vintages): 1 = yes (booked); 0 = no", S_TXT)
    put(132, "B", "num", 1, S_IN); put(132, "C", "text", "switch", S_TXT)
    put(133, "A", "text", "Build-to-go-live lag (months; DeploySum note 1)", S_TXT)
    put(133, "B", "num", 2, S_IN); put(133, "C", "text", "months", S_TXT)
    put(134, "A", "text", "Check: DeploySum go-lives Jan-Dec 27 (rows 19/20) less builds 2 months earlier (rows 24/25), all units (must be 0)", S_TXT)
    put(134, "B", "f", "SUMPRODUCT(ABS(DeploySum!F19:Q19-DeploySum!D24:O24))+SUMPRODUCT(ABS(DeploySum!F20:Q20-DeploySum!D25:O25))", S_UNIT)
    R["sw_current"], R["sw_q4"], R["lag"], R["lag_check"] = 131, 132, 133, 134

    r = 136; R["hdr"] = r
    put(r, "A", "text", "Build month", S_HDR)
    months(r, lambda c: f"DeploySum!{c}4", S_DATE)
    put(r, "R", "text", "FY2027", S_HDR); put(r, "S", "text", "Source / formula", S_HDR)
    src = [(137, "lift", "Lift payroll (Prod Payroll row 41)", 41), (138, "tray", "Tray payroll (Prod Payroll row 62)", 62),
           (139, "sup", "Supervisor payroll (Prod Payroll row 79)", 79), (140, "total", "Total production payroll (Prod Payroll row 81)", 81)]
    for rr, key, lab, prow in src:
        put(rr, "A", "text", lab, S_TXT)
        months(rr, lambda c, prow=prow: f"{PP}!{c}{prow}")
        put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON)
        R["pp_" + key] = rr
    put(137, "S", "text", "Prod Payroll rows 41 / 62 / 79 / 81, same columns (Sep-26..Dec-27); R = Jan-Dec 27", S_TXT)
    put(141, "A", "text", "Current lift staff at model rates (B9 x B5/12, raise M5 from M6, absence row 33, overtime row 35, tax B6, benefits B7)", S_TXT)
    months(141, lambda c: f"{PP}!$B$9*{PP}!$B$5/12*(1+IF({c}$136>={PP}!$M$6,{PP}!$M$5,0))*((1+{PP}!{c}$33+{PP}!{c}$35*{PP}!$M$10)*(1+{PP}!$B$6)+{PP}!$B$7)")
    put(141, "R", "f", "SUM(F141:Q141)", S_MON)
    put(142, "A", "text", "Current tray staff at model rates (C9 x C5/12, raise M5 from M6, absence row 54, overtime row 56, tax C6, benefits C7)", S_TXT)
    months(142, lambda c: f"{PP}!$C$9*{PP}!$C$5/12*(1+IF({c}$136>={PP}!$M$6,{PP}!$M$5,0))*((1+{PP}!{c}$54+{PP}!{c}$56*{PP}!$M$10)*(1+{PP}!$C$6)+{PP}!$C$7)")
    put(142, "R", "f", "SUM(F142:Q142)", S_MON)
    put(141, "S", "text", "Same cost formula as Prod Payroll rows 32-38 / 53-59 at the current headcount (no hires)", S_TXT)
    R["cur_lift"], R["cur_tray"] = 141, 142

    def q4(c):
        return "*$B$132" if c in Y26 else ""
    put(143, "A", "text", "Lift labour capitalised before supervisors (row 137 less (1 - B131) x row 141)", S_TXT)
    months(143, lambda c: f"({c}137-(1-$B$131)*{c}141){q4(c)}")
    put(144, "A", "text", "Tray labour capitalised before supervisors (row 138 less (1 - B131) x row 142)", S_TXT)
    months(144, lambda c: f"({c}138-(1-$B$131)*{c}142){q4(c)}")
    put(145, "A", "text", "Supervisor labour capitalised (row 139)", S_TXT)
    months(145, lambda c: f"{c}139{q4(c)}")
    put(146, "A", "text", "Supervisors to lifts (pro rata to rows 143 : 144)", S_TXT)
    months(146, lambda c: f"IF({c}143+{c}144=0,0,{c}145*{c}143/({c}143+{c}144))")
    put(147, "A", "text", "Capitalised labour - SlipLifts", S_BOLD)
    months(147, lambda c: f"{c}143+{c}146")
    put(148, "A", "text", "Capitalised labour - SlipTrays", S_BOLD)
    months(148, lambda c: f"{c}144+{c}145-{c}146")
    put(149, "A", "text", "Capitalised labour - total", S_BOLD)
    months(149, lambda c: f"{c}147+{c}148")
    for rr in range(143, 150):
        put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON)
    put(143, "S", "text", "B:E (Sep-Dec 26) x B132; B131 = 1 capitalises the current staff too", S_TXT)
    R.update(cap_lift_pre=143, cap_tray_pre=144, cap_sup=145, sup_to_lift=146, cap_lift=147, cap_tray=148, cap_total=149)

    put(150, "A", "text", "SlipLifts built (DeploySum row 24)", S_TXT)
    months(150, lambda c: f"DeploySum!{c}24", S_UNIT)
    put(151, "A", "text", "SlipTrays built (DeploySum row 25)", S_TXT)
    months(151, lambda c: f"DeploySum!{c}25", S_UNIT)
    for rr in (150, 151):
        put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_UNIT)
    for (pool, wait, cap, units, what) in ((152, 153, 147, 150, "Lift"), (154, 155, 148, 151, "Tray")):
        put(pool, "A", "text", f"{what} labour pooled into the month's builds (a month with no {what.lower()}s built rolls into the next build month)", S_TXT)
        put(wait, "A", "text", f"{what} labour waiting for the next {what.lower()} build", S_TXT)
        for i, c in enumerate(MCOL):
            prev = "" if i == 0 else f"+{MCOL[i - 1]}{wait}"
            put(pool, c, "f", f"IF({c}{units}>0,{c}{cap}{prev},0)", S_MON)
            put(wait, c, "f", f"IF({c}{units}>0,0,{c}{cap}{prev})", S_MON)
    put(156, "A", "text", "Labour per SlipLift built ($ per unit)", S_BOLD)
    months(156, lambda c: f"IF({c}150>0,{c}152/{c}150,0)")
    put(157, "A", "text", "Labour per SlipTray built ($ per unit)", S_BOLD)
    months(157, lambda c: f"IF({c}151>0,{c}154/{c}151,0)")
    put(158, "A", "text", "Go-live month of these builds (build month + B133)", S_TXT)
    months(158, lambda c: f"EOMONTH({c}136,$B$133)", S_DATE)
    put(159, "A", "text", "Labour in builds that go live in 2027 (Jan-Dec 27 vintages, depreciated from go-live)", S_TXT)
    months(159, lambda c: f"IF(YEAR({c}158)=2027,{c}152+{c}154,0)")
    put(160, "A", "text", "Labour in builds that go live in 2028 (CIP / inventory at Dec-27, no 2027 depreciation)", S_TXT)
    months(160, lambda c: f"IF(YEAR({c}158)>2027,{c}152+{c}154,0)")
    put(161, "A", "text", "Labour in builds that went live in 2026 (Nov-26 go-lives; inside the PO run rate, not added)", S_TXT)
    months(161, lambda c: f"IF(YEAR({c}158)<2027,{c}152+{c}154,0)")
    put(152, "S", "text", "Rows 152-161 run over all build months B:Q (Sep-26..Dec-27); totals in section 10", S_TXT)
    R.update(units_lift=150, units_tray=151, pool_lift=152, wait_lift=153, pool_tray=154, wait_tray=155, per_lift=156, per_tray=157,
             golive=158, lab_2027=159, lab_2028=160, lab_2026=161)

    r = 163; R["vhdr"] = r
    for col, t in (("A", "Capitalised labour by go-live vintage"), ("B", "GL"), ("C", "Go-live month #"), ("D", "Labour in vintage ($)"),
                   ("E", "$/month per row"), ("R", "FY2027"), ("S", "Source / formula")):
        put(r, col, "text", t, S_HDR)
    months(r, lambda c: f"{c}$27", S_DATE, Y27)
    MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    R["vint_lift"], R["vint_tray"] = [], []
    for (head, first, gl, life, per, units, key, nm) in ((164, 165, "$B$14", "$B$9", 156, 29, "vint_lift", "SlipLift"),
                                                        (177, 178, "$B$15", "$B$10", 157, 30, "vint_tray", "SlipTray")):
        put(head, "A", "text", f"New {nm}s (2027 go-lives): capitalised labour, one row per go-live month", S_BOLD)
        for k in range(12):
            rr = first + k
            put(rr, "A", "text", f"Labour - {nm} go-live {MON[k]}-27", S_TXT)
            put(rr, "B", "f", gl, S_GL)
            put(rr, "C", "num", k + 1, S_TXT)
            put(rr, "D", "f", f"INDEX($B${per}:$Q${per},4+$C{rr}-$B$133)*INDEX($F${units}:$Q${units},$C{rr})", S_MON)
            put(rr, "E", "f", f"D{rr}*(1-$B$12)/{life}", S_MON)
            months(rr, lambda c, rr=rr, life=life: f"IF(AND({c}$28>=$C{rr}+$B$13,{c}$28<$C{rr}+$B$13+{life}),$E{rr},0)", S_MON, Y27)
            put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON)
            R[key].append(rr)
        put(first, "S", "text", "D = labour per unit built (row %d, build month = go-live - B133) x units going live (row %d); E = D x (1 - salvage) / life; "
            "on from go-live month + start delay (same rule as rows 38-62)" % (per, units), S_TXT)

    r = 191
    put(r, "A", "text", "9. OUTPUT: capitalised-labour depreciation by GL (added to PO rows 20 / 21, columns U:AF)", S_SEC)
    for col, t in (("A", "GL line"), ("B", "GL"), ("C", "PO row"), ("R", "FY2027"), ("S", "Source / formula")):
        put(192, col, "text", t, S_HDR)
    months(192, lambda c: f"{c}$27", S_DATE, Y27)
    for rr, lab, gl, porow, tgt in ((193, "5011 - SlipLift Depreciation - capitalised production labour", "$B$14", 20, "-> PO!U20:AF20 (added to row 80)"),
                                    (194, "5012 - SlipCarrier Depreciation - capitalised production labour", "$B$15", 21, "-> PO!U21:AF21 (added to row 81)")):
        put(rr, "A", "text", lab, S_TXT); put(rr, "B", "f", gl, S_GL); put(rr, "C", "num", porow, S_TXT)
        months(rr, lambda c, rr=rr: f"SUMIF($B$165:$B$189,$B{rr},{c}$165:{c}$189)", S_MON, Y27)
        put(rr, "R", "f", f"SUM(F{rr}:Q{rr})", S_MON); put(rr, "S", "text", tgt, S_TXT)
    put(195, "A", "text", "Total capitalised-labour depreciation 2027", S_BOLD)
    months(195, lambda c: f"{c}193+{c}194", S_MON, Y27)
    put(195, "R", "f", "SUM(F195:Q195)", S_MON)
    put(196, "A", "text", "Check: rows 193:194 less all labour vintage rows (must be 0; a non-0 means a GL is not 5011/5012)", S_TXT)
    put(196, "R", "f", "R195-SUM(R165:R189)", S_MON)
    put(197, "A", "text", "Total depreciation 2027 incl. capitalised labour (R83 + R195) = PO rows 19-22", S_BOLD)
    put(197, "R", "f", "R83+R195", S_MON)
    R.update(out_lift=193, out_tray=194, out_total=195, out_check=196, dep_total=197)

    put(199, "A", "text", "10. MEMO: where the capitalised labour sits", S_SEC)
    put(200, "A", "text", "Item", S_HDR); put(200, "B", "text", "Value", S_HDR)
    memo = [
        (201, "pp_2027", "Production payroll 2027 (Prod Payroll!S81)", f"{PP}!S81"),
        (202, "cap_2027", "Capitalised labour from 2027 payroll (row 149, Jan-Dec 27)", "R149"),
        (203, "to_2027", "  in builds for 2027 go-lives (Feb-Oct 27 builds; Jan-27 rolls into Feb-27)", "SUM(F159:Q159)"),
        (204, "to_2028", "  in builds for 2028 go-lives (Nov-Dec 27 builds): CIP / inventory at Dec-27", "SUM(F160:Q160)"),
        (205, "waiting", "  waiting for a build at Dec-27", "Q153+Q155"),
        (206, "chk_2027", "Check: row 202 less rows 203:205 (must be 0)", "B202-SUM(B203:B205)"),
        (207, "cap_2026", "Capitalised labour from Sep-Dec 26 model months (row 149, B:E)", "SUM(B149:E149)"),
        (208, "from26_to_2027", "  in Jan-Feb 27 go-lives (Nov-Dec 26 builds; Oct-26 rolls into Nov-26)", "SUM(B159:E159)"),
        (209, "from26_to_2026", "  in Nov-26 go-lives (Sep-26 builds; in the PO run rate, not added)", "SUM(B161:E161)"),
        (210, "vint_total", "Depreciable labour in the 2027 go-live vintages (column D, rows 165-189)", "SUM(D165:D176)+SUM(D178:D189)"),
        (211, "chk_vint", "Check: row 210 less (row 203 + row 208) (must be 0)", "B210-B203-B208"),
        (212, "dep_2027", "2027 depreciation of capitalised labour (R195)", "R195"),
        (213, "nbv_end27", "Capitalised labour in service at Dec-27, net of 2027 depreciation (row 210 - row 212)", "B210-B212"),
        (214, "cur_2027", "Current staff at model rates, 2027 (rows 141:142)", "SUM(F141:Q142)"),
        (215, "avg_lift", "Average labour per SlipLift built in 2027 (rows 147, 150)", "SUM(F147:Q147)/SUM(F150:Q150)"),
        (216, "avg_tray", "Average labour per SlipTray built in 2027 (rows 148, 151)", "SUM(F148:Q148)/SUM(F151:Q151)"),
    ]
    for rr, key, lab, f in memo:
        put(rr, "A", "text", lab, S_TXT); put(rr, "B", "f", f, S_MON); R["m_" + key] = rr
    return rows, R


def cell_xml(ref, kind, content, style):
    if kind == "text":
        return f'<c r="{ref}" s="{style}" t="inlineStr"><is><t xml:space="preserve">{escape(content)}</t></is></c>'
    if kind == "num":
        return f'<c r="{ref}" s="{style}" t="n"><v>{content}</v></c>'
    return f'<c r="{ref}" s="{style}" t="n"><f aca="false">{escape(content)}</f><v>0</v></c>'


def colkey(col):
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def stage1(v7, out, spec_out):
    z = zipfile.ZipFile(v7)
    p_pp, p_ds, p_po = part_of(z, "Prod Payroll"), part_of(z, DS), part_of(z, "PO")
    ss = z.read("xl/sharedStrings.xml").decode()
    sst = [re.sub(r"<[^>]+>", "", s) for s in re.findall(r"<si>(.*?)</si>", ss, re.S)]
    from xml.sax.saxutils import unescape
    spec = {"xlookup": [], "ds_new": [], "po": [], "rowmap": None}

    # 1. Prod Payroll row 48
    x = z.read(p_pp).decode("utf-8")
    pat = re.compile(r'<c r="([B-Q])48"( s="\d+")? t="e"><f aca="false">(_xlfn\.xlookup\(MAX\(\$C\$8,([B-Q])47\),'
                     r'\$B\$20:\$O\$20,\$B\$20:\$O\$20,MAX\(\$B\$20:\$O\$20\),1\))</f><v>#NAME\?</v></c>')

    def rpp(m):
        col, s, old, col2 = m.group(1), m.group(2) or "", m.group(3), m.group(4)
        assert col == col2
        new = xlookup_equiv(f"MAX($C$8,{col}47)")
        spec["xlookup"].append({"cell": f"{col}48", "old": "=" + unescape(old), "new": "=" + new})
        return f'<c r="{col}48"{s}><f aca="false">{escape(new)}</f><v>0</v></c>'
    x, n = pat.subn(rpp, x)
    assert n == 16, ("expected 16 XLOOKUP cells in Prod Payroll row 48", n)
    assert "xlookup" not in x.lower(), "an XLOOKUP is left in Prod Payroll"
    new_pp = x

    # 2. Depreciation Schedule block
    x = z.read(p_ds).decode("utf-8")
    assert re.findall(r'<row r="(\d+)"', x)[-1] == "126" and '<dimension ref="A1:S126"/>' in x
    rows, rowmap = block()
    xml_rows = []
    for r in sorted(rows):
        cells = sorted(rows[r], key=lambda t: colkey(t[0]))
        assert len({c[0] for c in cells}) == len(cells), ("duplicate cell", r)
        xml_rows.append(f'<row r="{r}">' + "".join(cell_xml(f"{c}{r}", k, v, s) for c, k, v, s in cells) + "</row>")
        for c, k, v, s in cells:
            spec["ds_new"].append({"cell": f"{c}{r}", "kind": k, "content": ("=" + v) if k == "f" else v})
    last = max(rows)
    x = x.replace('<dimension ref="A1:S126"/>', f'<dimension ref="A1:S{last}"/>', 1)
    i = x.index("</sheetData>")
    new_ds = x[:i] + "".join(xml_rows) + x[i:]
    spec["rowmap"] = rowmap

    # 3. PO rows 20, 21 (5011, 5012) and 33-35 (5110-5140), months U:AF; notes in AI
    x = z.read(p_po).decode("utf-8")
    po_cols = [get_column_letter(c) for c in range(21, 33)]           # U..AF
    for row, outrow, base in ((20, rowmap["out_lift"], 80), (21, rowmap["out_tray"], 81)):
        for k, col in enumerate(po_cols):
            dcol = Y27[k]
            old_f = f"&apos;Depreciation Schedule&apos;!{dcol}{base}"
            pat = re.compile(rf'<c r="{col}{row}" s="(\d+)" t="n"><f aca="false">{re.escape(old_f)}</f><v>([^<]*)</v></c>')
            m = pat.search(x); assert m, (col, row)
            new_f = f"'Depreciation Schedule'!{dcol}{base}+'Depreciation Schedule'!{dcol}{outrow}"
            x = x[:m.start()] + f'<c r="{col}{row}" s="{m.group(1)}" t="n"><f aca="false">{escape(new_f, {chr(39): "&apos;"})}</f><v>{m.group(2)}</v></c>' + x[m.end():]
            spec["po"].append({"cell": f"{col}{row}", "old": "=" + unescape(old_f.replace("&apos;", "'")), "new": "=" + new_f})
    for row in (33, 34, 35):
        for col in po_cols:
            pat = re.compile(rf'<c r="{col}{row}" s="(\d+)" t="n"><v>([^<]*)</v></c>')
            m = pat.search(x); assert m, (col, row)
            float(m.group(2))
            new_f = f"(1-'Depreciation Schedule'!$B${rowmap['sw_current']})*{m.group(2)}"
            x = x[:m.start()] + f'<c r="{col}{row}" s="{m.group(1)}" t="n"><f aca="false">{escape(new_f, {chr(39): "&apos;"})}</f><v>{m.group(2)}</v></c>' + x[m.end():]
            spec["po"].append({"cell": f"{col}{row}", "old": m.group(2), "new": "=" + new_f})
    notes = {20: None, 21: None, 33: None, 34: None, 35: None}
    for row, outrow in ((20, rowmap["out_lift"]), (21, rowmap["out_tray"])):
        m = re.search(rf'<c r="AI{row}" s="(\d+)" t="s"><v>(\d+)</v></c>', x); assert m, row
        old = unescape(sst[int(m.group(2))]).replace("&apos;", "'")
        new = old + f" v8: + capitalised production labour ('Depreciation Schedule' row {outrow}, section 9)."
        x = x[:m.start()] + f'<c r="AI{row}" s="{m.group(1)}" t="inlineStr"><is><t xml:space="preserve">{escape(new)}</t></is></c>' + x[m.end():]
        spec["po"].append({"cell": f"AI{row}", "old": old, "new": new})
    gl = {33: "5110", 34: "5130", 35: "5140"}
    for row in (33, 34, 35):
        m = re.search(rf'<c r="AI{row}" s="(\d+)"/>', x); assert m, row
        new = (f"v8: {gl[row]} 2027 = 0, CAPITALISED (production labour capitalised, user-confirmed 2026-10-10). Formula = (1 - 'Depreciation "
               f"Schedule'!B{rowmap['sw_current']}) x the v7 amount (PO HC roster); the labour is in 5011/5012 depreciation via 'Depreciation Schedule' "
               f"sections 8-9. B{rowmap['sw_current']} = 0 restores the v7 amount (sensitivity).")
        x = x[:m.start()] + f'<c r="AI{row}" s="{m.group(1)}" t="inlineStr"><is><t xml:space="preserve">{escape(new)}</t></is></c>' + x[m.end():]
        spec["po"].append({"cell": f"AI{row}", "old": "", "new": new})
    new_po = x

    repl = {p_pp: new_pp, p_ds: new_ds, p_po: new_po}
    zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for info in z.infolist():
        data = z.read(info.filename)
        if info.filename in repl:
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = info.external_attr
            zo.writestr(zi, repl[info.filename].encode("utf-8"))
        else:
            zo.writestr(info, data)
    zo.close()
    spec["parts"] = {"Prod Payroll": p_pp, DS: p_ds, "PO": p_po}
    json.dump(spec, open(spec_out, "w"), indent=1)
    print(f"stage1 {out}: Prod Payroll XLOOKUP cells replaced {len(spec['xlookup'])}; Depreciation Schedule new cells {len(spec['ds_new'])} "
          f"(rows 128-{last}); PO cells edited {len(spec['po'])}; parts changed {sorted(repl)}")


# ------------------------------------------------------------------ final: value patch + Collation Notes
CELL_RX = re.compile(r'<c r="([A-Z]+\d+)"([^>]*?)(/>|>(.*?)</c>)', re.S)


def patch_values(xml, items):
    """items: {coord: (type, vtext)} type in n|str|b|e. Only formula cells are patched; formula and style kept."""
    done = set()

    def sub(m):
        coord = m.group(1)
        if coord not in items:
            return m.group(0)
        attrs, body = m.group(2), m.group(4) or ""
        fm = re.match(r"(<f[^>]*>[^<]*</f>|<f[^>]*/>)", body)
        assert fm, (coord, "patched cell must hold a formula")
        s = re.search(r' s="\d+"', attrs)
        typ, v = items[coord]
        tattr = {"n": ' t="n"', "str": ' t="str"', "b": ' t="b"', "e": ' t="e"'}[typ]
        done.add(coord)
        return f'<c r="{coord}"{s.group(0) if s else ""}{tattr}>{fm.group(1)}<v>{v}</v></c>'
    out = CELL_RX.sub(sub, xml)
    missing = set(items) - done
    assert not missing, ("cells to patch not found", sorted(missing)[:10])
    return out, len(done)


def final(stage, patch_json, out, notes_json=None):
    z = zipfile.ZipFile(stage)
    P = json.load(open(patch_json))
    by_part = {}
    for s, coord, typ, v in P["values"]:
        by_part.setdefault(part_of(z, s), {})[coord] = (typ, v)
    cn_part = part_of(z, "Collation Notes")
    assert cn_part not in by_part
    styles = z.read("xl/styles.xml").decode()
    zo = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    changed = []; npatched = 0; cn_info = ""
    for info in z.infolist():
        data = z.read(info.filename); new = None
        if info.filename in by_part:
            x, k = patch_values(data.decode("utf-8"), by_part[info.filename]); npatched += k
            new = x.encode("utf-8")
        elif info.filename == cn_part and notes_json:
            import cnotes_v6
            xml, last, nform = cnotes_v6.render(data.decode("utf-8"), json.load(open(notes_json)), styles)
            new = xml.encode("utf-8"); cn_info = f"; Collation Notes regenerated ({last} rows, {nform} formulas)"
        if new is None:
            zo.writestr(info, data)
        else:
            zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zi.compress_type = zipfile.ZIP_DEFLATED; zi.external_attr = info.external_attr
            zo.writestr(zi, new); changed.append(info.filename)
    zo.close()
    assert npatched == len(P["values"])
    print(f"final {out}: {npatched} cached values patched in {sorted(by_part)}; parts changed {changed}{cn_info}")


if __name__ == "__main__":
    if sys.argv[1] == "stage1":
        stage1(*sys.argv[2:5])
    elif sys.argv[1] == "final":
        final(*sys.argv[2:6])
    else:
        raise SystemExit(__doc__)
