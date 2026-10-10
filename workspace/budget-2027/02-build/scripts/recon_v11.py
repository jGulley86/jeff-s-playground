"""v11 recon-lite for paylist_v7.amounts(): the Logistics&WH seats with the two held by people on an HC roster
(LOG-01, MH-01) flagged, recomputed from the workbook (same arithmetic as figures_v10.py; no names read or written).

Usage:
  python3 -I recon_v11.py V11_XLSX OUT_RECON_JSON
"""
import sys, re, json
import openpyxl

X, OUT = sys.argv[1:3]
wb = openpyxl.load_workbook(X, data_only=True)
lw = wb["Logistics&WH Payroll"]
n = lambda v: float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
S27 = range(6, 18)  # F..Q = Jan..Dec 27
tax, ben = n(lw["B5"].value), n(lw["B6"].value)
nh_per = n(lw["B10"].value)
seats = []
for r in range(14, 23):
    lab = re.sub(r"\s*\(.*?\)\s*$", "", str(lw.cell(r, 1).value)).strip()
    code = lab.split()[0]
    sal27 = sum(n(lw.cell(r + 20, c).value) for c in S27)
    start = lw.cell(r, 21).value
    nh27 = nh_per if (start is not None and start.year == 2027) else 0.0
    seats.append({"annual": n(lw.cell(r, 24).value), "cost_2027": sal27 * (1 + tax + ben) + nh27,
                  "overlap": code in ("LOG-01", "MH-01")})
assert sum(1 for s in seats if s["overlap"]) == 2
json.dump({"lw": {"seats": seats}}, open(OUT, "w"), indent=1)
print(f"recon_v11: {len(seats)} Logistics&WH seats, 2 flagged as roster-held (LOG-01, MH-01)")
