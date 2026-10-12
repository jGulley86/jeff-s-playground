"""v13 recon-lite for paylist_v7.amounts(): the Logistics&WH seats; the seats excluded from the CO add-on because a
person on an HC roster holds/realizes them (LOG-01, MH-01, TL-01, MH-02) are flagged, so their amounts join the
per-person rounding guard. Same arithmetic as recon_v11 (no names read or written).

Usage:
  python3 -I recon_v13.py V13_XLSX OUT_RECON_JSON
"""
import sys, re, json
import openpyxl

X, OUT = sys.argv[1:3]
wb = openpyxl.load_workbook(X, data_only=True)
lw = wb["Logistics&WH Payroll"]
n = lambda v: float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0
S27 = range(6, 18)
tax, ben = n(lw["B5"].value), n(lw["B6"].value)
nh_per = n(lw["B10"].value)
seats = []
for r in range(14, 23):
    lab = re.sub(r"\s*\(.*?\)\s*$", "", str(lw.cell(r, 1).value)).strip()
    code = lab.split()[0]
    sal27 = sum(n(lw.cell(r + 20, c).value) for c in S27)
    start = lw.cell(r, 21).value
    nh27 = nh_per if (start is not None and start.year == 2027) else 0.0
    seats.append({"annual": n(lw.cell(r, 24).value), "cost_2027": sal27 * (1 + tax + ben) + nh27,
                  "overlap": code in ("LOG-01", "MH-01", "TL-01", "MH-02")})
json.dump({"lw": {"seats": seats}}, open(OUT, "w"), indent=1)
print(f"recon_v13: {len(seats)} seats, {sum(s['overlap'] for s in seats)} flagged as roster-held/realized -> {OUT}")
