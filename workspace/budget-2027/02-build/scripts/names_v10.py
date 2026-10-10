"""v10 scratch name list for privacy_v6.py (WORK dir only; never in the deliverables, deleted by run_build_v10.sh).

Usage:
  python3 -I names_v10.py V10_XLSX NAME_FILE OUT_NAMES

Names come from: the HC tabs (column B of person rows, i.e. rows 14-64 with a title in column C), Fld Eng Budget A13:A21
('Title – Name'), the Logistics&WH Payroll tab (names in brackets, 'backfill for ...', 'LOG-02 <name>', 'aligned to <name>'),
and NAME_FILE (the new CO employee, given by the user in chat). Nothing is printed except counts.
"""
import sys, re, json
import openpyxl

X, NF, OUT = sys.argv[1:4]
wb = openpyxl.load_workbook(X, data_only=True)
full = set()
for d in ("FO", "PO", "CO", "FE", "PE"):
    ws = wb[f"HC-{d}"]
    for r in range(14, 65):
        b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(b, str) and b.strip() and isinstance(c, str) and c.strip():
            full.add(b.strip())
feb = wb["Fld Eng Budget"]
for r in range(13, 22):
    parts = [p.strip() for p in str(feb.cell(r, 1).value).split("–")]
    if len(parts) >= 2 and re.fullmatch(r"[A-Z][\w'.-]+(?: [A-Z][\w'.-]+)+", parts[1]):
        full.add(parts[1])
single = set()
lw = wb["Logistics&WH Payroll"]
NAME = r"([A-Z][a-z'.-]+(?: [A-Z][a-z'.-]+)+)"
for row in lw.iter_rows():
    for c in row:
        if isinstance(c.value, str):
            for m in re.finditer(r"\(" + NAME + r"\)", c.value):
                full.add(m.group(1))
            for m in re.finditer(r"backfill for " + NAME, c.value):
                full.add(m.group(1))
            for m in re.finditer(r"LOG-0\d " + NAME, c.value):
                full.add(m.group(1))
            for m in re.finditer(r"aligned to ([A-Z][a-z]+)", c.value):
                single.add(m.group(1))
TITLE = {"logistics", "manager", "material", "handler", "handling", "team", "lead", "inventory", "control", "specialist", "coordinator",
         "warehouse", "supervisor"}
PLACEHOLDER = {"backfill", "tbd", "open", "new hire", "vacant"}
full = {f for f in full if f.strip().lower() not in PLACEHOLDER and not ({t.lower() for t in f.split()} & TITLE)}
new = open(NF, encoding="utf-8").read().strip()
full.add(new)
COMMON = {"field", "manager", "lead", "new", "hire", "budgeted", "material", "handler"}
tokens = {t.lower() for f in full for t in re.split(r"[\s,]+", f) if len(t) >= 2} | {s.lower() for s in single}
tokens -= COMMON
json.dump({"tokens": sorted(tokens), "full": sorted(full)}, open(OUT, "w"))
print(f"names_v10: {len(full)} full names, {len(tokens)} tokens (scratch file, WORK dir only)")
