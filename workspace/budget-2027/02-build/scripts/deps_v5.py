"""Static dependency trace: every formula cell downstream of one or more seed cells.

Usage (module):   from deps_v5 import downstream; cells, info = downstream(XLSX, [("Fld Eng Budget", "B21")])
Usage (script):   python3 -I deps_v5.py XLSX OUT_JSON SHEET!CELL [SHEET!CELL ...]

How: every formula in every sheet is tokenised (openpyxl Tokenizer); each OPERAND/RANGE token is resolved to a
rectangle (sheet, r1, c1, r2, c2): A1, $A$1, A1:B2, A:A, 1:1, Sheet!A1, 'Sheet name'!A1:B2. Bare names are looked up
in the workbook's defined names; a defined name that points at cells of this workbook is resolved to its rectangle(s).
A name that cannot be resolved stops the script (so nothing is silently missed). INDIRECT / OFFSET / external
references also stop the script: a static trace cannot follow them. The closure is computed to a fixed point.
"""
import sys, json, re, collections
import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

MAXR, MAXC = 1048576, 16384
CELL = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")
COL = re.compile(r"^\$?([A-Z]{1,3})$")
ROW = re.compile(r"^\$?(\d+)$")


def split_sheet(ref, here):
    if "!" in ref:
        s, a = ref.rsplit("!", 1)
        if s.startswith("'") and s.endswith("'"):
            s = s[1:-1].replace("''", "'")
        return s, a
    return here, ref


def rect(addr):
    """'A1', 'A1:B2', 'A:B', '1:2' -> (r1, c1, r2, c2) or None if not an address."""
    parts = addr.replace("$", "").split(":")
    if len(parts) == 1:
        m = CELL.match(parts[0])
        if not m:
            return None
        r, c = int(m.group(2)), column_index_from_string(m.group(1))
        return (r, c, r, c)
    if len(parts) != 2:
        return None
    a, b = parts
    ma, mb = CELL.match(a), CELL.match(b)
    if ma and mb:
        r1, c1 = int(ma.group(2)), column_index_from_string(ma.group(1))
        r2, c2 = int(mb.group(2)), column_index_from_string(mb.group(1))
        return (min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
    ca, cb = COL.match(a), COL.match(b)
    if ca and cb:
        c1, c2 = column_index_from_string(ca.group(1)), column_index_from_string(cb.group(1))
        return (1, min(c1, c2), MAXR, max(c1, c2))
    ra, rb = ROW.match(a), ROW.match(b)
    if ra and rb:
        return (min(int(ra.group(1)), int(rb.group(1))), 1, max(int(ra.group(1)), int(rb.group(1))), MAXC)
    return None


def downstream(path, seeds):
    wb = openpyxl.load_workbook(path)
    sheets = set(wb.sheetnames)
    names = {}
    for n, d in wb.defined_names.items():
        names[n.upper()] = d.attr_text
    refs = {}                         # (sheet, coord) -> list of (sheet, r1, c1, r2, c2)
    used_names = collections.Counter()
    nform = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    v = v.text
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                nform += 1
                up = v.upper()
                if "INDIRECT(" in up or "OFFSET(" in up:
                    raise SystemExit(f"volatile reference in {ws.title}!{c.coordinate}: {v[:80]}")
                out = []
                for t in Tokenizer(v).items:
                    if t.type != Token.OPERAND or t.subtype != Token.RANGE:
                        continue
                    s, a = split_sheet(t.value, ws.title)
                    if "[" in s:
                        raise SystemExit(f"external reference in {ws.title}!{c.coordinate}: {v[:80]}")
                    r = rect(a)
                    if r is not None:
                        if s not in sheets:
                            raise SystemExit(f"unknown sheet {s!r} in {ws.title}!{c.coordinate}")
                        out.append((s,) + r)
                        continue
                    nm = t.value.upper()
                    if nm in ("TRUE", "FALSE"):
                        continue
                    if nm not in names:
                        raise SystemExit(f"unresolved operand {t.value!r} in {ws.title}!{c.coordinate}: {v[:80]}")
                    used_names[nm] += 1
                    for piece in names[nm].split(","):
                        s2, a2 = split_sheet(piece.strip(), ws.title)
                        r2 = rect(a2)
                        if r2 is None or s2 not in sheets:
                            raise SystemExit(f"defined name {nm} -> {names[nm][:80]} not resolvable ({ws.title}!{c.coordinate})")
                        out.append((s2,) + r2)
                refs[(ws.title, c.coordinate)] = out
    # index readers by referenced sheet
    by_sheet = collections.defaultdict(list)
    for cell, rr in refs.items():
        for (s, r1, c1, r2, c2) in rr:
            by_sheet[s].append((r1, c1, r2, c2, cell))
    seen = set()
    todo = [(s, a) for s, a in seeds]
    while todo:
        s, a = todo.pop()
        m = CELL.match(a)
        r, c = int(m.group(2)), column_index_from_string(m.group(1))
        for (r1, c1, r2, c2, cell) in by_sheet.get(s, ()):
            if r1 <= r <= r2 and c1 <= c <= c2 and cell not in seen:
                seen.add(cell)
                todo.append(cell)
    info = {"formulas": nform, "defined_names_used": dict(used_names), "downstream": len(seen),
            "by_sheet": dict(sorted(collections.Counter(s for s, _ in seen).items()))}     # sorted: deterministic output
    return seen, info


if __name__ == "__main__":
    path, out = sys.argv[1:3]
    seeds = [tuple(x.rsplit("!", 1)) for x in sys.argv[3:]]
    cells, info = downstream(path, seeds)
    json.dump({"seeds": seeds, "info": info, "cells": sorted([list(x) for x in cells])}, open(out, "w"), indent=0)
    print("deps:", info)
