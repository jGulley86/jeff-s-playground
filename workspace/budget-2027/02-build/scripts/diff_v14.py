"""v14 confined diff: v13 -> v14 cell-by-cell (formulas, cached values, comments, defined names).

Rules (CONSTRAINTS of the v14 handoff): the diff must be confined to the Collation Notes sheet, and even there to
TEXT/value cells only - v14 changes no formula anywhere (H78:H81 and C21:F28 included). A violation = any formula
difference on any sheet, any value difference outside Collation Notes, any comment or defined-name difference, or
any 2026-column (B:N) difference on the six P&L sheets. Exit 1 on violation. HC-tab coordinates would be reported
without values (none expected).

Usage:
  python3 -I diff_v14.py V13_XLSX V14_XLSX OUT_JSON
"""
import sys, json, datetime
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import get_column_letter

OLD, NEW, OUT = sys.argv[1:4]
TOL = 1e-6
isnum = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)


def norm_f(v):
    if isinstance(v, ArrayFormula):
        return ("array", v.ref, v.text or "")
    return v if isinstance(v, str) and v.startswith("=") else None


def norm_v(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


of, ov = openpyxl.load_workbook(OLD), openpyxl.load_workbook(OLD, data_only=True)
nf, nv = openpyxl.load_workbook(NEW), openpyxl.load_workbook(NEW, data_only=True)
assert of.sheetnames == nf.sheetnames, "sheet set/order changed"

viol, cn_cells, cn_count = [], [], 0
for sh in of.sheetnames:
    o_f, o_v, n_f, n_v = of[sh], ov[sh], nf[sh], nv[sh]
    rows = max(o_f.max_row, n_f.max_row)
    cols = max(o_f.max_column, n_f.max_column)
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            fo, fn = norm_f(o_f.cell(r, c).value), norm_f(n_f.cell(r, c).value)
            a, b = norm_v(o_v.cell(r, c).value), norm_v(n_v.cell(r, c).value)
            if fo is None and fn is None:
                a2 = o_f.cell(r, c).value; b2 = n_f.cell(r, c).value
                a, b = norm_v(a2) if a2 is not None else a, norm_v(b2) if b2 is not None else b
            f_changed = fo != fn
            if isnum(a) and isnum(b):
                v_changed = abs(a - b) > TOL
            else:
                v_changed = (a if a != "" else None) != (b if b != "" else None)
            if not f_changed and not v_changed:
                continue
            coord = f"{get_column_letter(c)}{r}"
            if f_changed:
                viol.append(f"{sh}!{coord}: formula changed (v14 changes no formula)")
            if v_changed:
                if sh == "Collation Notes":
                    cn_count += 1
                    if len(cn_cells) < 200:
                        cn_cells.append(coord)
                else:
                    viol.append(f"{sh}!{coord}: value changed outside Collation Notes")
# comments
for sh in of.sheetnames:
    oc = {c.coordinate: c.comment.text for row in of[sh].iter_rows() for c in row if c.comment}
    ncm = {c.coordinate: c.comment.text for row in nf[sh].iter_rows() for c in row if c.comment}
    if oc != ncm:
        viol.append(f"{sh}: comments changed")
# defined names
dn_o = {k: str(v.value) for k, v in of.defined_names.items()}
dn_n = {k: str(v.value) for k, v in nf.defined_names.items()}
if dn_o != dn_n:
    viol.append("defined names changed")
# 2026 columns of the six P&L sheets (B:N = 2-14)
for sh in ("COO P&L", "FO", "PO", "CO", "FE", "PE"):
    for r in range(1, of[sh].max_row + 1):
        for c in range(2, 15):
            a, b = ov[sh].cell(r, c).value, nv[sh].cell(r, c).value
            same = abs(a - b) <= TOL if (isnum(a) and isnum(b)) else (a if a != "" else None) == (b if b != "" else None)
            if not same:
                viol.append(f"{sh}!{get_column_letter(c)}{r}: 2026 column changed")

json.dump({"violations": viol, "collation_notes_changed_cells": cn_count, "cells": cn_cells,
           "defined_names": len(dn_n)}, open(OUT, "w"), indent=1)
print(f"diff_v14: {cn_count} text/value cells changed on Collation Notes; 0 formula diffs expected; "
      f"violations: {len(viol)}")
for v in viol[:40]:
    print("VIOLATION", v)
sys.exit(1 if viol else 0)
