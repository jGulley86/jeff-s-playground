"""v10 checks. Exit 1 if any completion criterion fails.

Usage:
  python3 -I analyze_v10.py V9_XLSX V10_XLSX V10_RECALC_XLSX V9_RECALC_XLSX DIFF_JSON SPEC_JSON FIG_JSON NOTES_MD NOTES_JSON \
      PRIVACY_TXT PAYLIST_TXT OUT_JSON

  - diff v9 -> v10 confined to the areas of handoff v10: the new Logistics&WH Payroll sheet; HC-CO row 18 and its totals (rows 67-70,
    2027 columns M:Z); CO rows 67-69 U:AF (+ AG / AH / AJ), their subtotals (rows 73, 116, 117, 132) and ratio rows (137-142), CO
    AI67:AI69 notes; PO U33:AF35 (formulas; values unchanged); the COO P&L roll-up (same rows); Collation Notes (regenerated). HC-COO
    (organisation roll-up of the HC tabs, formulas unchanged) is reported separately as a downstream area. Any other cell fails.
    0 changes in the 2026 columns (B:N) of COO P&L / FO / PO / CO / FE / PE;
  - package: parts changed / added;
  - LibreOffice recalculation of v10 = stored values on every sheet except Collation Notes; 0 new error cells vs the v9 recalculation;
    the Logistics&WH tab has 0 errors;
  - COO P&L = FO + PO + CO + FE + PE on every dept-sum row (2027 months and AG); check row 133 = 0;
  - Collation Notes: live formulas = recalculation; first section = CONFIDENTIAL banner;
  - figures: LOG-01 not double counted (CO check 0), MH-01 only in CO (PO base drops exactly PO HC r20; booked PO 5110-5140 = 0),
    contribution = v9 - logistics less LOG-01 - new CO employee; privacy_v6 / paylist_v10 found nothing.
"""
import sys, json, re, zipfile, datetime, collections
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string as ci
from openpyxl.worksheet.formula import ArrayFormula

V9, V10, LO10, LO9, DIFF, SPEC, FIG, MD, NJ, PRIV, PAY, OUT = sys.argv[1:13]
BANNER = "CONFIDENTIAL – contains named compensation data (HC-* tabs); restrict distribution"
ERR = re.compile(r"^#(REF!|NAME\?|VALUE!|DIV/0!|N/A|NUM!|NULL!)$")
fails = []; res = {}


def check(cond, msg):
    if not cond:
        fails.append(msg)


def norm(v):
    if isinstance(v, ArrayFormula):
        return ("array", v.ref, v.text)
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return v


isnum = lambda x: isinstance(x, (int, float)) and not isinstance(x, bool)
f9 = openpyxl.load_workbook(V9); v9 = openpyxl.load_workbook(V9, data_only=True)
f10 = openpyxl.load_workbook(V10); v10 = openpyxl.load_workbook(V10, data_only=True)
lo10 = openpyxl.load_workbook(LO10, data_only=True); lo9 = openpyxl.load_workbook(LO9, data_only=True)
F = json.load(open(FIG)); d = json.load(open(DIFF))
LW = "Logistics&WH Payroll"
PL = ("COO P&L", "FO", "PO", "CO", "FE", "PE")
U, AF, AG, AH, AJ = ci("U"), ci("AF"), ci("AG"), ci("AH"), ci("AJ")


def area(s, r, c):
    col = get_column_letter(c)
    if s == "HC-CO":
        if r == 18: return "HC-CO row 18: new CO employee (inputs B18:E18 + the row's formulas)"
        if r in (67, 68, 69, 70) and 13 <= c <= 26: return "HC-CO rows 67-70 (roster totals), 2027 columns M:Z"
    if s == "CO":
        if r in (67, 68, 69) and U <= c <= AF: return "CO U67:AF69 (6610 / 6630 / 6640, 2027): typed values -> formulas"
        if r in (67, 68, 69) and c in (AG, AH, AJ): return "CO rows 67-69 AG / AH / AJ (totals, ratio, YOY)"
        if col == "AI" and r in (67, 68, 69): return "CO AI67:AI69 (notes column)"
        if r in (73, 116, 117, 132) and U <= c <= AJ: return "CO subtotals rows 73, 116, 117, 132 (U:AJ)"
        if 137 <= r <= 142 and U <= c <= AG: return "CO ratio rows 137-142 (U:AG)"
    if s == "PO":
        if r in (33, 34, 35) and U <= c <= AF: return "PO U33:AF35 (5110-5140 B131 base; formulas, values unchanged)"
    if s == "COO P&L":
        if r in (67, 68, 69, 73, 116, 117, 132) and U <= c <= AJ: return "COO P&L roll-up rows 67-69, 73, 116, 117, 132 (U:AJ)"
        if 137 <= r <= 142 and U <= c <= AG: return "COO P&L ratio rows 137-142 (U:AG)"
    if s == "HC-COO":
        if 13 <= c <= 26: return "HC-COO roll-up of the HC tabs, 2027 columns M:Z (downstream of HC-CO; formulas unchanged; not in the handoff list)"
    return None


areas = collections.Counter(); outside = []; cols26 = 0
for s in f9.sheetnames:
    if s == "Collation Notes":
        continue
    a_f, b_f, a_v, b_v = f9[s], f10[s], v9[s], v10[s]
    mr = max(a_f.max_row, b_f.max_row); mc = max(a_f.max_column, b_f.max_column)
    for r in range(1, mr + 1):
        for c in range(1, mc + 1):
            fa, fb = norm(a_f.cell(r, c).value), norm(b_f.cell(r, c).value)
            va, vb = norm(a_v.cell(r, c).value), norm(b_v.cell(r, c).value)
            if fa == fb and va == vb:
                continue
            if s in PL and 2 <= c <= 14:
                cols26 += 1
            ar = area(s, r, c)
            if ar is None:
                outside.append(f"{s}!{get_column_letter(c)}{r}")
            else:
                areas[(s, ar)] += 1
check(not outside, f"changes outside the allowed areas: {outside[:10]}")
check(cols26 == 0, f"{cols26} changes in 2026 columns")
check(d["sheets_only_in_new"] == [LW] and not d["sheets_only_in_old"], "sheet set")
check([len(d["comments_added"]), len(d["comments_removed"]), len(d["comments_changed"])] == [0, 0, 0], "comments changed")
res["areas"] = [{"sheet": s, "area": a, "cells": n} for (s, a), n in areas.items()] + [{"sheet": LW, "area": "new sheet (copied from the SD book)", "cells": sum(1 for row in f10[LW].iter_rows() for c in row if c.value is not None)}]
res["outside"] = len(outside); res["cols26"] = cols26
res["diff"] = {"total": d["total"], "comments": [len(d["comments_added"]), len(d["comments_removed"]), len(d["comments_changed"])],
               "only_new": d["sheets_only_in_new"]}

# package parts
z9, z10 = zipfile.ZipFile(V9), zipfile.ZipFile(V10)
n9, n10 = set(z9.namelist()), set(z10.namelist())
res["parts_added"] = sorted(n10 - n9)
res["parts_changed"] = sorted(p for p in n9 & n10 if z9.read(p) != z10.read(p))
check(not (n9 - n10), "parts removed")

# recalculation vs stored, errors
mx = 0.0; ncells = 0; nonnum = []
err10 = set(); err9 = set()
for s in v10.sheetnames:
    for row in lo10[s].iter_rows():
        for c in row:
            if isinstance(c.value, str) and ERR.match(c.value):
                err10.add((s, c.coordinate))
    if s in lo9.sheetnames:
        for row in lo9[s].iter_rows():
            for c in row:
                if isinstance(c.value, str) and ERR.match(c.value):
                    err9.add((s, c.coordinate))
    if s == "Collation Notes":
        continue
    a, b = v10[s], lo10[s]
    for row in a.iter_rows():
        for c in row:
            x, y = c.value, b[c.coordinate].value
            if x is None and y is None:
                continue
            if isnum(x) and isnum(y):
                ncells += 1; mx = max(mx, abs(x - y))
            elif norm(x) != norm(y) and not (x in (None, "") and y in (None, "")):
                nonnum.append(f"{s}!{c.coordinate}")
res.update(recalc_max_diff=mx, recalc_cells=ncells, recalc_nonnum=len(nonnum), errors_v10=len(err10), errors_v9=len(err9),
           new_errors=len(err10 - err9), lw_errors=sum(1 for s, _ in err10 if s == LW))
check(mx <= 1e-6, f"stored vs recalculated max diff {mx}")
check(not nonnum, f"non-numeric stored vs recalculated differences {nonnum[:5]}")
check(not (err10 - err9), f"new errors {sorted(err10 - err9)[:5]}")
check(res["lw_errors"] == 0, "Logistics&WH errors")

# COO P&L = sum of depts
c_f = f10["COO P&L"]; L = lo10
smax = 0.0; nrows = 0
for r in range(1, c_f.max_row + 1):
    fv = c_f.cell(r, U).value
    if isinstance(fv, str) and re.fullmatch(rf"=SUM\(FO!U{r},PO!U{r},CO!U{r},FE!U{r},PE!U{r}\)", fv):
        nrows += 1
        for c in list(range(U, AF + 1)) + [AG]:
            tot = sum((L[dp].cell(r, c).value or 0) for dp in ("FO", "PO", "CO", "FE", "PE") if isnum(L[dp].cell(r, c).value))
            v = L["COO P&L"].cell(r, c).value
            smax = max(smax, abs((v if isnum(v) else 0) - tot))
r133 = max(abs(L["COO P&L"].cell(133, c).value or 0) for c in range(U, AF + 1))
res.update(sum_depts_max=smax, sum_depts_rows=nrows, row133_max=r133)
check(nrows > 50 and smax < 1e-6, f"COO P&L != sum of depts ({smax})")
check(r133 < 1e-6, "check row 133")

# Collation Notes formulas = recalculation; banner
cnf, cnv, cnl = f10["Collation Notes"], v10["Collation Notes"], lo10["Collation Notes"]
cnmax = 0.0; ncn = 0; cnbad = []
for row in cnf.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith("="):
            ncn += 1
            x, y = cnv[c.coordinate].value, cnl[c.coordinate].value
            if isnum(x) and isnum(y):
                cnmax = max(cnmax, abs(x - y))
            elif x != y:
                cnbad.append(c.coordinate)
NJd = json.load(open(NJ))
res.update(cn_formulas=ncn, cn_max_diff=cnmax, cn_bad=cnbad)
check(not cnbad and cnmax < 1e-6, "Collation Notes formulas differ from the recalculation")
check(NJd["sections"][0]["heading"] == BANNER, "banner")

# figures
co_chk = max(abs(F["co"][g]["check"]) for g in ("6610", "6630", "6640"))
recon = F["c9"] - F["lw"]["ex_total"] - F["newco"]["cost_2027"] - F["c10"]
res["fig"] = {"co_check_max": co_chk, "po_base_drop_minus_r20": F["po"]["base_v9"] - F["po"]["base_v10"] - F["po"]["r20_cost"],
              "po_booked": F["po"]["booked_v10"], "contribution_bridge_residual": recon}
check(co_chk < 1e-6, "LOG-01 double counted / CO check")
check(abs(res["fig"]["po_base_drop_minus_r20"]) < 1e-6 and abs(F["po"]["booked_v10"]) < 1e-9, "MH-01 not removed from the PO base")
check(abs(recon) < 1e-6, "contribution bridge")
check(round(F["lw"]["total_2027"], 2) == 547441.22 and round(F["lw"]["log01_cost"]) == 148949 and round(F["lw"]["mh01_cost"]) == 61963
      and round(F["lw"]["ex_total"]) == 398492, "Logistics&WH verified figures not reproduced")
# CO / PO formulas
co_f = all(isinstance(f10["CO"].cell(r, c).value, str) and f10["CO"].cell(r, c).value.startswith("=") for r in (67, 68, 69) for c in range(U, AF + 1))
check(co_f, "CO U67:AF69 not all formulas")
check(all("'Logistics&WH Payroll'!" in f10["CO"].cell(r, c).value and "34" in f10["CO"].cell(r, c).value for r in (67, 68, 69) for c in range(U, AF + 1)),
      "CO formulas do not exclude LOG-01 (row 34)")
check(all("'HC-PO'!" in f10["PO"].cell(r, c).value for r in (33, 34, 35) for c in range(U, AF + 1)), "PO base formulas")
# privacy / per-person
priv = open(PRIV).read(); pay = open(PAY).read()
check(re.search(r"hits 0\b", priv) is not None, "privacy_v6 found names")
check(re.search(r"exact hits 0;", pay) is not None, "paylist found exact per-person amounts")
md = open(MD, encoding="utf-8").read()
for h in ("## Changes from v9", "## Payroll reconciliation (v10)", "I-17 RESOLVED", "Sign-flip statement (computed, v10)", "Range (v10)",
          "the B131 sensitivity"):
    check(h in md, f"md lacks {h!r}")
res["criteria"] = collections.OrderedDict([
    ("COO P&L = sum of the 5 depts (every dept-sum row, 2027 months and AG)", f"max difference {smax:.1e} over {nrows} rows; row 133 max {r133:.1e}"),
    ("0 new errors; Logistics&WH computes in-book", f"new error cells {len(err10 - err9)}; Logistics&WH errors {res['lw_errors']}; total "
     f"{F['lw']['total_2027']:,.2f}; B60 {F['lw']['check_B60']:.0f}"),
    ("LOG-01 not double counted", f"CO 6610-6640 - HC-CO roster - (Logistics&WH less row 34) = {co_chk:.1e} (each GL)"),
    ("MH-01 only in CO", f"PO B131 base drop - PO HC r20 cost = {res['fig']['po_base_drop_minus_r20']:.1e}; booked PO 5110-5140 = "
     f"{F['po']['booked_v10']:.0f}; MH-01 in CO 6610-6640 via the model"),
    ("Diff confined (v9 -> v10)", f"outside the allowed areas {len(outside)}; 2026 columns {cols26}; HC-COO roll-up reported separately"),
    ("Every new value is a formula", "CO U67:AF69 and PO U33:AF35 are formulas; Logistics&WH keeps the SD formulas; HC-CO r18 holds the "
     "roster inputs (name, title, salary, start) like every roster row"),
    ("Every figure computed by script", "figures_v10.py from the workbook and 7 LibreOffice scenario recalculations; carried v9 effects listed per row"),
    ("Stored values = LibreOffice recalculation", f"max {mx:.1e} over {ncells:,} numeric cells; non-numeric {len(nonnum)}"),
    ("No names in the md / Collation Notes; per-person amounts rounded in the md", "privacy_v6 hits 0; paylist exact hits 0"),
])
res["fails"] = fails
json.dump(res, open(OUT, "w"), indent=1, default=str)
print(f"analyze_v10: areas {len(areas)}, outside {len(outside)}, 2026 cols {cols26}, new errors {len(err10 - err9)}, recalc max {mx:.1e}, "
      f"sum-of-depts max {smax:.1e}, CN formulas {ncn}; fails {len(fails)}")
for f in fails:
    print("FAIL:", f)
sys.exit(1 if fails else 0)
