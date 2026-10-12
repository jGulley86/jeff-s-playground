"""v14 build: documentation corrections on v13 - Collation Notes TEXT cells only; no budget value changes.

Handoff 2026-10-12 (orchestrator): fix critic v13 B1 + O1-O4 and the verifier's stale v12-era text; AG132 stays
(646,753.85). Every edit below is a string cell on the Collation Notes sheet (plus the appended v14 block rows
876-888). The deliverable is produced later by ONE full LibreOffice recalculation + engine save (the v12/v13
technique); diff_v14.py proves the v13 -> v14 diff confined to the Collation Notes sheet.

Edits (critic v13 / verifier v13):
  B1  E13, C41, D381, H381 (+ E381, B858): the F20 = Dec-27 flip ALONE creates an expense + capitalised overlap
      (HC-PO r16 stays in the 8-person B131 base; the Prod Payroll model counts by number; asset overstated ~70k),
      and the left-at-Nov-26 branch (r16 72,060.60/yr; B131 = 0 base 590,317.51 -> 518,256.91; the (1,205,750.10)
      scenario and LB2 (1,714,922) 72,060.60 too negative; 9-of-9: asset 70,056, 2027 depreciation 2,821).
      H381's "budget the role outside the capitalised labour (set F20=Dec-27)" corrected to the three-way ask.
  O1  A255: HC-COO double-lists the Team Lead person Sep-Nov-26 (2026-side only; 2027 unaffected).
  O2  G90, B857: MFS <-> MH-02 asymmetry (out 61,962.70 at model rates, in 100,304.86 at 85,000 + CO HC burden,
      net +38,342.16 inside the B+D (71,048.29)).
  O3  B857: A-46 reversal recipe (the -'Logistics&WH Payroll'!<col>39 term per month column of CO rows 67-69,
      x B5/B6 on 68/69; the H78:H80 co-edit; flipping back ADDS 61,962.70 of cost).
  O4  A33, A71, A873: one hires/moves count - 11 (4 CO, 5 PE, 1 PO, 1 FE) with stated scope.
  V   E17, A32, A35, C43, C47, C60, C61, F357, D377:G377, D380, D8, A114, A1 (+ C323, C336, sweep-found,
      same-vintage): stale v12-era text -> v13 values.
  New v14 block appended at rows 876-888 (no row insertions; every existing address stays stable).

No clock, no file properties, no dates are read (F1 rule). Names are neither read nor written.

Usage:
  python3 -I build_v14.py V13_XLSX OUT_XLSX
"""
import sys, os, zipfile
import openpyxl


def pin_coreprops(ref_xlsx, out_xlsx):
    """openpyxl force-stamps docProps/core.xml 'modified' with the wall clock at save, which would make the build
    date-dependent (F1 rule) and the two passes differ. Replace the script-written INTERMEDIATE's core.xml with the
    reference file's bytes (zip-level, fixed entry date). The deliverable is the later LibreOffice engine save,
    which carries docProps unchanged (v12-proven), so the shipped docProps stay v13's."""
    core = zipfile.ZipFile(ref_xlsx).read("docProps/core.xml")
    tmp = out_xlsx + ".tmp"
    with zipfile.ZipFile(out_xlsx) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in zi.infolist():
            if info.filename == "docProps/core.xml":
                zo.writestr(zipfile.ZipInfo("docProps/core.xml", date_time=(1980, 1, 1, 0, 0, 0)), core)
            else:
                zo.writestr(info, zi.read(info.filename))
    os.replace(tmp, out_xlsx)


N = lambda v: float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def preflight(wb, wbv):
    """Assert the base is v13 (cached values via wbv = data_only load; the old texts this build replaces via wb),
    so edits land on the right cells. All numeric constants below are verifier-confirmed (04-verify_report_v13.md)
    or recomputed here from the book."""
    cn, cnv = wb["Collation Notes"], wbv["Collation Notes"]
    ok_v = lambda ad, exp, tol=0.01: abs(N(cnv[ad].value) - exp) <= tol or _die(f"{ad} != {exp}")
    ok_t = lambda ad, sub: (isinstance(cn[ad].value, str) and sub in cn[ad].value) or _die(f"{ad} lacks {sub!r}")

    def _die(msg):
        raise AssertionError("preflight: " + msg)

    assert abs(N(wbv["COO P&L"]["AG132"].value) - (-646753.85)) <= 0.01, "AG132 != (646,753.85)"
    for ad, exp in (("D41", -72060.60), ("D42", 2821.18), ("D43", 10105.91), ("D60", -91476.67),
                    ("D61", -33629.76), ("B867", -1205750.10), ("B868", -773662.29), ("B869", -707448.94),
                    ("B870", -1286153.81), ("B871", -1239164.67), ("E94", 166414.40), ("F94", 210022.54)):
        ok_v(ad, exp)
    # recompute the O2 asymmetry and I-37 components from the book (never printed; guards the hard-coded text)
    co = wbv["HC-CO"]
    t, b = N(co["D6"].value), N(co["D7"].value)
    r35 = sum(N(co.cell(35, c).value) for c in range(13, 25)) * (1 + t + b)
    assert abs(r35 - 100304.86) <= 0.02, "r35 2027 cost != 100,304.86"
    lw = wbv["Logistics&WH Payroll"]
    mh02 = sum(N(lw.cell(39, c).value) for c in range(6, 18)) * (1 + N(lw["B5"].value) + N(lw["B6"].value))
    assert abs(mh02 - 61962.70) <= 0.02, "MH-02 2027 cost != 61,962.70"
    assert abs((r35 - mh02) - 38342.16) <= 0.03, "net swap != 38,342.16"
    gap = (N(co["Z67"].value) * (0.0765 - t)                                    # CO roster       20,906
           + 166414.40 * (0.0765 - N(lw["B5"].value))                           # 5 model seats    1,415
           + N(wbv["HC-PE"]["Z67"].value) * (0.0765 - N(wbv["HC-PE"]["D6"].value))  # PE roster    64,102
           + sum(N(wbv["HC-PO"].cell(41, c).value) for c in range(13, 25))
           * (0.0765 - N(wbv["HC-PO"]["D6"].value)))                            # PM               5,054
    assert abs(gap - 91476.67) <= 1.0, "I-37 CO+logistics+PE+PM gap != 91,477"
    assert abs((3036558 + 919884) * 0.0085 - 33629.76) <= 1.0, "FO/FE gap basis"
    assert abs(-646753.85 - 2448290 - (-3095043.85)) <= 0.01, "2026-policy contribution"
    assert abs(590317.51 - 72060.60 - 518256.91) <= 0.005, "base less r16"
    # the old texts about to be replaced/extended (fail loudly if the base is not v13)
    ok_t("A1", "(build v13, draft for review)")
    ok_t("D8", "Decision #6).")
    ok_t("E13", "+67,917 CO cost) if the person should be costed in CO through 2027")
    ok_t("E17", "(399,860)")
    ok_t("A32", "(2,289,154)")
    ok_t("A33", "13 hires/moves")
    ok_t("A35", "'Carried from v9' have inputs v10 does not touch")
    ok_t("C41", "Flip HC-CO F20 to Dec-27 to cost the person in CO instead (one cell).")
    ok_t("C43", "319,393")
    ok_t("C47", "398,492 in the headline")
    ok_t("C60", "CO HC 15,525 + logistics 2,715 + PE 43,967")
    ok_t("C61", "3,100,031")
    ok_t("A71", "4 CO adds, 5 PE adds, the PO Production Manager, the FE Associate")
    ok_t("G90", "excluded by formula")
    ok_t("A114", "02-build_notes_v13.md")
    ok_t("A255", "costs him once, in FE.")
    ok_t("C323", "PO/CO/PE gap up to 77,882")
    ok_t("C336", "+19,396 at CO HC rates (caveat)")
    ok_t("F357", "(398,492) booked in the headline")
    ok_t("D377", "(62,207) (v9 (56,801))")
    ok_t("E377", "PO +21,081")
    ok_t("F377", "(34,169)")
    ok_t("G377", "impact 62,207")
    ok_t("D380", "-> (399,860)")
    ok_t("D381", "set HC-CO F20 = Dec-27 (one cell, +67,917 of CO cost at CO HC rates).")
    ok_t("E381", "The 2 people together at PO HC rates: 134,563")
    ok_t("H381", "budget the role outside the capitalised labour (set HC-CO F20 = Dec-27, or keep an expense line).")
    ok_t("B857", "re-add row 39 to CO!U67:AF69 (+61,962.70).")
    ok_t("B858", "HC-CO F20 = Dec-27 (one cell, +67,917.01 at CO HC rates). See I-41.")
    ok_t("A873", "10 user-approved 2026-10-12 hires/moves (4 CO, 5 PE, 1 PO)")
    assert cn["A874"].value is not None and cn["A875"].value in (None, "") and cn["A876"].value in (None, ""), \
        "v14 block target rows 876+ not empty"


# ------------------------------------------------------------------------------------------- the corrected texts
SET = {
 "A1": "2027 COO Budget - Collation Notes (build v14 = v13 figures unchanged + documentation corrections; draft "
       "for review)",
 "E13": "Resolved. I-41 restated in v13, completed in v14: the Material Handling Team Lead is on HC-PO r16 (in the "
        "capitalised B131 base and, by headcount, in the by-number Prod Payroll model) and on HC-CO r20 with end "
        "month Nov-26 (2027 cost 0, the user's file value); no 2027 expense line carries the role and seat TL-01 is "
        "excluded: (72,061) if it should stay an expense. Flipping HC-CO F20 = Dec-27 ALONE books +67,917 in CO but "
        "removes nothing from the capitalised side (r16 stays in the 8-person B131 base; the model counts by "
        "number), so the role would be carried twice - expense + capitalised robot cost - with the asset overstated "
        "~70k. If instead the person LEFT at Nov-26 (the backfill reading), r16, the live B131 = 0 base (590,317.51 "
        "-> 518,256.91 without r16), lower bound 2 and the by-number model all carry a departed person: the B131 = "
        "0 scenario (1,205,750) and LB2 (1,714,922) are then 72,061 too negative, and on the 9-of-9 reading the "
        "asset is ~70,056 too high (2027 depreciation 2,821). See I-41 (D381/H381). Logistics at CO HC burden "
        "rates: +10,106.",
 "E17": "Current staff expensed (sensitivity): (558,996) -> (1,205,750.10). Idle-month labour expensed (B135 = 0): "
        "(126,908) -> (773,662.29) (with the current staff expensed too: (1,286,153.81)). All Sep-Dec 26 labour "
        "capitalised (B132 = 1; only with a 2026 restatement): (60,695) -> (707,448.94). Each is one switch on the "
        "Depreciation Schedule; the workbook recalculates. (v14: endpoints restated on the v13 headline; the deltas "
        "are unchanged.)",
 "A32": "(0b) Year on year is not like-for-like: 2026 expensed production labour, 2027 capitalises it. Policy "
        "effect carried from v10 (the labour block is untouched in v11-v13): +2,448,290 (PO 5110-5140 at B131 = 0 "
        "590,318 + the ramp 1,962,389 - booked labour depreciation 104,417; v9 +2,510,793): on the 2026 policy the "
        "2027 contribution would be about (3,095,044).",
 "A33": "(1) What the figure is. The bottom line, (646,754), is a COO contribution: the company's recognized "
        "revenue (15,563,646, all of it) less the costs of the five COO departments only (COGS + Opex of FO, PO, "
        "CO, FE, PE: 16,210,399). Costs outside COO are not in this workbook, so it is not company Net Income. The "
        "workbook row keeps its GL label 'Net Income' (COO P&L!A132). It is NEGATIVE in v13: the user-approved "
        "2026-10-12 additions - 11 hires/moves (4 CO, 5 PE, 1 PO, 1 FE; counted as the roster rows added on the "
        "four HC tabs: the Material Handler's move counts once, as its CO row, and the HC-CO r18 title edit and the "
        "HC-PO r20 blanking are not additions) plus two CO GL lines - put 2027 COO costs above the plan's "
        "recognized revenue.",
 "A35": "(3) Open exposures, each on its own against the headline, (646,754). 'n/d' = not determinable from the "
        "workbook. Rows marked 'Carried from v9' / 'recalculated on v10' keep those provenance labels: their inputs "
        "are untouched by every later build (the v9 -> v13 diffs show 0 changes on them), so each effect stands "
        "against the v13 headline (column E = headline + effect). The scenario-switch rows, I-41, I-37 and the "
        "burden-rate row are recomputed on v13.",
 "C41": "v13: the Material Handling Team Lead sits on HC-PO r16 (inside the capitalised B131 base, 0 expense while "
        "B131 = 1; the Prod Payroll model also counts its current staff by number) AND on HC-CO r20 end-dated "
        "Nov-26 (2027 cost 0, the user's file value); seat TL-01 is excluded from the CO add-on as realized on the "
        "roster. So no 2027 expense line carries the role. Three branches (v14): (1) warehouse role, person STAYS "
        "(7-of-9 reading): its 2027 cost at PO HC rates should be an expense - this row, (72,061); flipping HC-CO "
        "F20 to Dec-27 ALONE is not that fix - it books +67,917 in CO while r16 stays in the 8-person B131 base and "
        "the by-number model keeps capitalising the role: carried twice (expense + robot cost), asset overstated "
        "~70k. (2) Person LEFT at Nov-26 (the backfill reading; the Material Flow Supervisor, 85,000 Nov-26 in CO, "
        "replaces the warehouse role): no booked expense effect, but r16, the live B131 = 0 base (590,317.51 -> "
        "518,256.91 without it), LB2 and the by-number model all carry a departed person - the (1,205,750) scenario "
        "and LB2 (1,714,922) are 72,061 too negative, and on the 9-of-9 reading the asset is ~70,056 too high "
        "(2027 depreciation 2,821). (3) Production role, person stays: booked treatment right (capitalised), 0.",
 "C43": "v10 books the logistics seats at the model's own rates (payroll tax 6.8%, benefits 16.4%: 'the WH and "
        "logistics numbers'). CO HC rates are 5.22% / 11.91%. v13: 5 model seats remain; on their 166,414 of 2027 "
        "logistics salaries in CO: burden 38,608 booked vs 28,502 at CO HC rates. Both CO HC rates are unverified "
        "(I-37).",
 "C47": "RESOLVED (user decision 2026-10-10): the Logistics&WH model is booked in CO 6610-6640, less seat LOG-01 "
        "(already on CO HC r17). v13: less the roster-realized MH-01/TL-01/MH-02 as well - 210,023 in the headline "
        "(5 model seats + 5,000 new-hire cost); v10-v12 (398,492) (the 7 new seats 336,530 + MH-01 61,963 at model "
        "rates). v9 showed it as a downside of (336,530) to (471,093) depending on the reading.",
 "C60": "CO, PE and the PO Production Manager's payroll tax below 7.65% (upper bound on every salary dollar). v13 "
        "base: the CO HC roster (incl. the 2026-10-12 adds) at the CO HC rate (5.22%), the 5 remaining logistics "
        "seats at the model rate (6.8%), the PE roster (incl. 5 adds) at 3.57%, and the PM at the PO rate (3.64%): "
        "CO HC 20,906 + logistics 1,415 + PE 64,102 + PM 5,054 = 91,477. v12: (62,207); v9: (56,801).",
 "C61": "FO and FE payroll tax uses the SD field rate 6.8%, also below 7.65% (same upper-bound basis; FO salary "
        "base 3,036,558 incl. the ramp - one tech out since v13, B11 = 13 - and FE 919,884).",
 "A71": "Conclusion: the user additions of 2026-10-12 (brief v7: CO 7700 +24,000 and 8600 +78,000; brief v8: the HC "
        "roster merge - 11 hires/moves: 4 CO, 5 PE, 1 PO (the Production Manager), 1 FE (the Associate; roster row "
        "only - he was already costed in FE), counted as roster rows added on the four HC tabs, the CO Team Lead "
        "row with no 2027 cost among them; I-28 resolved) move the 2027 COO contribution from 159,136 (v12) to "
        "(646,754): NEGATIVE at the booked policy. Bridge: CO GL adds (102,000); CO payroll net (71,048); PE "
        "payroll (554,865); PO Production Manager (~155,000) (one person; exact in the v13 block below); I-28 "
        "resolution +78,529. Treat (646,754) as a point estimate inside (1,275,819) to 706,820 (lower bound 2: "
        "(1,714,922)) until Decisions #2 (I-32) and #10 (I-40) are answered, I-10 is confirmed, and the PO owner "
        "answers the restated I-41 three-way (left at Nov-26 / warehouse / production - see H381).",
 "A114": "Full notes: 02-build/02-build_notes_v14.md (= the v13 notes with the v14 documentation corrections "
         "applied; all figures unchanged). Figures: figures_v13.py (from the workbook, the v12 workbook and 5 "
         "LibreOffice scenario recalculations); v13 scenario values recomputed, carried effects marked per row.",
 "C323": "v13: CO/PE/PM gap up to 91,477; FO/FE up to 33,630 (v7: PO/CO/PE 77,882, FO/FE 34,169 - figures_v7)",
 "D323": "I-37 row (v13 recompute); figures_v7.py (original)",
 "C336": "+10,106 at CO HC rates on the 5 remaining seats (caveat; v10-v12: +19,396 on 8 seats)",
 "F357": "(210,023) booked in the headline (v13: 5 model seats + 5,000 new-hire cost; v10-v12 (398,492))",
 "D377": "The HC T&B employer payroll tax rates (A-30, stored values) are all below the 7.65% employer FICA (6.2% "
         "Social Security + 1.45% Medicare): FO 6.847%, PO 3.637%, CO 5.221%, FE 4.069%, PE 3.571%. A rate below "
         "7.65% is possible only if salaries above the Social Security wage base pull the average down, and the "
         "rates would normally also carry unemployment taxes. They are unverified. CO and PE payroll tax in the "
         "budget uses these rates, and since v13 the PO Production Manager is expensed at the PO rate, so these "
         "lines may be under-budgeted; FO and FE use the SD rate 6.8%, also below 7.65%. FE's I-34 upside rests on "
         "the FE rate. v8: PO production payroll is capitalised (PO 5110-5140 = 0), so the capitalised labour's "
         "part is no longer an expensed P&L risk; the model's own 7.5% applies. v13 bases (the roster adds, the "
         "live PE wiring, the PM): CO + PE + PM gap at 7.65% (91,477) (v12 (62,207), v9 (56,801)).",
 "E377": "At 7.65% on every 2027 salary dollar (upper bound), v13 bases: CO HC roster +20,906, logistics seats (5, "
         "at 6.8%) +1,415, PE +64,102, PO Production Manager +5,054 = +91,477 of tax; FO +25,811 and FE +7,819 "
         "against the SD 6.8% (FO salary base 3,036,558, FE 919,884). I-34 at an FE tax rate of 7.65%: +76,600 "
         "instead of +109,539.",
 "F377": "up to (91,477) CO/PE/PM (v12 (62,207)); up to (33,630) FO/FE",
 "G377": "impact 91,477 ($50k-250k)",
 "D380": None,  # built below from the old text (only the stale tail sentence is replaced)
 "D381": "v13 restatement, completed in v14. The two material-handling roles of PO HC: r20 Material Handler is in "
         "CO (HC-CO r19; v10 decision, v13 roster move completed). r16 Material Handling Team Lead stays on HC-PO "
         "(inside the live B131 base = the 8 production staff; the Prod Payroll model also counts its current staff "
         "by number) AND is listed on HC-CO r20 with end month Nov-26 (2027 cost 0; the user's file value, kept per "
         "the brief). Booked (B131 = 1): PO 5110-5140 = 0, so no expense line carries r16, and seat TL-01 is "
         "excluded from the CO add-on as realized on the roster. Three branches: (1) r16 is warehouse staff and the "
         "person STAYS (7-of-9 reading): its cost should stay an expense, (72,061); setting HC-CO F20 = Dec-27 "
         "ALONE is not that fix - it books 67,917.01 of CO expense while r16 stays in the B131 base and the "
         "by-number model keeps capitalising the role: carried twice (expense + robot cost), asset overstated ~70k. "
         "(2) The person LEFT at Nov-26 (the backfill reading - the Material Flow Supervisor replaces the warehouse "
         "role in CO): booked expense unchanged, but r16, the live B131 = 0 base (590,317.51; 518,256.91 without "
         "r16), LB2 and the by-number model all carry a departed person: the B131 = 0 scenario (1,205,750.10) and "
         "lower bound 2 (1,714,922) are 72,060.60 too negative, and on the 9-of-9 reading the capitalised asset is "
         "~70,056 too high (2027 depreciation 2,821). (3) r16 is production staff and stays: the booked treatment "
         "(capitalised; under B131 = 0 expensed inside PO 5110-5140) is right.",
 "H381": "PO owner, three-way: did the person leave at Nov-26 (the user's HC-CO end date), stay in a warehouse "
         "role, or stay in production? If WAREHOUSE and staying: budget the role as an expense AND take it out of "
         "the capitalised side - HC-CO F20 = Dec-27 alone only adds +67,917 of CO expense (r16 stays in the B131 "
         "base and the by-number model: carried twice, asset overstated ~70k); r16 must also leave the capitalised "
         "labour (end-date/move r16 and, on the 9-of-9 reading, one head off Prod Payroll's current-staff count). "
         "If LEFT at Nov-26: end-date HC-PO r16 at Nov-26 (B131 = 0 base 590,318 -> 518,257) and take the departed "
         "head out of the by-number model (asset (70,056); 2027 depreciation +2,821). If PRODUCTION and staying: "
         "keep the booked treatment (capitalised; no expense line).",
 "A873": "YoY footnote (extended for the new hires): 2026 vs 2027 is not like-for-like. 2026 expensed production "
         "labour (2027 capitalises it), and the 2027 payroll adds the user-approved 2026-10-12 roster additions - "
         "11 hires/moves (4 CO, 5 PE, 1 PO, 1 FE), counted as roster rows added on the four HC tabs (the FE row "
         "feeds no P&L line; the CO Team Lead row has no 2027 cost) - plus the CO 7700/8600 lines that have little "
         "or no 2026 spend. The Sep-Dec 26 costs of the new starters are on the HC tabs only - the 2026 P&L columns "
         "are book data and were not restated. On the 2026 policy AND before these additions the 2027 contribution "
         "would differ by both effects; see the md for the policy-effect recompute.",
}

APPEND = {
 "D8": " v14 note: the current contribution is (646,754) (v13, the user's 2026-10-12 additions - see the v13 block "
       "below); the capitalisation policy itself is unchanged since v9.",
 "A255": " v14 addition to this inventory (critic v13 O1): the Material Handling Team Lead is listed on BOTH HC-PO "
         "r16 (Sep-26..Dec-27) and HC-CO r20 (Sep-26..Nov-26), so HC-COO's roster-level payroll and headcount count "
         "that person twice in Sep-Nov-26 (2026-side columns only); from Dec-26, and for all of 2027, the HC-CO row "
         "has ended and costs 0 - 2027 totals are unaffected.",
 "G90": " v14 note - the swap is asymmetric: the seat goes out at the model's 61,962.70 (this row) while the "
        "realized backfill comes in on the roster at 85,000 salary + CO HC burden = 100,304.86 of 2027 cost; net "
        "+38,342.16 of cost inside the B+D (71,048.29) term.",
 "B857": " v14 (reversal recipe + asymmetry): to flip A-46 back, delete the -'Logistics&WH Payroll'!<col>39 term in "
         "each of the 12 month columns U:AF of CO row 67 and the row-39 term inside the bracketed subtraction of "
         "rows 68/69 (it is multiplied by B5/B6 there), and co-edit the H78:H80 check formulas on this sheet (their "
         "S39 terms) the same way; flipping back ADDS the 61,962.70 of 2027 cost (the contribution falls by it). "
         "The booked swap is asymmetric: MH-02 out at the model's 61,962.70, the Material Flow Supervisor in at "
         "85,000 + CO HC burden = 100,304.86; net +38,342.16 inside the B+D (71,048.29).",
 "B858": " v14: the flip books the person in CO but removes nothing from the capitalised side (HC-PO r16 stays in "
         "the B131 base; the Prod Payroll model counts by number) - see I-41 for the three branches, incl. the "
         "left-at-Nov-26 branch (B131 = 0 base 590,317.51 -> 518,256.91 without r16).",
 "E381": " v14: r16 alone 72,060.60/yr at PO HC rates (the restated exposure); the HC-CO F20 = Dec-27 flip books "
         "67,917.01/yr at CO HC rates; B131 = 0 base without r16: 518,256.91.",
}

D380_OLD_TAIL = ("v10: the B131 base excludes PO HC r20 (MH-01, in CO): sensitivity (558,996) -> (399,860) "
                 "(v9 (621,499) -> 65,936).")
D380_NEW_TAIL = ("v10: the B131 base excludes PO HC r20 (MH-01, in CO). v13 (headline (646,754)): sensitivity "
                 "(558,996) -> (1,205,750) (v9 (621,499) -> 65,936; v10-v12 -> (399,860) on the then-headline).")

V14_BLOCK = [  # rows 876-888, columns A..C
 ("Changes from v13 (v14) - documentation corrections only; NO budget value changes (AG132 unchanged (646,753.85); "
  "the v13 -> v14 diff is confined to this sheet - 0 formula and 0 value changes anywhere else)", None, None),
 ("Sources: critic v13 (B1 blocking + O1-O4) and verifier v13 (stale v12-era text; md headline GP row; "
  "severed-reference wording). Every v13 figure, scenario value and bridge term is unchanged and holds for v14.",
  None, None),
 ("Item", "Fix", "Cells (this sheet); md = 02-build_notes_v14.md"),
 ("B1 (critic)",
  "I-41/A-47 completed: (i) flipping HC-CO F20 = Dec-27 ALONE books 67,917.01 in CO while HC-PO r16 stays in the "
  "8-person B131 base and the by-number Prod Payroll model keeps capitalising the role - carried twice, asset "
  "overstated ~70k (model avg 70,056/head); (ii) new left-at-Nov-26 branch: r16 (72,060.60/yr), the live B131 = 0 "
  "base (590,317.51 -> 518,256.91 without r16), the (1,205,750.10) scenario, LB2 (1,714,922) and the by-number "
  "model all carry a departed person (9-of-9: asset 70,056, 2027 depreciation 2,821); (iii) H381's 'budget the "
  "role outside the capitalised labour (set F20=Dec-27)' was inaccurate - corrected to the three-way ask.",
  "E13, C41, D381, H381, E381, B858; md I-41, A-47, caveat (iii), roster-merge note"),
 ("O1", "HC-COO double-listing added to the inventory: the Team Lead person is on HC-PO r16 and HC-CO r20, so "
        "HC-COO's roster payroll and headcount count the person twice Sep-Nov-26 (2026-side only; 2027 unaffected "
        "- the CO row costs 0).", "A255; md HC-COO roll-up"),
 ("O2", "MFS <-> MH-02 asymmetry stated: seat out at the model's 61,962.70, hire in at 85,000 + CO HC burden = "
        "100,304.86, net +38,342.16 of cost inside the B+D (71,048.29).", "G90, B857; md decision 2"),
 ("O3", "A-46 reversal recipe: delete the -'Logistics&WH Payroll'!<col>39 term in each month column U:AF of CO row "
        "67 and the row-39 term in the bracketed subtraction of rows 68/69 (x B5/B6), co-edit H78:H80 (S39 terms); "
        "flipping back ADDS 61,962.70 of 2027 cost.", "B857; md A-46"),
 ("O4", "One hires/moves count: 11 (4 CO, 5 PE, 1 PO, 1 FE) = roster rows added on the four HC tabs (the HC-CO r18 "
        "title edit and the HC-PO r20 blanking are not counted; the MH move counts once, as the CO row; the FE row "
        "feeds no P&L line; the CO TL row has no 2027 cost).", "A33, A71, A873; md footnote (a) + conclusion"),
 ("Verifier", "md headline Gross Profit row corrected: 2027 v12 column 3,960,032 and v13 - v12 change +78,529 "
              "(was 4,038,561 / 0).", "md only"),
 ("Verifier", "Stale v12-era text replaced with v13 values: E17 (scenario endpoints (1,205,750.10) / (773,662.29) "
              "/ (1,286,153.81) / (707,448.94)); A32 ((3,095,044) on the 2026 policy); A35 (provenance note); C43 "
              "(5-seat logistics base 166,414 / 38,608 / 28,502); C47 + F357 (I-17 booked 210,023); C60 + row 377 "
              "(I-37 91,477 = CO 20,906 + logistics 1,415 + PE 64,102 + PM 5,054); C61 (FO base 3,036,558); D380 "
              "((1,205,750) scenario); D8 (current-contribution note); C323 + C336 (same-vintage rows, found by "
              "extending the sweep).", "listed cells"),
 ("Verifier", "Severed-reference claims rescoped ('0 anywhere' was imprecise): 0 in formulas, caches and defined "
              "names; 6 historical text mentions of the Excel error marker remain on this sheet (E343, D364, D365, "
              "E365 - v2/v3 cleanup records; this row avoids the literal marker to keep that count exact).",
  "md wording"),
 ("Sweep", "Scripted sweep (sweep_v14.py) of rows 8-71, 357, 375-381 (+ 323/336, builder-extended) for the known "
           "stale values: pass 1 (on v13) flagged the cells above; pass 2 (on this file) clean.",
  "scripts/run_build_v14.sh"),
 ("Open (user)", "Not acted on - awaiting user answers: (1) PO!R24:R26 production freight 'YOY Growth' in the "
                 "user's edited file vs 'Even'/0 carried since v1 (would add 60,056.60 of 2027 COGS; AG132 = about "
                 "(706,810)); (2) HC-PE!A22 'transfer?' flag blank in the user's file (no number change; I-35's "
                 "~255k upside rests on it); (3) the I-41 three-way (H381).", "md 'Not done / limits'"),
]


def main(v13_path, out_path):
    wb = openpyxl.load_workbook(v13_path)
    wbv = openpyxl.load_workbook(v13_path, data_only=True)
    preflight(wb, wbv)
    cn = wb["Collation Notes"]
    n = 0
    for ad, txt in SET.items():
        if ad == "D380":
            old = cn["D380"].value
            assert old.endswith(D380_OLD_TAIL), "D380 tail changed"
            txt = old[: -len(D380_OLD_TAIL)] + D380_NEW_TAIL
        cn[ad] = txt
        n += 1
    for ad, tail in APPEND.items():
        cn[ad] = cn[ad].value + tail
        n += 1
    r = 876
    for row in V14_BLOCK:
        for i, v in enumerate(row):
            if v is not None:
                cn.cell(row=r, column=1 + i, value=v)
        r += 1
    wb.save(out_path)
    pin_coreprops(v13_path, out_path)
    print(f"build_v14: {n} Collation Notes cells updated in place + v14 block rows 876-{r - 1} -> {out_path} "
          "(docProps pinned to v13's; no other sheet touched)")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
